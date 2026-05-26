from __future__ import annotations

import io
import re
import zipfile

from app.modules.ai_scoring.firewall import extract_pdf_text_with_firewall
from app.modules.ai_scoring.schema import TextExtractionResult


def extract_resume_text(content: bytes, file_type: str) -> TextExtractionResult:
    if file_type == "pdf":
        return extract_pdf_text_with_firewall(content)
    if file_type == "docx":
        return _extract_docx_text(content)
    if file_type == "doc":
        return _extract_doc_text(content)
    raise ValueError(f"Unsupported resume file type: {file_type}")


def _extract_docx_text(content: bytes) -> TextExtractionResult:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to analyze DOCX resumes") from exc

    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_cells: list[str] = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    table_cells.append(text)

    text = "\n".join([*paragraphs, *table_cells])
    return TextExtractionResult(
        extracted_text=text,
        sanitized_text=text,
        parser_warnings=[
            {
                "type": "limited_firewall_coverage",
                "message": "DOCX text was extracted, but PDF hidden glyph checks do not apply to this format.",
            }
        ],
    )


def _extract_doc_text(content: bytes) -> TextExtractionResult:
    text = _extract_doc_text_with_olefile(content)
    if not text.strip():
        text = _extract_binary_text_fallback(content)
    if not text.strip():
        raise RuntimeError("Unable to extract text from DOC resume")
    return TextExtractionResult(
        extracted_text=text,
        sanitized_text=text,
        parser_warnings=[
            {
                "type": "limited_firewall_coverage",
                "message": "DOC text was extracted, but PDF hidden glyph checks do not apply to this format.",
            },
            {
                "type": "legacy_doc_parser",
                "message": "Legacy DOC extraction may be less precise than PDF or DOCX extraction.",
            },
        ],
    )


def _extract_doc_text_with_olefile(content: bytes) -> str:
    try:
        import olefile
    except ImportError:
        return ""
    if not olefile.isOleFile(io.BytesIO(content)):
        return ""
    try:
        with olefile.OleFileIO(io.BytesIO(content)) as ole:
            if not ole.exists("WordDocument"):
                return ""
            data = ole.openstream("WordDocument").read()
    except OSError:
        return ""
    return _extract_binary_text_fallback(data)


def _extract_binary_text_fallback(content: bytes) -> str:
    decoded_parts: list[str] = []
    for encoding in ("utf-8", "utf-16le", "latin-1"):
        decoded = content.decode(encoding, errors="ignore")
        cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]+", " ", decoded)
        cleaned = re.sub(r"\s{2,}", " ", cleaned)
        if len(cleaned.strip()) > 40:
            decoded_parts.append(cleaned.strip())
    return max(decoded_parts, key=len, default="")


def detect_resume_file_type(url: str, content_type: str | None, content: bytes) -> str:
    lower_url = url.lower().split("?", 1)[0]
    lower_content_type = (content_type or "").lower()
    if lower_content_type == "application/pdf" or lower_url.endswith(".pdf") or content.startswith(b"%PDF"):
        return "pdf"
    if (
        "wordprocessingml.document" in lower_content_type
        or lower_url.endswith(".docx")
        or zipfile.is_zipfile(io.BytesIO(content))
    ):
        return "docx"
    if (
        lower_content_type in {"application/msword", "application/vnd.ms-word"}
        or lower_url.endswith(".doc")
        or content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    ):
        return "doc"
    return "unsupported"

