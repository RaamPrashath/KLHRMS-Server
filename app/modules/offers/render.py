from __future__ import annotations

import asyncio
import base64
import mimetypes
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from html import escape, unescape
from html.parser import HTMLParser
from io import BytesIO
from typing import Any
from urllib.parse import urlparse

import httpx

from app.models.organization import Organization
from app.models.recruitment import Candidate, JobPosting, OfferLetter
from app.modules.offers.schema import (
    ALLOWED_VARIABLE_TOKENS,
    COMPENSATION_VARIABLE_TOKENS,
    VARIABLE_PATTERN,
)

A4_WIDTH_PX = 794
A4_HEIGHT_PX = 1123
PAGE_PADDING_TOP_PX = 90
PAGE_PADDING_BOTTOM_PX = 90
PAGE_PADDING_X_PX = 76
BODY_FONT_SIZE_PX = 14
BODY_LINE_HEIGHT_PX = 27
CORNER_MARK_TOP_PX = 28
CORNER_MARK_LEFT_PX = 28
CORNER_MARK_HEIGHT_PX = 68
CORNER_MARK_WIDTH_PX = 19
BRAND_GAP_FROM_CORNER_MARK_PX = 20
LOGO_MAX_HEIGHT_PX = 56
CONTENT_WIDTH_PX = A4_WIDTH_PX - PAGE_PADDING_X_PX * 2
CONTENT_HEIGHT_PX = A4_HEIGHT_PX - PAGE_PADDING_TOP_PX - PAGE_PADDING_BOTTOM_PX
HEADER_BRAND_LEFT_PX = (
    CORNER_MARK_LEFT_PX
    + CORNER_MARK_WIDTH_PX
    + BRAND_GAP_FROM_CORNER_MARK_PX
    - PAGE_PADDING_X_PX
)
HEADER_BRAND_TOP_PX = (
    CORNER_MARK_TOP_PX
    + (CORNER_MARK_HEIGHT_PX - LOGO_MAX_HEIGHT_PX) / 2
    - PAGE_PADDING_TOP_PX
)
MAX_INLINE_IMAGE_BYTES = 5 * 1024 * 1024
IMG_SRC_PATTERN = re.compile(r'(<img\b[^>]*\bsrc=)(["\'])(.*?)(\2)', re.IGNORECASE | re.DOTALL)
PARAGRAPH_PATTERN = re.compile(r"<p\b[^>]*>(.*?)</p>", re.IGNORECASE | re.DOTALL)
HEADING_PATTERN = re.compile(r"<h[1-6]\b[^>]*>(.*?)</h[1-6]>", re.IGNORECASE | re.DOTALL)
TAG_PATTERN = re.compile(r"<[^>]+>")
DEFAULT_KOVAN_LOGO_SVG = (
    "data:image/svg+xml;base64,"
    "PHN2ZyB4bWxucz0iaHR0cDovL3d3dy53My5vcmcvMjAwMC9zdmciIHdpZHRoPSIxNzYiIGhlaWdodD0iNTYiIHZpZXdCb3g9IjAgMCAxNzYgNTYiPjxwYXRoIGQ9Ik0xIDFoMTl2NTRIMXoiIGZpbGw9IiMwMDg3NEEiLz48cGF0aCBkPSJNMzUgMTVoMTN2MjZIMzV6IiBmaWxsPSIjMDBBNjU0Ii8+PHBhdGggZD0iTTUwIDE1aDEzdjI2SDUweiIgZmlsbD0iI0VBNDMzNSIvPjxwYXRoIGQ9Ik02NSAxNWgxM3YyNkg2NXoiIGZpbGw9IiNGQkJDMDEiLz48cGF0aCBkPSJNNzcuNSAyOGwtMTIgMTNWMzNMMzUgNDFWMTVsMzAgMjYgMTItMTN6IiBmaWxsPSIjNDI4NUY0IiBvcGFjaXR5PSIuOTUiLz48dGV4dCB4PSI5MiIgeT0iMzUiIGZvbnQtZmFtaWx5PSJBcmlhbCwgSGVsdmV0aWNhLCBzYW5zLXNlcmlmIiBmb250LXNpemU9IjIzIiBmb250LXdlaWdodD0iNzAwIiBmaWxsPSIjNmU2ZTczIj5Lb3ZhbiBMYWJzPC90ZXh0Pjwvc3ZnPg=="
)


class OfferRenderError(ValueError):
    pass


def render_offer_html(
    *,
    template_snapshot: dict[str, Any],
    candidate: Candidate,
    job_posting: JobPosting,
    offer_letter: OfferLetter,
    organization: Organization,
    generated_date: datetime,
) -> str:
    first_name = (candidate.firstName or "").strip()
    last_name = (candidate.lastName or "").strip()
    if not first_name:
        raise OfferRenderError("Candidate first name is missing")
    if not last_name:
        raise OfferRenderError("Candidate last name is missing")

    template = _snapshot_object(template_snapshot, "template")
    sections = _snapshot_sections(template_snapshot)
    footer_html = str(template.get("footerHtml") or "")
    html_parts = []
    for section in sorted(sections, key=lambda item: int(item.get("order") or 0)):
        section_key = str(section.get("sectionKey") or "")
        section_html = str(section.get("html") or "")
        if section_key == "metadata":
            section_html = _normalize_metadata_header_html(section_html, template)
        html_parts.append(
            f'<section data-section="{escape(section_key)}">'
            f"{section_html}"
            "</section>"
        )
    if footer_html:
        footer_html = _repair_footer_asset_html(footer_html, template)
        html_parts.append(f"<footer>{footer_html}</footer>")
    body_html = "".join(html_parts)

    replacements = _build_replacements(
        candidate=candidate,
        job_posting=job_posting,
        offer_letter=offer_letter,
        generated_date=generated_date,
        html=body_html,
    )
    rendered_body = _replace_variables(body_html, replacements)
    _ensure_no_unknown_variables(rendered_body)

    return f"""<!DOCTYPE html>
<html>
<head>
  <meta charset="utf-8" />
  <style>
    @page {{ size: {A4_WIDTH_PX}px {A4_HEIGHT_PX}px; margin: 0; }}
    * {{ box-sizing: border-box; }}
    html {{
      width: {A4_WIDTH_PX}px;
      margin: 0;
      background: #ffffff;
      -webkit-print-color-adjust: exact;
      print-color-adjust: exact;
    }}
    body {{
      width: {A4_WIDTH_PX}px;
      margin: 0;
      color: #1d1d1f;
      font-family: Inter, -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
      font-size: {BODY_FONT_SIZE_PX}px;
      line-height: {BODY_LINE_HEIGHT_PX}px;
      background: #ffffff;
      -webkit-print-color-adjust: exact;
      print-color-adjust: exact;
    }}
    article {{
      position: relative;
      width: {A4_WIDTH_PX}px;
      min-height: {A4_HEIGHT_PX}px;
      padding: {PAGE_PADDING_TOP_PX}px {PAGE_PADDING_X_PX}px;
      background: #ffffff;
      page-break-after: always;
    }}
    article::before {{
      content: "";
      position: absolute;
      top: {CORNER_MARK_TOP_PX}px;
      left: {CORNER_MARK_LEFT_PX}px;
      width: {CORNER_MARK_WIDTH_PX}px;
      height: {CORNER_MARK_HEIGHT_PX}px;
      background: #00874a;
      z-index: 2;
    }}
    article::after {{
      content: "";
      position: absolute;
      right: {PAGE_PADDING_X_PX}px;
      bottom: 32px;
      width: 100px;
      height: 6px;
      background: linear-gradient(to right, #1d1d1f 0, #1d1d1f 64%, transparent 64%, transparent 68%, #00874a 68%, #00874a 100%);
      z-index: 2;
    }}
    article:last-child {{ page-break-after: auto; }}
    .offer-content {{ width: {CONTENT_WIDTH_PX}px; min-height: {CONTENT_HEIGHT_PX}px; }}
    .offer-page-corner-mark {{
      position: fixed;
      top: {CORNER_MARK_TOP_PX}px;
      left: {CORNER_MARK_LEFT_PX}px;
      width: {CORNER_MARK_WIDTH_PX}px;
      height: {CORNER_MARK_HEIGHT_PX}px;
      background: #00874a;
    }}
    .offer-page-bottom-mark {{
      position: fixed;
      right: {PAGE_PADDING_X_PX}px;
      bottom: 32px;
      width: 100px;
      height: 6px;
      background: linear-gradient(to right, #1d1d1f 0, #1d1d1f 64%, transparent 64%, transparent 68%, #00874a 68%, #00874a 100%);
    }}

    .offer-content section {{ margin: 0; }}
    .offer-content > * + * {{ margin-top: 16px; }}
    .offer-content h1, .offer-content h2 {{ font-size: 20px; line-height: 26px; margin: 0 0 14px; font-weight: 600; }}
    .offer-content h3 {{ font-size: 16px; line-height: 23px; margin: 14px 0 9px; font-weight: 600; }}
    .offer-content p {{ margin: 0 0 12px; }}
    .offer-content strong, .offer-content b {{ font-weight: 700; }}
    .offer-content em, .offer-content i {{ font-style: italic; }}
    .offer-content img {{ max-width: 100%; object-fit: contain; }}
    .offer-content a {{ color: #00874a; text-decoration: underline; }}
    .offer-content ul, .offer-content ol {{ margin: 0 0 12px; padding-left: 22px; }}
    .offer-content table {{ width: 100%; border-collapse: collapse; margin: 16px 0; page-break-inside: avoid; }}
    .offer-content th, .offer-content td {{ border: 1px solid #e5e5ea; padding: 8px; vertical-align: top; }}
    .offer-content th {{ background: #f2f2f7; font-weight: 600; }}
    .offer-content footer {{ margin-top: 22px; padding-top: 0; color: #1d1d1f; font-size: 13px; line-height: 19px; page-break-inside: avoid; break-inside: avoid; }}
    .offer-content footer p {{ margin: 0; line-height: 19px; }}
    .offer-letter-footer {{ display: grid; grid-template-columns: minmax(0, 1fr) auto; column-gap: 56px; align-items: end; page-break-inside: avoid; break-inside: avoid; }}
    .offer-letter-header {{ position: relative; min-height: 116px; margin-bottom: 24px; }}
    .offer-letter-header-brand {{ position: absolute; left: {HEADER_BRAND_LEFT_PX}px; top: {HEADER_BRAND_TOP_PX}px; }}
    .offer-letter-logo {{ max-width: 176px; max-height: 56px; object-fit: contain; }}
    .offer-letter-header-meta {{ position: absolute; right: 0; top: 48px; text-align: right; font-size: 13px; line-height: 24px; }}
    .offer-letter-header-meta p {{ margin: 0 0 4px; }}
    .offer-letter-header h1 {{ position: absolute; left: 0; right: 0; bottom: 0; margin: 0; text-align: center; white-space: nowrap; font-size: 16px; line-height: 22px; font-weight: 600; }}
    .offer-signature-slot {{ grid-column: 1; grid-row: 1; margin: 0 0 12px; page-break-inside: avoid; break-inside: avoid; }}
    .offer-signature-slot img {{ display: block; max-width: 128px; max-height: 80px; object-fit: contain; margin: 0; }}
    .offer-signature-slot p {{ margin: 0; line-height: 18px; }}
    .offer-signature-name {{ margin: 0; font-weight: 600; }}
    .offer-footer-address {{ grid-column: 1; grid-row: 2; margin: 0; }}
    .offer-footer-address p {{ margin: 0; line-height: 19px; }}
    .offer-content .offer-footer-website {{ grid-column: 2; grid-row: 2; align-self: center; float: none; margin: 0 0 4px; color: #1d1d1f; font-size: 16px; line-height: 20px; font-weight: 700; text-decoration: none; white-space: nowrap; }}
    [data-page-break="true"], .offer-page-break {{ break-after: page; page-break-after: always; height: 0; overflow: hidden; }}
  </style>
</head>
<body>
  <article>
    <main class="offer-content">{rendered_body}</main>
  </article>
</body>
</html>"""


def _normalize_metadata_header_html(html: str, template: dict[str, Any]) -> str:
    if "offer-letter-header" in html:
        return _repair_header_asset_html(html, template)

    paragraphs = [_html_text(match.group(1)) for match in PARAGRAPH_PATTERN.finditer(html)]
    headings = [_html_text(match.group(1)) for match in HEADING_PATTERN.finditer(html)]
    title = headings[0] if headings else "Offer Letter"
    generated_date = paragraphs[0] if len(paragraphs) >= 1 else ""
    location = paragraphs[1] if len(paragraphs) >= 2 else ""

    return (
        '<div class="offer-letter-header">'
        '<div class="offer-letter-header-brand">'
        f"{_logo_html(str(template.get('logoUrl') or ''))}"
        "</div>"
        '<div class="offer-letter-header-meta">'
        f"<p>{escape(generated_date)}</p>"
        f"<p>{escape(location)}</p>"
        "</div>"
        f"<h1>{escape(title)}</h1>"
        "</div>"
    )


def _repair_header_asset_html(html: str, template: dict[str, Any]) -> str:
    if "offer-letter-logo" in html:
        return html
    logo = _logo_html(str(template.get("logoUrl") or ""))
    return re.sub(
        r'(<div\b[^>]*class=["\'][^"\']*\boffer-letter-header-brand\b[^"\']*["\'][^>]*>)',
        rf"\1{logo}",
        html,
        count=1,
        flags=re.IGNORECASE,
    )


def _repair_footer_asset_html(html: str, template: dict[str, Any]) -> str:
    signature_url = str(template.get("signatureUrl") or "").strip()
    if not signature_url or "<img" in html.lower():
        return html
    image = f'<img src="{escape(signature_url, quote=True)}" alt="" />'
    repaired = re.sub(
        r'(<div\b[^>]*class=["\'][^"\']*\boffer-signature-slot\b[^"\']*["\'][^>]*>)',
        rf"\1{image}",
        html,
        count=1,
        flags=re.IGNORECASE,
    )
    if repaired != html:
        return repaired
    return f'<div class="offer-signature-slot">{image}</div>{html}'


def _logo_html(logo_url: str) -> str:
    src = logo_url.strip() or DEFAULT_KOVAN_LOGO_SVG
    return f'<img class="offer-letter-logo" src="{escape(src, quote=True)}" alt="Kovan Labs" />'


def _html_text(value: str) -> str:
    without_tags = TAG_PATTERN.sub(" ", value)
    return " ".join(unescape(without_tags).split())


async def render_offer_pdf(html: str) -> bytes:
    html = await _inline_offer_pdf_images(html)
    if sys.platform == "win32":
        return await asyncio.to_thread(_render_offer_pdf_in_windows_thread, html)
    return await _render_offer_pdf_with_playwright(html)


def render_offer_docx(html: str, *, title: str = "Offer Letter") -> bytes:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ModuleNotFoundError as exc:
        raise RuntimeError("python-docx is required to export offer letters as Word files") from exc

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_before = Pt(0)
    styles["Normal"].paragraph_format.space_after = Pt(0)
    styles["Normal"].paragraph_format.line_spacing = 1.0

    blocks = _extract_offer_text_blocks(html)
    if not blocks:
        document.add_paragraph(title)
    has_seen_footer_address = False
    for block in blocks:
        if block.kind == "heading":
            paragraph = document.add_heading(block.text, level=2)
        elif block.kind == "list":
            paragraph = document.add_paragraph(block.text, style="List Bullet")
        else:
            paragraph = document.add_paragraph(block.text)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        if "offer-footer-address" in block.classes and not has_seen_footer_address:
            paragraph.paragraph_format.space_before = Pt(8)
            has_seen_footer_address = True

    output = BytesIO()
    document.save(output)
    return output.getvalue()


@dataclass(frozen=True)
class _OfferTextBlock:
    kind: str
    text: str
    classes: tuple[str, ...] = ()


class _OfferTextExtractor(HTMLParser):
    block_tags = {"p", "div", "section", "footer", "tr"}
    heading_tags = {"h1", "h2", "h3", "h4", "h5", "h6"}
    list_tags = {"li"}
    skip_tags = {"script", "style"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[_OfferTextBlock] = []
        self._parts: list[str] = []
        self._kind_stack: list[str] = []
        self._tag_class_stack: list[tuple[str, tuple[str, ...]]] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self.skip_tags:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag != "br":
            classes = _html_classes(attrs)
            self._tag_class_stack.append((tag, classes))
        if tag in self.heading_tags:
            self._flush()
            self._kind_stack.append("heading")
        elif tag in self.list_tags:
            self._flush()
            self._kind_stack.append("list")
        elif tag in self.block_tags:
            self._flush()
            self._kind_stack.append("paragraph")
        elif tag == "br":
            self._parts.append("\n")
        elif tag in {"td", "th"}:
            self._parts.append("  ")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.skip_tags and self._skip_depth:
            self._skip_depth -= 1
            return
        if self._skip_depth:
            return
        if tag in self.heading_tags | self.list_tags | self.block_tags:
            self._flush()
            if self._kind_stack:
                self._kind_stack.pop()
        self._pop_tag_classes(tag)

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self._parts.append(data)

    def close(self) -> None:
        super().close()
        self._flush()

    def _flush(self) -> None:
        text = " ".join("".join(self._parts).replace("\xa0", " ").split())
        self._parts = []
        if text:
            kind = self._kind_stack[-1] if self._kind_stack else "paragraph"
            self.blocks.append(_OfferTextBlock(kind=kind, text=text, classes=self._active_classes()))

    def _active_classes(self) -> tuple[str, ...]:
        classes: list[str] = []
        for _tag, tag_classes in self._tag_class_stack:
            classes.extend(tag_classes)
        return tuple(classes)

    def _pop_tag_classes(self, tag: str) -> None:
        for index in range(len(self._tag_class_stack) - 1, -1, -1):
            if self._tag_class_stack[index][0] == tag:
                del self._tag_class_stack[index]
                return


def _html_classes(attrs: list[tuple[str, str | None]]) -> tuple[str, ...]:
    for name, value in attrs:
        if name == "class" and value:
            return tuple(value.split())
    return ()


def _extract_offer_text_blocks(html: str) -> list[_OfferTextBlock]:
    extractor = _OfferTextExtractor()
    extractor.feed(html)
    extractor.close()
    return extractor.blocks


async def _inline_offer_pdf_images(html: str) -> str:
    image_sources = {
        unescape(match.group(3)).strip()
        for match in IMG_SRC_PATTERN.finditer(html)
        if unescape(match.group(3)).strip().startswith(("http://", "https://"))
    }
    if not image_sources:
        return html

    replacements: dict[str, str] = {}
    async with httpx.AsyncClient(follow_redirects=True, timeout=15.0) as client:
        for source in image_sources:
            response = await client.get(source)
            response.raise_for_status()
            content = response.content
            if len(content) > MAX_INLINE_IMAGE_BYTES:
                raise RuntimeError(f"Offer PDF image is too large to embed: {source}")

            content_type = response.headers.get("content-type", "").split(";", 1)[0].strip()
            if not content_type.startswith("image/"):
                guessed_type = mimetypes.guess_type(urlparse(source).path)[0]
                content_type = guessed_type if guessed_type and guessed_type.startswith("image/") else ""
            if not content_type:
                raise RuntimeError(f"Offer PDF image has an unsupported content type: {source}")

            encoded = base64.b64encode(content).decode("ascii")
            replacements[source] = f"data:{content_type};base64,{encoded}"

    def replace_src(match: re.Match[str]) -> str:
        prefix, quote, raw_src, suffix_quote = match.group(1), match.group(2), match.group(3), match.group(4)
        source = unescape(raw_src).strip()
        replacement = replacements.get(source)
        if replacement is None:
            return match.group(0)
        return f"{prefix}{quote}{replacement}{suffix_quote}"

    return IMG_SRC_PATTERN.sub(replace_src, html)


async def _render_offer_pdf_with_playwright(html: str) -> bytes:
    try:
        from playwright.async_api import async_playwright
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Playwright is not installed. Install the server playwright package and Chromium browser."
        ) from exc

    playwright = await async_playwright().start()
    browser = None
    try:
        browser = await playwright.chromium.launch(headless=True)
        page = await browser.new_page(
            viewport={"width": A4_WIDTH_PX, "height": A4_HEIGHT_PX},
            device_scale_factor=1,
        )
        await page.emulate_media(media="screen")
        await page.set_content(html, wait_until="networkidle")
        await _wait_for_stable_offer_pdf_assets(page)
        return await page.pdf(
            width=f"{A4_WIDTH_PX}px",
            height=f"{A4_HEIGHT_PX}px",
            margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
            print_background=True,
            prefer_css_page_size=True,
            scale=1,
        )
    finally:
        if browser is not None:
            await browser.close()
        await playwright.stop()


async def _wait_for_stable_offer_pdf_assets(page: Any) -> None:
    failed_images = await page.evaluate(
        """async () => {
          if (document.fonts && document.fonts.ready) {
            await document.fonts.ready;
          }
          const failed = [];
          await Promise.all(
            Array.from(document.images).map((image) => {
              const settle = () => {
                if (!image.naturalWidth && !image.naturalHeight) {
                  failed.push(image.currentSrc || image.src);
                }
              };
              if (image.complete) {
                settle();
                return Promise.resolve();
              }
              if (image.decode) return image.decode().catch(() => undefined).then(settle);
              return new Promise((resolve) => {
                image.addEventListener('load', () => { settle(); resolve(); }, { once: true });
                image.addEventListener('error', () => { settle(); resolve(); }, { once: true });
              });
            })
          );
          return failed;
        }"""
    )
    if failed_images:
        failed = ", ".join(str(source) for source in failed_images[:3])
        raise RuntimeError(f"Offer PDF image(s) failed to load: {failed}")


def _render_offer_pdf_in_windows_thread(html: str) -> bytes:
    if not hasattr(asyncio, "ProactorEventLoop"):
        return asyncio.run(_render_offer_pdf_with_playwright(html))

    loop = asyncio.ProactorEventLoop()
    try:
        asyncio.set_event_loop(loop)
        return loop.run_until_complete(_render_offer_pdf_with_playwright(html))
    finally:
        loop.run_until_complete(loop.shutdown_asyncgens())
        asyncio.set_event_loop(None)
        loop.close()


def _snapshot_object(snapshot: dict[str, Any], key: str) -> dict[str, Any]:
    value = snapshot.get(key)
    return value if isinstance(value, dict) else {}


def _snapshot_sections(snapshot: dict[str, Any]) -> list[dict[str, Any]]:
    sections = snapshot.get("sections")
    return [item for item in sections if isinstance(item, dict)] if isinstance(sections, list) else []


def _build_replacements(
    *,
    candidate: Candidate,
    job_posting: JobPosting,
    offer_letter: OfferLetter,
    generated_date: datetime,
    html: str,
) -> dict[str, str]:
    tokens = {match.group(1) for match in VARIABLE_PATTERN.finditer(html)}
    unknown_tokens = sorted(tokens - ALLOWED_VARIABLE_TOKENS)
    if unknown_tokens:
        raise OfferRenderError(f"Unknown offer variable token(s): {', '.join(unknown_tokens)}")

    requisition = job_posting.requisition
    if tokens & COMPENSATION_VARIABLE_TOKENS and (
        requisition is None
        or requisition.salaryMin is None
        or requisition.salaryMax is None
        or not (requisition.currency or "").strip()
    ):
        raise OfferRenderError("Job salary data is missing")

    salary_min = requisition.salaryMin if requisition is not None else None
    salary_max = requisition.salaryMax if requisition is not None else offer_letter.salary
    currency = (requisition.currency if requisition is not None else offer_letter.currency) or offer_letter.currency

    return {
        "candidate.firstName": escape((candidate.firstName or "").strip()),
        "candidate.lastName": escape((candidate.lastName or "").strip()),
        "offer.generatedDate": escape(generated_date.strftime("%d %b %Y")),
        "job.salaryMin": escape(_format_amount(salary_min)),
        "job.salaryMax": escape(_format_amount(salary_max)),
        "job.currency": escape((currency or "").strip()),
    }


def _replace_variables(value: str, replacements: dict[str, str]) -> str:
    return VARIABLE_PATTERN.sub(lambda match: replacements.get(match.group(1), match.group(0)), value)


def _ensure_no_unknown_variables(value: str) -> None:
    remaining = sorted({match.group(1) for match in VARIABLE_PATTERN.finditer(value)})
    if remaining:
        raise OfferRenderError(f"Unknown offer variable token(s): {', '.join(remaining)}")


def _format_amount(value: float | int | None) -> str:
    if value is None:
        return ""
    if float(value).is_integer():
        return f"{int(value):,}"
    return f"{float(value):,.2f}"


def _image(class_name: str, url: str) -> str:
    if not url:
        return ""
    return f'<img class="{class_name}" src="{escape(url, quote=True)}" alt="" />'


def _signature(signature_url: str, signatory_name: str, signatory_title: str) -> str:
    if not signature_url and not signatory_name and not signatory_title:
        return ""
    image = _image("", signature_url)
    name = f'<p class="signature-name">{escape(signatory_name)}</p>' if signatory_name else ""
    title = f'<p class="signature-title">{escape(signatory_title)}</p>' if signatory_title else ""
    return f'<div class="signature">{image}{name}{title}</div>'
