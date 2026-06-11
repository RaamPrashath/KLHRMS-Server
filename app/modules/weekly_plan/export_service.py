"""Plan report export service — XLSX, PDF, CSV."""

from __future__ import annotations

import csv
import html
import io
import re
from collections import defaultdict
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.modules.weekly_plan.schema import PlanExportRequest, PlanExportRow
from app.modules.weekly_plan.locations import PLAN_LOCATION_MAP

_HEADER_FILL = "1F4E78"
_TITLE_FILL = "D9EAF7"
_SUBTOTAL_FILL = "D9EAF7"
_THIN_BORDER = Side(style="thin", color="D9E2EC")
_GRID_BORDER = Border(left=_THIN_BORDER, right=_THIN_BORDER, top=_THIN_BORDER, bottom=_THIN_BORDER)
_REPORT_HEADERS = ["Date", "Day", "Location", "Project"]


def _employee_name(employee) -> str:
    return employee.name.strip() or employee.email or employee.id


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", " ", name).strip() or "Employee"
    base = cleaned[:31]
    candidate = base
    counter = 2
    while candidate in used:
        suffix = f" {counter}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate)
    return candidate


def _date_label(value: date) -> str:
    return value.isoformat()


def _day_label(value: date) -> str:
    return value.strftime("%A")


def _location_label(row: PlanExportRow) -> str:
    return PLAN_LOCATION_MAP.get(row.work_location, {}).get("label", row.work_location)


def _project_label(row: PlanExportRow) -> str:
    return row.project or ""


def _rows_by_employee(rows: list[PlanExportRow]) -> dict[str, list[PlanExportRow]]:
    grouped: dict[str, list[PlanExportRow]] = defaultdict(list)
    for row in rows:
        grouped[row.user_id].append(row)
    for employee_rows in grouped.values():
        employee_rows.sort(key=lambda item: item.date)
    return grouped


def _employees_for_export(payload: PlanExportRequest) -> list:
    if payload.employees:
        return sorted(payload.employees, key=lambda item: _employee_name(item).lower())
    seen: dict[str, PlanExportRequest] = {}
    for row in payload.rows:
        key = row.user_id
        if key not in seen:
            seen[key] = row
    return sorted((e for e in payload.employees if e.id in {r.user_id for r in payload.rows}),
                  key=lambda item: _employee_name(item).lower()) if payload.employees else []


def _apply_report_sheet_layout(ws, last_row: int) -> None:
    widths = [14, 14, 22, 64]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for row in ws.iter_rows(min_row=1, max_row=last_row, max_col=len(_REPORT_HEADERS)):
        for cell in row:
            cell.border = _GRID_BORDER


def generate_plan_xlsx(payload: PlanExportRequest) -> bytes:
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    grouped = _rows_by_employee(payload.rows)
    used_sheet_names: set[str] = set()
    employees = sorted(payload.employees, key=lambda item: _employee_name(item).lower()) if payload.employees else []

    if not employees:
        employees = [type("_Employee", (), {"id": "empty", "name": "No selected employees", "email": None})()]

    for employee in employees:
        employee_rows = grouped.get(employee.id, [])
        ws = wb.create_sheet(_safe_sheet_name(_employee_name(employee), used_sheet_names))

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_REPORT_HEADERS))
        title_cell = ws.cell(row=1, column=1, value=f"Name: {_employee_name(employee)}")
        title_cell.font = Font(bold=True, size=12, color="1D1D1F")
        title_cell.fill = PatternFill("solid", fgColor=_TITLE_FILL)
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 24

        for col, header in enumerate(_REPORT_HEADERS, start=1):
            cell = ws.cell(row=2, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 24

        for row_index, row in enumerate(employee_rows, start=3):
            values = [
                _date_label(row.date),
                _day_label(row.date),
                _location_label(row),
                _project_label(row),
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_index, column=col, value=value)
                cell.alignment = Alignment(
                    horizontal="center" if col <= 2 else "left",
                    vertical="top",
                    wrap_text=True,
                )
            ws.row_dimensions[row_index].height = 22

        _apply_report_sheet_layout(ws, len(employee_rows) + 2)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _paragraph(value: str, style: ParagraphStyle) -> Paragraph:
    escaped = html.escape(value or "").replace("\n", "<br/>")
    return Paragraph(escaped or "&nbsp;", style)


def generate_plan_pdf(payload: PlanExportRequest) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("ReportCell", parent=styles["Normal"], fontSize=8, leading=10)
    title_style = ParagraphStyle("ReportTitle", parent=styles["Heading2"], fontSize=12, leading=14)

    grouped = _rows_by_employee(payload.rows)
    story = []
    employees = sorted(payload.employees, key=lambda item: _employee_name(item).lower()) if payload.employees else []

    for employee_index, employee in enumerate(employees):
        if employee_index:
            story.append(PageBreak())
        story.append(Paragraph(f"Name: {_employee_name(employee)}", title_style))
        if payload.periodLabel:
            story.append(Paragraph(html.escape(payload.periodLabel), normal))
        story.append(Spacer(1, 4 * mm))

        table_data: list[list[object]] = [_REPORT_HEADERS]
        for row in grouped.get(employee.id, []):
            table_data.append([
                _date_label(row.date),
                _day_label(row.date),
                _location_label(row),
                _project_label(row),
            ])

        table = Table(
            table_data,
            colWidths=[24 * mm, 24 * mm, 38 * mm, 108 * mm],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_HEADER_FILL}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E2EC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ])
        )
        story.append(table)

    if not story:
        story.append(Paragraph("No selected employees", title_style))
    doc.build(story)
    return buffer.getvalue()


def generate_plan_csv(payload: PlanExportRequest) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Employee", "Date", "Day", "Location", "Project"])

    employees = sorted(payload.employees, key=lambda item: _employee_name(item).lower()) if payload.employees else []
    grouped = _rows_by_employee(payload.rows)

    if not employees:
        employees = [type("_Employee", (), {"id": "empty", "name": "No selected employees", "email": None})()]

    for employee in employees:
        employee_rows = grouped.get(employee.id, [])
        for row in employee_rows:
            writer.writerow([
                _employee_name(employee),
                _date_label(row.date),
                _day_label(row.date),
                _location_label(row),
                _project_label(row),
            ])

    return output.getvalue().encode("utf-8-sig")


def generate_plan_export(payload: PlanExportRequest) -> bytes:
    if payload.format == "xlsx":
        return generate_plan_xlsx(payload)
    if payload.format == "csv":
        return generate_plan_csv(payload)
    return generate_plan_pdf(payload)
