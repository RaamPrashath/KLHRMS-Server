from __future__ import annotations

import csv
import hashlib
import io
import secrets
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import Literal

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import Select, extract, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import joinedload, selectinload

from app.models.asset import Asset
from app.models.asset_assignment import AssetAssignment
from app.models.asset_category_definition import AssetCategoryDefinition
from app.models.asset_category_field_definition import AssetCategoryFieldDefinition
from app.models.asset_custom_field_value import AssetCustomFieldValue
from app.models.asset_id_definition import AssetIdDefinition
from app.models.asset_maintenance_log import AssetMaintenanceLog
from app.models.asset_notification import AssetNotification
from app.models.asset_purchase_requisition import AssetPurchaseRequisition
from app.models.asset_unit import AssetUnit
from app.models.base import generate_uuid
from app.models.member import Member
from app.models.organization import Organization
from app.models.role import Role
from app.models.user import User
from app.modules.notifications.service import NotificationCreateInput, create_notification_batch
from app.shared.notifications.email import (
    send_asset_warranty_expiry_alert,
    send_asset_replacement_notification,
    send_temp_replacement_reminder,
)
from app.modules.assets.schema import (
    ASSET_CONDITIONS,
    ASSET_STATUS_ALIASES,
    ASSET_STATUSES,
    CRITICALITY_TIERS,
    MAINTENANCE_STATUSES,
    MAINTENANCE_TYPES,
    REPORT_TYPES,
    SWAP_MODES,
    TICKET_MODES,
    AssetCategoryCreate,
    AssetCategoryResponse,
    AssetCategoryUpdate,
    AssetBrandModelAnalyticsResponse,
    AssetBrandModelAnalyticsRow,
    AssetDashboardResponse,
    AssetDetailResponse,
    EmployeeAssetViewItem,
    EmployeeAssetViewResponse,
    AssetFilters,
    AssetIdCreate,
    AssetIdResponse,
    AssetIdUpdate,
    AssetIssueRequest,
    AssetIssueResponse,
    AssetListResponse,
    AssetLookupOption,
    AssetMaintenanceCreateRequest,
    AssetMaintenanceSummary,
    AssetMaintenanceUpdateRequest,
    AssetMetaResponse,
    AssetOsDistributionResponse,
    AssetOsDistributionRow,
    AssetRevokeSwapRequest,
    AssetProvideRecordSummary,
    AssetReportRequest,
    AssetReturnRequest,
    AssetSwapExecutionResponse,
    AssetSwapPreviewResponse,
    AssetStatusCount,
    AssetSummary,
    AssetUnitResponse,
    AssetUnitSummary,
    AssetUpsertRequest,
    AvailableAssetGroupResponse,
    BulkAssetCreateRequest,
    CategoryFieldDefinitionCreate,
    CategoryFieldDefinitionResponse,
    CategoryFieldDefinitionUpdate,
    CustomFieldValueResponse,
    HelpdeskTicketCreateRequest,
    MaintenanceTicketResponse,
    MonthlyTrend,
    MyTicketResponse,
    RecentActivityItem,
    ReturnedAssetSummary,
    SwapAvailabilityOption,
    TicketAlertItem,
    WarrantyExpirationFeedItem,
    WarrantyExpirationFeedResponse,
    ReplacementRecord,
    ReplacementProvideRequest,
    ReplacementRaiseAppraisalRequest,
    SetReturnDateRequest,
    MemberTicketSummary,
)
from app.shared.deps.organization_member import MemberContext


def _to_title(value: str) -> str:
    return value.lower().replace("_", " ").title()


def _normalize_asset_status(value: str | None) -> str | None:
    if value is None:
        return None
    normalized = value.strip().upper().replace("-", "_").replace(" ", "_")
    return ASSET_STATUS_ALIASES.get(normalized, normalized)


def _is_assigned_status(value: str | None) -> bool:
    return _normalize_asset_status(value) == "ASSIGNED"


def _is_maintenance_status(value: str | None) -> bool:
    return _normalize_asset_status(value) == "IN_MAINTENANCE"


def _normalized_token(value: str | None) -> str:
    return " ".join((value or "").strip().split())


def _normalize_key(value: str | None) -> str:
    return _normalized_token(value).lower().replace("-", " ")


def _is_laptop_asset(asset: Asset) -> bool:
    category_key = _normalize_key(asset.category)
    name_key = _normalize_key(asset.name)
    model_key = _normalize_key(asset.model)
    keywords = ("laptop", "notebook", "macbook", "thinkpad", "elitebook", "probook")
    return "laptop" in category_key or any(
        keyword in name_key or keyword in model_key for keyword in keywords
    )


def _is_laptop_category_name(value: str | None) -> bool:
    normalized = _normalize_key(value)
    return normalized in {"laptop", "laptops"}


async def _get_laptop_category(
    db: AsyncSession, organization_id: str
) -> AssetCategoryDefinition | None:
    result = await db.execute(
        select(AssetCategoryDefinition)
        .where(
            AssetCategoryDefinition.organizationId == organization_id,
            AssetCategoryDefinition.isActive.is_(True),
        )
        .options(joinedload(AssetCategoryDefinition.fields))
        .order_by(AssetCategoryDefinition.name.asc())
    )
    categories = result.unique().scalars().all()
    return next((category for category in categories if _is_laptop_category_name(category.name)), None)


async def _ensure_laptop_os_field(
    db: AsyncSession, organization_id: str
) -> AssetCategoryFieldDefinition | None:
    laptop_category = await _get_laptop_category(db, organization_id)
    if laptop_category is None:
        return None

    existing_field = next(
        (
            field
            for field in (laptop_category.fields or [])
            if _normalize_key(field.fieldName) == "os"
        ),
        None,
    )
    if existing_field is not None:
        if existing_field.fieldType != "SELECT":
            existing_field.fieldType = "SELECT"
        merged_options = ["Windows", "macOS", "Linux", "ChromeOS", "Ubuntu", "Other"]
        existing_values = ((existing_field.fieldOptions or {}).get("options") or [])[:]
        for option in merged_options:
            if option not in existing_values:
                existing_values.append(option)
        existing_field.fieldOptions = {"options": existing_values}
        await db.commit()
        return existing_field

    display_order = max((field.displayOrder for field in (laptop_category.fields or [])), default=-1) + 1
    field = AssetCategoryFieldDefinition(
        categoryId=laptop_category.id,
        fieldName="OS",
        fieldType="SELECT",
        fieldOptions={"options": ["Windows", "macOS", "Linux", "ChromeOS", "Ubuntu", "Other"]},
        isRequired=False,
        displayOrder=display_order,
    )
    db.add(field)
    await db.commit()
    return field


def _resolve_asset_brand(asset: Asset) -> str:
    if asset.brand and asset.brand.strip():
        return asset.brand.strip()

    brand_field_names = {"brand", "manufacturer", "make"}
    for field_value in asset.customFieldValues or []:
        field_name = _normalize_key(
            field_value.fieldDefinition.fieldName if field_value.fieldDefinition else None
        )
        if field_name in brand_field_names and _normalized_token(field_value.value):
            return _to_title(_normalized_token(field_value.value))

    name_value = _normalized_token(asset.name)
    generic_name_tokens = {"laptop", "notebook", "computer", "device", "asset", "temporary"}
    if name_value and name_value.split()[0].lower() not in generic_name_tokens:
        return name_value.split()[0].title()

    model_value = _normalized_token(asset.model)
    if model_value:
        return model_value.split()[0].title()

    return "Unknown"


def _resolve_asset_model(asset: Asset, brand: str) -> str:
    raw_model = _normalized_token(asset.model) or _normalized_token(asset.name) or "Unknown Model"
    brand_key = _normalize_key(brand)
    model_key = _normalize_key(raw_model)
    if brand_key and model_key.startswith(brand_key):
        trimmed = raw_model[len(brand) :].strip(" -_/")
        if trimmed:
            return trimmed
    return raw_model


def _is_temporary_laptop(asset: Asset, brand: str, model: str) -> bool:
    combined = " ".join(
        filter(
            None,
            [
                _normalize_key(asset.name),
                _normalize_key(asset.assetCode),
                _normalize_key(asset.model),
                _normalize_key(brand),
                _normalize_key(model),
            ],
        )
    )
    markers = ("temporary", "temp", "loaner", "spare")
    return any(marker in combined for marker in markers)


def _normalize_os_name(value: str | None) -> str:
    normalized = _normalize_key(value)
    aliases = {
        "windows": "Windows",
        "windows 10": "Windows",
        "windows 11": "Windows",
        "macos": "macOS",
        "mac os": "macOS",
        "os x": "macOS",
        "linux": "Linux",
        "ubuntu": "Ubuntu",
        "chromeos": "ChromeOS",
        "chrome os": "ChromeOS",
    }
    if not normalized:
        return "Unknown"
    return aliases.get(normalized, _to_title(normalized))


def _resolved_downtime_hours(log: AssetMaintenanceLog) -> int | None:
    if log.estimatedDowntimeHours is not None:
        return log.estimatedDowntimeHours
    if log.expectedCompletionDate and log.serviceDate:
        return max((log.expectedCompletionDate - log.serviceDate).days, 0) * 24
    return None


def _criticality_rank(tier: str | None) -> int:
    normalized = (tier or "STANDARD").strip().upper()
    if normalized == "MISSION_CRITICAL":
        return 3
    if normalized == "BUSINESS_CRITICAL":
        return 2
    return 1


def _requires_replacement_validation(log: AssetMaintenanceLog, asset: Asset | None) -> bool:
    downtime_hours = _resolved_downtime_hours(log)
    normalized_condition = (log.conditionBeforeMaintenance or asset.condition if asset else None) or ""
    return bool(
        downtime_hours is not None
        and downtime_hours > 24
        and (
            _criticality_rank(log.operationalCriticalityTier) >= 2
            or normalized_condition in {"DAMAGED", "NEEDS_REPAIR"}
        )
    )


def _should_manage_assets(role: Role | None) -> bool:
    if role is None:
        return False

    permissions = role.permissions or {}
    assets_permissions = permissions.get("assets") or {}
    if assets_permissions.get("edit") == "organization":
        return True
    return "admin" in (role.name or "").strip().lower()


async def _notify_general_helpdesk_admins(
    db: AsyncSession,
    ctx: MemberContext,
    ticket: AssetMaintenanceLog,
) -> None:
    if ticket.ticketMode != "GENERAL_HELP_REQUEST":
        return

    admins = await _members_with_asset_admin_scope(db, ctx.organization.id)
    title = ticket.subject or _to_title(ticket.maintenanceType)
    await create_notification_batch(
        db,
        [
            NotificationCreateInput(
                organization_id=ctx.organization.id,
                member_id=member.id,
                type="HELPDESK_REQUEST_CREATED",
                category="helpdesk",
                title="New helpdesk request needs attention",
                message=f"{title} was submitted and is now waiting in the shared queue.",
                action_url=f"/{ctx.organization.slug}/maintenance",
                entity_type="HELPDESK_TICKET",
                entity_id=ticket.id,
                metadata={"ticketId": ticket.ticketId, "status": ticket.status},
            )
            for member in admins
            if member.id != ctx.member.id
        ],
    )


async def _notify_general_helpdesk_requester_if_completed(
    db: AsyncSession,
    ctx: MemberContext,
    ticket: AssetMaintenanceLog,
    previous_status: str | None,
) -> None:
    if ticket.ticketMode != "GENERAL_HELP_REQUEST":
        return
    if previous_status == "COMPLETED" or ticket.status != "COMPLETED":
        return
    if not ticket.loggedByMemberId:
        return

    await create_notification_batch(
        db,
        [
            NotificationCreateInput(
                organization_id=ctx.organization.id,
                member_id=ticket.loggedByMemberId,
                type="HELPDESK_REQUEST_COMPLETED",
                category="helpdesk",
                title="Your helpdesk request was resolved",
                message=f"{ticket.subject or _to_title(ticket.maintenanceType)} was marked as completed.",
                action_url=f"/{ctx.organization.slug}/helpdesk",
                entity_type="HELPDESK_TICKET",
                entity_id=ticket.id,
                metadata={"ticketId": ticket.ticketId, "status": ticket.status},
            )
        ],
    )


async def _ensure_tracking_unit(db: AsyncSession, asset: Asset) -> AssetUnit:
    if asset.units:
        return asset.units[0]

    tracking_unit = AssetUnit(
        assetId=asset.id,
        serialNumber=asset.serialNumber,
        status=asset.status,
        currentHolderMemberId=None,
        condition=asset.condition,
        warrantyExpiryDate=asset.warrantyExpiryDate,
        reminderCompleted=False,
    )
    db.add(tracking_unit)
    await db.flush()
    asset.units.append(tracking_unit)
    return tracking_unit


async def _active_assignment_map_for_assets(
    db: AsyncSession, asset_ids: list[str]
) -> dict[str, AssetAssignment]:
    if not asset_ids:
        return {}

    result = await db.execute(
        select(AssetAssignment)
        .where(
            AssetAssignment.assetId.in_(asset_ids),
            AssetAssignment.returnDate.is_(None),
        )
        .options(
            joinedload(AssetAssignment.member).joinedload(Member.user),
            joinedload(AssetAssignment.providedByMember).joinedload(Member.user),
            joinedload(AssetAssignment.receivedByMember).joinedload(Member.user),
        )
    )
    assignments = result.unique().scalars().all()
    return {assignment.assetId: assignment for assignment in assignments}


async def _members_with_asset_admin_scope(
    db: AsyncSession, organization_id: str
) -> list[Member]:
    result = await db.execute(
        select(Member)
        .where(Member.organizationId == organization_id)
        .options(
            joinedload(Member.user),
            joinedload(Member.role),
        )
    )
    members = result.unique().scalars().all()
    return [member for member in members if member.user and _should_manage_assets(member.role)]


async def _list_available_laptop_units(
    db: AsyncSession, organization_id: str, exclude_unit_id: str | None = None
) -> list[tuple[Asset, AssetUnit]]:
    result = await db.execute(
        select(AssetUnit)
        .join(Asset, Asset.id == AssetUnit.assetId)
        .where(
            Asset.organizationId == organization_id,
            Asset.deletedAt.is_(None),
            AssetUnit.status.in_(("AVAILABLE",)),
        )
        .options(
            joinedload(AssetUnit.asset).joinedload(Asset.customFieldValues).joinedload(
                AssetCustomFieldValue.fieldDefinition
            ),
            joinedload(AssetUnit.asset).joinedload(Asset.units),
        )
    )
    rows = result.unique().scalars().all()
    available: list[tuple[Asset, AssetUnit]] = []
    for unit in rows:
        asset = unit.asset
        if asset is None or not _is_laptop_asset(asset):
            continue
        if exclude_unit_id and unit.id == exclude_unit_id:
            continue
        available.append((asset, unit))
    return available


async def _build_swap_preview(
    db: AsyncSession, organization_id: str, log: AssetMaintenanceLog, asset: Asset
) -> AssetSwapPreviewResponse:
    active_assignment = next((record for record in asset.provisions if record.returnDate is None), None)
    assigned_member_name = None
    if active_assignment and active_assignment.member and active_assignment.member.user:
        assigned_member_name = active_assignment.member.user.name or active_assignment.member.user.email

    source_unit = None
    if log.assetUnitId:
        source_unit = next((unit for unit in asset.units if unit.id == log.assetUnitId), None)

    available_laptops = await _list_available_laptop_units(db, organization_id, exclude_unit_id=log.assetUnitId)

    asset_brand = _resolve_asset_brand(asset)
    asset_model = _resolve_asset_model(asset, asset_brand)
    exact_match_units: list[tuple[Asset, AssetUnit]] = []
    temporary_units: list[tuple[Asset, AssetUnit]] = []
    exact_match_unit_ids: set[str] = set()
    temporary_unit_ids: set[str] = set()
    for candidate_asset, candidate_unit in available_laptops:
        candidate_brand = _resolve_asset_brand(candidate_asset)
        candidate_model = _resolve_asset_model(candidate_asset, candidate_brand)
        if candidate_model == asset_model:
            exact_match_units.append((candidate_asset, candidate_unit))
            exact_match_unit_ids.add(candidate_unit.id)
        if _is_temporary_laptop(candidate_asset, candidate_brand, candidate_model):
            temporary_units.append((candidate_asset, candidate_unit))
            temporary_unit_ids.add(candidate_unit.id)

    any_available_units: list[tuple[Asset, AssetUnit]] = []
    for candidate_asset, candidate_unit in available_laptops:
        if candidate_unit.id not in exact_match_unit_ids and candidate_unit.id not in temporary_unit_ids:
            any_available_units.append((candidate_asset, candidate_unit))

    requires_validation = _requires_replacement_validation(log, asset)
    recommended_mode: str | None = None
    if requires_validation:
        if exact_match_units:
            recommended_mode = "PERMANENT_REPLACEMENT"
        elif temporary_units:
            recommended_mode = "TEMPORARY_BACKUP"
        elif any_available_units:
            recommended_mode = "ANY_AVAILABLE"

    downtime_hours = _resolved_downtime_hours(log)
    if not requires_validation:
        reason = "Repair window is within 24 hours, so a swap is optional."
    elif exact_match_units:
        reason = "Repair exceeds 24 hours and matching model stock is available for a permanent swap."
    elif temporary_units:
        reason = "Repair exceeds 24 hours and exact model stock is unavailable, so a temporary backup is recommended."
    elif any_available_units:
        reason = "Repair exceeds 24 hours but no matching model or temporary stock is available. A different model replacement is offered as a fallback."
    else:
        reason = "Repair exceeds 24 hours but no replacement inventory is currently available."

    return AssetSwapPreviewResponse(
        maintenanceId=log.id,
        assetId=asset.id,
        assetUnitId=log.assetUnitId,
        currentAssetStatus=_normalize_asset_status(source_unit.status if source_unit else asset.status)
        or asset.status,
        currentCondition=log.conditionBeforeMaintenance or asset.condition,
        assignedMemberId=active_assignment.memberId if active_assignment else None,
        assignedMemberName=assigned_member_name,
        brand=asset.brand,
        model=asset.model,
        operationalCriticalityTier=log.operationalCriticalityTier,
        estimatedDowntimeHours=downtime_hours,
        requiresReplacementValidation=requires_validation,
        recommendedMode=recommended_mode,
        reason=reason,
        options=[
            SwapAvailabilityOption(
                mode="PERMANENT_REPLACEMENT",
                label="Exact Model Replacement",
                available=bool(exact_match_units),
                availableCount=len(exact_match_units),
                assetUnitIds=[unit.id for _, unit in exact_match_units[:8]],
                serialNumbers=[unit.serialNumber or "Unit" for _, unit in exact_match_units[:8]],
                recommended=recommended_mode == "PERMANENT_REPLACEMENT",
            ),
            SwapAvailabilityOption(
                mode="TEMPORARY_BACKUP",
                label="Temporary Backup",
                available=bool(temporary_units),
                availableCount=len(temporary_units),
                assetUnitIds=[unit.id for _, unit in temporary_units[:8]],
                serialNumbers=[unit.serialNumber or "Unit" for _, unit in temporary_units[:8]],
                recommended=recommended_mode == "TEMPORARY_BACKUP",
            ),
            SwapAvailabilityOption(
                mode="ANY_AVAILABLE",
                label="Any Available Laptop",
                available=bool(any_available_units),
                availableCount=len(any_available_units),
                assetUnitIds=[unit.id for _, unit in any_available_units[:8]],
                serialNumbers=[unit.serialNumber or "Unit" for _, unit in any_available_units[:8]],
                recommended=recommended_mode == "ANY_AVAILABLE",
            ),
        ],
    )


async def _generate_ticket_id(db: AsyncSession) -> str:
    for _ in range(10):
        ticket_id = f"#{secrets.randbelow(1_000_000):06d}"
        result = await db.execute(
            select(AssetMaintenanceLog.id).where(AssetMaintenanceLog.ticketId == ticket_id)
        )
        if result.scalar_one_or_none() is None:
            return ticket_id

    raise HTTPException(status_code=500, detail="Unable to generate a unique ticket ID")


def _csv_to_pdf_bytes(report_type: str, csv_text: str) -> bytes:
    rows = list(csv.reader(io.StringIO(csv_text)))
    if not rows:
        rows = [["(no data)"]]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"{_to_title(report_type)} Report",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Title",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        spaceAfter=4,
        textColor=colors.HexColor("#111827"),
    )
    meta_style = ParagraphStyle(
        "Meta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#6b7280"),
        spaceAfter=10,
    )
    cell_style = ParagraphStyle(
        "Cell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#111827"),
    )
    header_style = ParagraphStyle(
        "Header",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=10,
        textColor=colors.whitesmoke,
    )

    story: list = [
        Paragraph(f"{_to_title(report_type)} Report", title_style),
        Paragraph(
            f"Generated at {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}",
            meta_style,
        ),
        Spacer(1, 4),
    ]

    header = [Paragraph(str(col), header_style) for col in rows[0]]
    body = [
        [Paragraph((str(col) if col is not None else "").replace("\n", " "), cell_style) for col in row]
        for row in rows[1:]
    ]
    table_data = [header, *body]

    page_width, _page_height = A4
    available_width = page_width - 24 * mm
    col_count = max(len(rows[0]), 1)
    equal_col = available_width / col_count
    col_widths = [equal_col] * col_count

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8.5),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 7),
                ("TOPPADDING", (0, 0), (-1, 0), 7),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 1), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, 0), "LEFT"),
                ("ALIGN", (0, 1), (-1, -1), "LEFT"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f8fa")]),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e5e7eb")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 1), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
            ]
        )
    )
    story.append(table)

    doc.build(story)
    buffer.seek(0)
    return buffer.read()


def _csv_to_xlsx_bytes(report_type: str, csv_text: str) -> bytes:
    rows = list(csv.reader(io.StringIO(csv_text)))
    if not rows:
        rows = [["(no data)"]]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = _to_title(report_type)[:31] or "Report"

    header_fill = PatternFill("solid", fgColor="FF1F2937")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="FF111827")
    meta_font = Font(name="Calibri", size=10, italic=True, color="FF6B7280")
    cell_font = Font(name="Calibri", size=10, color="FF111827")
    alt_fill = PatternFill("solid", fgColor="FFF7F8FA")
    thin = Side(style="thin", color="FFE5E7EB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    worksheet.cell(row=1, column=1, value=f"{_to_title(report_type)} Report").font = title_font
    worksheet.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=max(len(rows[0]), 1)
    )
    worksheet.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center")

    worksheet.cell(
        row=2,
        column=1,
        value=f"Generated at {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}",
    ).font = meta_font
    worksheet.merge_cells(
        start_row=2, start_column=1, end_row=2, end_column=max(len(rows[0]), 1)
    )
    worksheet.cell(row=2, column=1).alignment = Alignment(horizontal="left", vertical="center")

    header_row_index = 4
    for col_index, header_value in enumerate(rows[0], start=1):
        cell = worksheet.cell(row=header_row_index, column=col_index, value=header_value)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = left_align
        cell.border = border

    for row_offset, row in enumerate(rows[1:], start=1):
        zebra = alt_fill if row_offset % 2 == 0 else None
        for col_index, value in enumerate(row, start=1):
            text = "" if value is None else str(value).replace("\n", " ")
            cell = worksheet.cell(row=header_row_index + row_offset, column=col_index, value=text)
            cell.font = cell_font
            cell.alignment = left_align
            cell.border = border
            if zebra is not None:
                cell.fill = zebra

    col_count = max(len(rows[0]), 1)
    for col_index in range(1, col_count + 1):
        letter = get_column_letter(col_index)
        max_length = 0
        for row in rows:
            if col_index - 1 < len(row):
                cell_value = "" if row[col_index - 1] is None else str(row[col_index - 1])
                max_length = max(max_length, len(cell_value))
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 40)

    worksheet.row_dimensions[1].height = 22
    worksheet.row_dimensions[2].height = 18
    worksheet.row_dimensions[header_row_index].height = 24

    worksheet.freeze_panes = worksheet.cell(row=header_row_index + 1, column=1)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _to_float(value: Decimal | float | int | None) -> float | None:
    if value is None:
        return None
    return float(value)


async def _get_asset_or_404(db: AsyncSession, organization_id: str, asset_id: str) -> Asset:
    result = await db.execute(
        select(Asset)
        .where(
            Asset.id == asset_id,
            Asset.organizationId == organization_id,
            Asset.deletedAt.is_(None),
        )
        .options(
            joinedload(Asset.provisions),
            joinedload(Asset.maintenanceLogs),
            joinedload(Asset.units),
            joinedload(Asset.customFieldValues).joinedload(AssetCustomFieldValue.fieldDefinition),
        )
    )
    asset = result.unique().scalar_one_or_none()
    if asset is None:
        raise HTTPException(status_code=404, detail="Asset not found")
    return asset


async def _get_member_or_404(db: AsyncSession, organization_id: str, member_id: str) -> Member:
    result = await db.execute(
        select(Member)
        .where(Member.id == member_id, Member.organizationId == organization_id)
        .options(joinedload(Member.user))
    )
    member = result.unique().scalar_one_or_none()
    if member is None:
        raise HTTPException(status_code=404, detail="Employee not found")
    return member


async def _get_maintenance_or_404(
    db: AsyncSession, asset_id: str, maintenance_id: str
) -> AssetMaintenanceLog:
    result = await db.execute(
        select(AssetMaintenanceLog).where(
            AssetMaintenanceLog.id == maintenance_id,
            AssetMaintenanceLog.assetId == asset_id,
        )
    )
    log = result.scalar_one_or_none()
    if log is None:
        raise HTTPException(status_code=404, detail="Maintenance record not found")
    return log


async def _get_category_or_404(
    db: AsyncSession, organization_id: str, category_id: str
) -> AssetCategoryDefinition:
    result = await db.execute(
        select(AssetCategoryDefinition)
        .where(
            AssetCategoryDefinition.id == category_id,
            AssetCategoryDefinition.organizationId == organization_id,
            AssetCategoryDefinition.isActive.is_(True),
        )
        .options(joinedload(AssetCategoryDefinition.fields))
    )
    category = result.unique().scalar_one_or_none()
    if category is None:
        raise HTTPException(status_code=404, detail="Asset category not found")
    return category


def _derive_asset_status(units: list[AssetUnit]) -> str:
    if not units:
        return "AVAILABLE"
    statuses = {_normalize_asset_status(u.status) or "AVAILABLE" for u in units}
    if all(s == "AVAILABLE" for s in statuses):
        return "AVAILABLE"
    if all(s in ("RETIRED", "DISPOSED") for s in statuses):
        return "RETIRED"
    if "PENDING_RETURN" in statuses:
        return "PENDING_RETURN"
    if "ASSIGNED" in statuses:
        return "ASSIGNED"
    if "IN_MAINTENANCE" in statuses:
        return "IN_MAINTENANCE"
    if "DAMAGED" in statuses:
        return "DAMAGED"
    if "LOST" in statuses:
        return "LOST"
    return "AVAILABLE"


def _unit_summary(units: list[AssetUnit]) -> AssetUnitSummary | None:
    if not units:
        return None
    summary = AssetUnitSummary(total=len(units))
    for unit in units:
        normalized_status = _normalize_asset_status(unit.status)
        if normalized_status == "AVAILABLE":
            summary.available += 1
        elif normalized_status == "ASSIGNED":
            summary.provided += 1
        elif normalized_status in {"IN_MAINTENANCE", "PENDING_RETURN"}:
            summary.underMaintenance += 1
        elif normalized_status == "DAMAGED":
            summary.damaged += 1
    return summary


async def _active_provision_map(
    db: AsyncSession, organization_id: str
) -> dict[str, AssetAssignment]:
    result = await db.execute(
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.organizationId == organization_id,
            Asset.deletedAt.is_(None),
            AssetAssignment.returnDate.is_(None),
        )
        .options(
            joinedload(AssetAssignment.member).joinedload(Member.user),
            joinedload(AssetAssignment.providedByMember).joinedload(Member.user),
            joinedload(AssetAssignment.receivedByMember).joinedload(Member.user),
        )
    )
    provisions = result.unique().scalars().all()
    return {provision.assetId: provision for provision in provisions}


async def _open_maintenance_counts(db: AsyncSession, asset_ids: list[str]) -> dict[str, int]:
    if not asset_ids:
        return {}

    result = await db.execute(
        select(AssetMaintenanceLog.assetId, func.count(AssetMaintenanceLog.id))
        .where(
            AssetMaintenanceLog.assetId.in_(asset_ids),
            AssetMaintenanceLog.status.in_(("OPEN", "IN_PROGRESS")),
        )
        .group_by(AssetMaintenanceLog.assetId)
    )
    return {asset_id: count for asset_id, count in result.all()}


def _asset_summary(
    asset: Asset,
    active_provision: AssetAssignment | None,
    open_maintenance_count: int | None = None,
) -> AssetSummary:
    holder = (
        active_provision.member.user
        if active_provision and active_provision.member and active_provision.member.user
        else None
    )
    if open_maintenance_count is None:
        open_maintenance_count = sum(
            1 for log in asset.maintenanceLogs if log.status in {"OPEN", "IN_PROGRESS"}
        )

    custom_fields = [
        CustomFieldValueResponse(
            fieldDefinitionId=cfv.fieldDefinitionId,
            fieldName=cfv.fieldDefinition.fieldName if cfv.fieldDefinition else "",
            fieldType=cfv.fieldDefinition.fieldType if cfv.fieldDefinition else "",
            value=cfv.value,
        )
        for cfv in asset.customFieldValues
        if cfv.fieldDefinition
    ]

    return AssetSummary(
        id=asset.id,
        assetCode=asset.assetCode,
        name=asset.name,
        category=asset.category,
        categoryDefinitionId=asset.categoryDefinitionId,
        serialNumber=asset.serialNumber,
        brand=asset.brand,
        model=asset.model,
        purchaseDate=asset.purchaseDate,
        purchasePrice=_to_float(asset.purchasePrice),
        warrantyExpiryDate=asset.warrantyExpiryDate,
        condition=asset.condition,
        status=_normalize_asset_status(asset.status) or asset.status,
        location=asset.location,
        notes=asset.notes,
        quantity=asset.quantity,
        createdAt=asset.createdAt,
        updatedAt=asset.updatedAt,
        currentHolderMemberId=active_provision.memberId if active_provision else None,
        currentHolderName=holder.name if holder else None,
        currentHolderEmail=holder.email if holder else None,
        openMaintenanceCount=open_maintenance_count,
        unitSummary=_unit_summary(asset.units),
        customFields=custom_fields,
    )


def _provide_record_summary(record: AssetAssignment) -> AssetProvideRecordSummary:
    holder = record.member.user if record.member and record.member.user else None
    provider = (
        record.providedByMember.user
        if record.providedByMember and record.providedByMember.user
        else None
    )
    receiver = (
        record.receivedByMember.user
        if record.receivedByMember and record.receivedByMember.user
        else None
    )
    return AssetProvideRecordSummary(
        id=record.id,
        assetUnitId=record.assetUnitId,
        memberId=record.memberId,
        memberName=holder.name if holder else None,
        memberEmail=holder.email if holder else None,
        providedByMemberId=record.providedByMemberId,
        providedByName=provider.name if provider else None,
        providedDate=record.providedDate,
        conditionWhileProviding=record.conditionWhileProviding,
        provideNotes=record.provideNotes,
        returnDate=record.returnDate,
        returnedCondition=record.returnedCondition,
        receivedByMemberId=record.receivedByMemberId,
        receivedByName=receiver.name if receiver else None,
        returnNotes=record.returnNotes,
        replacementAssignmentId=record.replacementAssignmentId,
        handoverRequestedAt=record.handoverRequestedAt,
        handoverCompletedAt=record.handoverCompletedAt,
        handoverConditionNotes=record.handoverConditionNotes,
    )


def _assignment_status_for_employee(asset: Asset, assignment: AssetAssignment) -> str:
    if assignment.assetUnitId:
        unit = next((item for item in (asset.units or []) if item.id == assignment.assetUnitId), None)
        if unit is not None:
            return _normalize_asset_status(unit.status) or unit.status
    return _normalize_asset_status(asset.status) or asset.status


def _employee_asset_state(asset: Asset, assignment: AssetAssignment) -> tuple[str, str]:
    assignment_status = _assignment_status_for_employee(asset, assignment)
    if assignment.returnDate is None and (
        assignment.replacementAssignmentId or assignment_status == "PENDING_RETURN"
    ):
        return ("RETURN_PENDING", "Return Pending")
    if assignment.returnDate is None:
        return ("CURRENT_ASSIGNED", "Assigned")
    if assignment.returnDate is not None and assignment_status == "IN_MAINTENANCE":
        return ("RETURNED_IN_REPAIR", "Returned / In Repair")
    return ("RETURNED", "Returned")


def _maintenance_asset_lifecycle(asset: Asset | None) -> tuple[str | None, str | None]:
    if asset is None:
        return (None, None)

    active_assignment = next((record for record in (asset.provisions or []) if record.returnDate is None), None)
    historical_return_exists = any(record.returnDate is not None for record in (asset.provisions or []))
    normalized_status = _normalize_asset_status(asset.status) or asset.status

    if active_assignment is not None and active_assignment.replacementAssignmentId:
        return ("PENDING_RETURN", "Return Pending")
    if historical_return_exists and normalized_status == "IN_MAINTENANCE":
        return ("RETURNED_IN_REPAIR", "Returned / In Repair")
    if historical_return_exists and normalized_status == "AVAILABLE":
        return ("RETURNED_READY", "Returned / Ready")
    if active_assignment is not None:
        return ("ASSIGNED", "Assigned")
    if normalized_status is None:
        return (None, None)
    return (normalized_status, _to_title(normalized_status))


def _is_return_request_ticket(log: AssetMaintenanceLog) -> bool:
    return _normalized_token(log.subject).lower() == "asset return request"


def _employee_asset_view_item(
    asset: Asset,
    assignment: AssetAssignment,
    active_provision: AssetAssignment | None,
) -> EmployeeAssetViewItem:
    asset_summary = _asset_summary(asset, active_provision)
    employee_state, employee_status_label = _employee_asset_state(asset, assignment)
    return EmployeeAssetViewItem(
        assignmentId=assignment.id,
        assetId=asset.id,
        assetUnitId=assignment.assetUnitId,
        assetCode=asset_summary.assetCode,
        name=asset_summary.name,
        category=asset_summary.category,
        categoryDefinitionId=asset_summary.categoryDefinitionId,
        serialNumber=asset_summary.serialNumber,
        brand=asset_summary.brand,
        model=asset_summary.model,
        purchaseDate=asset_summary.purchaseDate,
        purchasePrice=asset_summary.purchasePrice,
        warrantyExpiryDate=asset_summary.warrantyExpiryDate,
        condition=asset_summary.condition,
        status=asset_summary.status,
        location=asset_summary.location,
        notes=asset_summary.notes,
        quantity=asset_summary.quantity,
        createdAt=asset_summary.createdAt,
        updatedAt=asset_summary.updatedAt,
        currentHolderMemberId=asset_summary.currentHolderMemberId,
        currentHolderName=asset_summary.currentHolderName,
        currentHolderEmail=asset_summary.currentHolderEmail,
        openMaintenanceCount=asset_summary.openMaintenanceCount,
        unitSummary=asset_summary.unitSummary,
        customFields=asset_summary.customFields,
        providedDate=assignment.providedDate,
        returnDate=assignment.returnDate,
        returnedCondition=assignment.returnedCondition,
        returnNotes=assignment.returnNotes,
        handoverRequestedAt=assignment.handoverRequestedAt,
        handoverCompletedAt=assignment.handoverCompletedAt,
        handoverConditionNotes=assignment.handoverConditionNotes,
        replacementAssignmentId=assignment.replacementAssignmentId,
        employeeState=employee_state,
        employeeStatusLabel=employee_status_label,
    )


def _maintenance_summary(log: AssetMaintenanceLog) -> AssetMaintenanceSummary:
    actor = log.loggedByMember.user if log.loggedByMember and log.loggedByMember.user else None
    return AssetMaintenanceSummary(
        id=log.id,
        ticketId=log.ticketId,
        ticketMode=log.ticketMode,
        assetId=log.assetId,
        assetUnitId=log.assetUnitId,
        category=log.category,
        subject=log.subject,
        attachmentsMetadata=log.attachmentsMetadata or [],
        maintenanceType=log.maintenanceType,
        issueDescription=log.issueDescription,
        serviceDate=log.serviceDate,
        expectedCompletionDate=log.expectedCompletionDate,
        estimatedDowntimeHours=_resolved_downtime_hours(log),
        operationalCriticalityTier=log.operationalCriticalityTier,
        completedDate=log.completedDate,
        cost=_to_float(log.cost),
        status=log.status,
        conditionBeforeMaintenance=log.conditionBeforeMaintenance,
        conditionAfterMaintenance=log.conditionAfterMaintenance,
        notes=log.notes,
        replacementDecision=log.replacementDecision,
        replacementAssetUnitId=log.replacementAssetUnitId,
        loggedByMemberId=log.loggedByMemberId,
        loggedByName=actor.name if actor else None,
    )


def _ticket_title(log: AssetMaintenanceLog) -> str:
    if log.subject and log.subject.strip():
        return log.subject.strip()
    if log.asset and log.asset.name:
        return log.asset.name
    if log.category and log.category.strip():
        return _to_title(log.category)
    return "Help Request"


def _my_ticket_response(log: AssetMaintenanceLog) -> MyTicketResponse:
    return MyTicketResponse(
        id=log.id,
        ticketId=log.ticketId,
        ticketMode=log.ticketMode,
        assetId=log.assetId,
        assetName=log.asset.name if log.asset else None,
        assetCode=log.asset.assetCode if log.asset else None,
        category=log.category,
        subject=log.subject,
        attachmentsMetadata=log.attachmentsMetadata or [],
        maintenanceType=log.maintenanceType,
        issueDescription=log.issueDescription,
        status=log.status,
        serviceDate=log.serviceDate.isoformat() if log.serviceDate else "",
        createdAt=log.createdAt.isoformat() if log.createdAt else "",
        updatedAt=log.updatedAt.isoformat() if log.updatedAt else "",
    )


def _mark_ticket_cancelled(log: AssetMaintenanceLog, ctx: MemberContext) -> None:
    log.cancelledByMemberId = ctx.member.id


def _maintenance_ticket_response(log: AssetMaintenanceLog) -> MaintenanceTicketResponse:
    asset_lifecycle_status, asset_lifecycle_status_label = _maintenance_asset_lifecycle(log.asset)
    return MaintenanceTicketResponse(
        id=log.id,
        ticketId=log.ticketId,
        ticketMode=log.ticketMode,
        assetId=log.assetId,
        assetUnitId=log.assetUnitId,
        assetName=log.asset.name if log.asset else None,
        assetCode=log.asset.assetCode if log.asset else None,
        assetCondition=log.asset.condition if log.asset else None,
        category=log.category,
        subject=log.subject,
        attachmentsMetadata=log.attachmentsMetadata or [],
        maintenanceType=log.maintenanceType,
        issueDescription=log.issueDescription,
        status=log.status,
        serviceDate=log.serviceDate.isoformat() if log.serviceDate else "",
        expectedCompletionDate=(
            log.expectedCompletionDate.isoformat() if log.expectedCompletionDate else None
        ),
        estimatedDowntimeHours=_resolved_downtime_hours(log),
        operationalCriticalityTier=log.operationalCriticalityTier,
        replacementDecision=log.replacementDecision,
        createdAt=log.createdAt.isoformat() if log.createdAt else "",
        loggedByMemberId=log.loggedByMemberId,
        loggedByName=log.loggedByMember.user.name
        if log.loggedByMember and log.loggedByMember.user
        else None,
        loggedByEmail=log.loggedByMember.user.email
        if log.loggedByMember and log.loggedByMember.user
        else None,
        cancelledByMemberId=log.cancelledByMemberId,
        cancelledByName=log.cancelledByMember.user.name
        if log.cancelledByMember and log.cancelledByMember.user
        else None,
        assetLifecycleStatus=asset_lifecycle_status,
        assetLifecycleStatusLabel=asset_lifecycle_status_label,
    )


# ── Asset ID CRUD ────────────────────────────────────────────────────────────


async def create_asset_id(
    db: AsyncSession,
    ctx: MemberContext,
    payload: AssetIdCreate,
) -> AssetIdResponse:
    existing = await db.execute(
        select(AssetIdDefinition).where(
            AssetIdDefinition.organizationId == ctx.organization.id,
            AssetIdDefinition.assetIdName == payload.assetIdName.strip(),
            AssetIdDefinition.isActive.is_(True),
        )
    )
    if existing.scalar_one_or_none() is not None:
        raise HTTPException(status_code=409, detail="An asset ID with this name already exists")
    asset_id = AssetIdDefinition(
        organizationId=ctx.organization.id,
        assetIdName=payload.assetIdName.strip(),
    )
    db.add(asset_id)
    await db.commit()
    return AssetIdResponse(
        id=asset_id.id,
        assetIdName=asset_id.assetIdName,
        isActive=asset_id.isActive,
    )


async def list_asset_ids(db: AsyncSession, ctx: MemberContext) -> list[AssetIdResponse]:
    result = await db.execute(
        select(AssetIdDefinition)
        .where(
            AssetIdDefinition.organizationId == ctx.organization.id,
            AssetIdDefinition.isActive.is_(True),
        )
        .order_by(AssetIdDefinition.assetIdName.asc())
    )
    items = result.scalars().all()
    return [AssetIdResponse(id=a.id, assetIdName=a.assetIdName, isActive=a.isActive) for a in items]


async def update_asset_id(
    db: AsyncSession,
    ctx: MemberContext,
    asset_id_id: str,
    payload: AssetIdUpdate,
) -> AssetIdResponse:
    result = await db.execute(
        select(AssetIdDefinition).where(
            AssetIdDefinition.id == asset_id_id,
            AssetIdDefinition.organizationId == ctx.organization.id,
            AssetIdDefinition.isActive.is_(True),
        )
    )
    entry = result.scalar_one_or_none()
    if entry is None:
        raise HTTPException(status_code=404, detail="Asset ID not found")
    if payload.assetIdName is not None:
        dupe = await db.execute(
            select(AssetIdDefinition).where(
                AssetIdDefinition.organizationId == ctx.organization.id,
                AssetIdDefinition.assetIdName == payload.assetIdName.strip(),
                AssetIdDefinition.id != asset_id_id,
                AssetIdDefinition.isActive.is_(True),
            )
        )
        if dupe.scalar_one_or_none() is not None:
            raise HTTPException(status_code=409, detail="An asset ID with this name already exists")
        entry.assetIdName = payload.assetIdName.strip()
    await db.commit()
    return AssetIdResponse(
        id=entry.id,
        assetIdName=entry.assetIdName,
        isActive=entry.isActive,
    )


async def delete_asset_id(db: AsyncSession, ctx: MemberContext, asset_id_id: str) -> None:
    result = await db.execute(
        select(AssetIdDefinition).where(
            AssetIdDefinition.id == asset_id_id,
            AssetIdDefinition.organizationId == ctx.organization.id,
            AssetIdDefinition.isActive.is_(True),
        )
    )
    entry = result.scalar_one_or_none()
    if entry is None:
        raise HTTPException(status_code=404, detail="Asset ID not found")
    entry.isActive = False
    await db.commit()


# ── Category CRUD ─────────────────────────────────────────────────────────────


async def create_category(
    db: AsyncSession,
    ctx: MemberContext,
    payload: AssetCategoryCreate,
) -> AssetCategoryResponse:
    existing = await db.execute(
        select(AssetCategoryDefinition).where(
            AssetCategoryDefinition.organizationId == ctx.organization.id,
            AssetCategoryDefinition.name == payload.name.strip(),
        )
    )
    if existing.scalar_one_or_none() is not None:
        raise HTTPException(status_code=409, detail="A category with this name already exists")

    category = AssetCategoryDefinition(
        organizationId=ctx.organization.id,
        name=payload.name.strip(),
        assetCode=payload.assetCode.strip() if payload.assetCode else None,
        description=payload.description.strip() if payload.description else None,
    )
    db.add(category)
    await db.commit()
    if _is_laptop_category_name(category.name):
        await _ensure_laptop_os_field(db, ctx.organization.id)

    return AssetCategoryResponse(
        id=category.id,
        name=category.name,
        assetCode=category.assetCode,
        description=category.description,
        isActive=category.isActive,
        fields=[],
    )


async def list_categories(db: AsyncSession, ctx: MemberContext) -> list[AssetCategoryResponse]:
    await _ensure_laptop_os_field(db, ctx.organization.id)
    result = await db.execute(
        select(AssetCategoryDefinition)
        .where(
            AssetCategoryDefinition.organizationId == ctx.organization.id,
            AssetCategoryDefinition.isActive.is_(True),
        )
        .options(joinedload(AssetCategoryDefinition.fields))
        .order_by(AssetCategoryDefinition.name.asc())
    )
    categories = result.unique().scalars().all()
    return [
        AssetCategoryResponse(
            id=c.id,
            name=c.name,
            assetCode=c.assetCode,
            description=c.description,
            isActive=c.isActive,
            fields=[
                CategoryFieldDefinitionResponse(
                    id=f.id,
                    categoryId=f.categoryId,
                    fieldName=f.fieldName,
                    fieldType=f.fieldType,
                    fieldOptions=f.fieldOptions,
                    isRequired=f.isRequired,
                    displayOrder=f.displayOrder,
                )
                for f in (c.fields or [])
            ],
        )
        for c in categories
    ]


async def update_category(
    db: AsyncSession,
    ctx: MemberContext,
    category_id: str,
    payload: AssetCategoryUpdate,
) -> AssetCategoryResponse:
    category = await _get_category_or_404(db, ctx.organization.id, category_id)
    if payload.name is not None:
        existing = await db.execute(
            select(AssetCategoryDefinition).where(
                AssetCategoryDefinition.organizationId == ctx.organization.id,
                AssetCategoryDefinition.name == payload.name.strip(),
                AssetCategoryDefinition.id != category_id,
            )
        )
        if existing.scalar_one_or_none() is not None:
            raise HTTPException(status_code=409, detail="A category with this name already exists")
        category.name = payload.name.strip()
    if payload.description is not None:
        category.description = payload.description.strip() if payload.description else None
    if payload.assetCode is not None:
        category.assetCode = payload.assetCode.strip() if payload.assetCode else None
    await db.commit()
    if _is_laptop_category_name(category.name):
        await _ensure_laptop_os_field(db, ctx.organization.id)

    fields = [
        CategoryFieldDefinitionResponse(
            id=f.id,
            categoryId=f.categoryId,
            fieldName=f.fieldName,
            fieldType=f.fieldType,
            fieldOptions=f.fieldOptions,
            isRequired=f.isRequired,
            displayOrder=f.displayOrder,
        )
        for f in (category.fields or [])
    ]
    return AssetCategoryResponse(
        id=category.id,
        name=category.name,
        assetCode=category.assetCode,
        description=category.description,
        isActive=category.isActive,
        fields=fields,
    )


async def delete_category(db: AsyncSession, ctx: MemberContext, category_id: str) -> None:
    category = await _get_category_or_404(db, ctx.organization.id, category_id)
    assets_with_category = await db.execute(
        select(Asset)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.categoryDefinitionId == category_id,
            Asset.deletedAt.is_(None),
        )
        .limit(1)
    )
    if assets_with_category.scalar_one_or_none() is not None:
        raise HTTPException(
            status_code=422, detail="Cannot delete category that has assets assigned to it"
        )
    category.isActive = False
    await db.commit()


# ── Category Field CRUD ───────────────────────────────────────────────────────


async def create_category_field(
    db: AsyncSession,
    ctx: MemberContext,
    category_id: str,
    payload: CategoryFieldDefinitionCreate,
) -> CategoryFieldDefinitionResponse:
    category = await _get_category_or_404(db, ctx.organization.id, category_id)
    if _normalize_key(payload.fieldName) == "os" and not _is_laptop_category_name(category.name):
        raise HTTPException(
            status_code=422,
            detail='The "OS" custom field is reserved for the Laptop category',
        )

    field = AssetCategoryFieldDefinition(
        categoryId=category_id,
        fieldName=payload.fieldName.strip(),
        fieldType=payload.fieldType,
        fieldOptions={"options": payload.fieldOptions} if payload.fieldOptions else None,
        isRequired=payload.isRequired,
        displayOrder=payload.displayOrder,
    )
    db.add(field)
    await db.commit()

    return CategoryFieldDefinitionResponse(
        id=field.id,
        categoryId=field.categoryId,
        fieldName=field.fieldName,
        fieldType=field.fieldType,
        fieldOptions=field.fieldOptions,
        isRequired=field.isRequired,
        displayOrder=field.displayOrder,
    )


async def update_category_field(
    db: AsyncSession,
    ctx: MemberContext,
    field_id: str,
    payload: CategoryFieldDefinitionUpdate,
) -> CategoryFieldDefinitionResponse:
    result = await db.execute(
        select(AssetCategoryFieldDefinition)
        .join(
            AssetCategoryDefinition,
            AssetCategoryDefinition.id == AssetCategoryFieldDefinition.categoryId,
        )
        .where(
            AssetCategoryFieldDefinition.id == field_id,
            AssetCategoryDefinition.organizationId == ctx.organization.id,
        )
        .options(joinedload(AssetCategoryFieldDefinition.category))
    )
    field = result.unique().scalar_one_or_none()
    if field is None:
        raise HTTPException(status_code=404, detail="Category field not found")
    category_name = field.category.name if field.category else None
    incoming_field_name = payload.fieldName.strip() if payload.fieldName is not None else field.fieldName
    if _normalize_key(incoming_field_name) == "os" and not _is_laptop_category_name(category_name):
        raise HTTPException(
            status_code=422,
            detail='The "OS" custom field is reserved for the Laptop category',
        )

    if payload.fieldName is not None:
        field.fieldName = payload.fieldName.strip()
    if payload.fieldType is not None:
        field.fieldType = payload.fieldType
    if payload.fieldOptions is not None:
        field.fieldOptions = {"options": payload.fieldOptions} if payload.fieldOptions else None
    if payload.isRequired is not None:
        field.isRequired = payload.isRequired
    if payload.displayOrder is not None:
        field.displayOrder = payload.displayOrder
    await db.commit()

    return CategoryFieldDefinitionResponse(
        id=field.id,
        categoryId=field.categoryId,
        fieldName=field.fieldName,
        fieldType=field.fieldType,
        fieldOptions=field.fieldOptions,
        isRequired=field.isRequired,
        displayOrder=field.displayOrder,
    )


async def delete_category_field(db: AsyncSession, ctx: MemberContext, field_id: str) -> None:
    result = await db.execute(
        select(AssetCategoryFieldDefinition)
        .join(
            AssetCategoryDefinition,
            AssetCategoryDefinition.id == AssetCategoryFieldDefinition.categoryId,
        )
        .where(
            AssetCategoryFieldDefinition.id == field_id,
            AssetCategoryDefinition.organizationId == ctx.organization.id,
        )
    )
    field = result.unique().scalar_one_or_none()
    if field is None:
        raise HTTPException(status_code=404, detail="Category field not found")

    await db.execute(
        select(AssetCustomFieldValue)
        .where(AssetCustomFieldValue.fieldDefinitionId == field_id)
        .limit(1)
    )
    await db.delete(field)
    await db.commit()


# ── Asset CRUD (modified) ──────────────────────────────────────────────────────


async def list_assets(
    db: AsyncSession, ctx: MemberContext, filters: AssetFilters
) -> AssetListResponse:
    scope = getattr(ctx, "scope", "organization")

    active_provision_subquery = (
        select(
            AssetAssignment.assetId.label("assetId"),
            AssetAssignment.memberId.label("memberId"),
        )
        .where(AssetAssignment.returnDate.is_(None))
        .subquery()
    )

    query: Select = (
        select(Asset)
        .where(Asset.organizationId == ctx.organization.id, Asset.deletedAt.is_(None))
        .order_by(Asset.updatedAt.desc(), Asset.createdAt.desc())
    )
    joined_active_provision = False

    if scope == "self":
        query = query.join(
            active_provision_subquery, active_provision_subquery.c.assetId == Asset.id
        ).where(
            active_provision_subquery.c.memberId == ctx.member.id,
        )
        joined_active_provision = True

    if filters.currentHolderMemberId:
        if not joined_active_provision:
            query = query.join(
                active_provision_subquery, active_provision_subquery.c.assetId == Asset.id
            )
            joined_active_provision = True
        query = query.where(active_provision_subquery.c.memberId == filters.currentHolderMemberId)

    if filters.search:
        term = f"%{filters.search.strip()}%"
        query = query.where(
            or_(
                Asset.name.ilike(term),
                Asset.assetCode.ilike(term),
                Asset.serialNumber.ilike(term),
                Asset.model.ilike(term),
                Asset.location.ilike(term),
            )
        )
    if filters.category:
        query = query.where(Asset.category == filters.category)
    if filters.categoryDefinitionId:
        query = query.where(Asset.categoryDefinitionId == filters.categoryDefinitionId)
    if filters.status:
        if filters.status == "ASSIGNED":
            query = query.where(Asset.status.in_(("ASSIGNED", "PROVIDED")))
        elif filters.status == "IN_MAINTENANCE":
            query = query.where(Asset.status.in_(("IN_MAINTENANCE", "UNDER_MAINTENANCE")))
        else:
            query = query.where(Asset.status == filters.status)

    total_result = await db.execute(
        select(func.count()).select_from(query.order_by(None).subquery())
    )
    total = total_result.scalar_one()

    rows = await db.execute(
        query.offset((filters.page - 1) * filters.page_size)
        .limit(filters.page_size)
        .options(
            joinedload(Asset.provisions),
            joinedload(Asset.units),
            joinedload(Asset.customFieldValues).joinedload(AssetCustomFieldValue.fieldDefinition),
        )
    )
    assets = rows.unique().scalars().all()

    active_map = await _active_provision_map(db, ctx.organization.id)
    maintenance_counts = await _open_maintenance_counts(db, [asset.id for asset in assets])
    items = [
        _asset_summary(asset, active_map.get(asset.id), maintenance_counts.get(asset.id, 0))
        for asset in assets
    ]

    return AssetListResponse(
        items=items,
        total=total,
        page=filters.page,
        page_size=filters.page_size,
        overdue_count=0,
    )


async def get_employee_asset_view(
    db: AsyncSession,
    ctx: MemberContext,
) -> EmployeeAssetViewResponse:
    result = await db.execute(
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
            AssetAssignment.memberId == ctx.member.id,
        )
        .options(
            joinedload(AssetAssignment.member).joinedload(Member.user),
            joinedload(AssetAssignment.providedByMember).joinedload(Member.user),
            joinedload(AssetAssignment.receivedByMember).joinedload(Member.user),
            joinedload(AssetAssignment.asset).joinedload(Asset.units),
            joinedload(AssetAssignment.asset).joinedload(Asset.maintenanceLogs),
            joinedload(AssetAssignment.asset)
            .joinedload(Asset.customFieldValues)
            .joinedload(AssetCustomFieldValue.fieldDefinition),
            joinedload(AssetAssignment.asset)
            .joinedload(Asset.provisions)
            .joinedload(AssetAssignment.member)
            .joinedload(Member.user),
        )
        .order_by(AssetAssignment.providedDate.desc(), AssetAssignment.createdAt.desc())
    )
    assignments = result.unique().scalars().all()

    current: list[EmployeeAssetViewItem] = []
    previous: list[EmployeeAssetViewItem] = []

    for assignment in assignments:
        asset = assignment.asset
        if asset is None:
            continue
        active_provision = next((record for record in asset.provisions if record.returnDate is None), None)
        item = _employee_asset_view_item(asset, assignment, active_provision)
        if item.employeeState == "CURRENT_ASSIGNED":
            current.append(item)
        else:
            previous.append(item)

    return EmployeeAssetViewResponse(current=current, previous=previous)


async def get_asset(db: AsyncSession, ctx: MemberContext, asset_id: str) -> AssetDetailResponse:
    asset = await _get_asset_or_404(db, ctx.organization.id, asset_id)

    provisions_result = await db.execute(
        select(AssetAssignment)
        .where(AssetAssignment.assetId == asset_id)
        .options(
            joinedload(AssetAssignment.member).joinedload(Member.user),
            joinedload(AssetAssignment.providedByMember).joinedload(Member.user),
            joinedload(AssetAssignment.receivedByMember).joinedload(Member.user),
        )
        .order_by(AssetAssignment.providedDate.desc())
    )
    provisions = provisions_result.unique().scalars().all()
    active_provision = next((record for record in provisions if record.returnDate is None), None)

    scope = getattr(ctx, "scope", "organization")
    if scope == "self" and not any(record.memberId == ctx.member.id for record in provisions):
        raise HTTPException(status_code=404, detail="Asset not found")

    logs_result = await db.execute(
        select(AssetMaintenanceLog)
        .where(AssetMaintenanceLog.assetId == asset_id)
        .options(joinedload(AssetMaintenanceLog.loggedByMember).joinedload(Member.user))
        .order_by(AssetMaintenanceLog.createdAt.desc())
    )
    logs = logs_result.unique().scalars().all()

    units_result = await db.execute(select(AssetUnit).where(AssetUnit.assetId == asset_id))
    units = units_result.scalars().all()

    unit_responses = [
        AssetUnitResponse(
            id=u.id,
            assetId=u.assetId,
            serialNumber=u.serialNumber,
            status=_normalize_asset_status(u.status) or u.status,
            currentHolderMemberId=u.currentHolderMemberId,
            currentHolderName="",
            condition=u.condition,
        )
        for u in units
    ]

    base = _asset_summary(asset, active_provision)
    return AssetDetailResponse(
        **base.model_dump(),
        activeProvision=_provide_record_summary(active_provision) if active_provision else None,
        assetHistory=[_provide_record_summary(record) for record in provisions],
        maintenanceHistory=[_maintenance_summary(log) for log in logs],
        units=unit_responses,
    )


async def upsert_asset(
    db: AsyncSession,
    ctx: MemberContext,
    payload: AssetUpsertRequest,
    asset_id: str | None = None,
) -> AssetDetailResponse:
    is_new = asset_id is None
    if is_new:
        asset = Asset(id=generate_uuid(), organizationId=ctx.organization.id)
        db.add(asset)
    else:
        asset = await _get_asset_or_404(db, ctx.organization.id, asset_id)
    previous_warranty_expiry = asset.warrantyExpiryDate if not is_new else None

    # Check active provision via query to avoid lazy relationship access
    if not is_new:
        provision_result = await db.execute(
            select(AssetAssignment)
            .where(
                AssetAssignment.assetId == asset.id,
                AssetAssignment.returnDate.is_(None),
            )
            .limit(1)
        )
        active_provision = provision_result.scalar_one_or_none()
    else:
        active_provision = None

    if active_provision is not None and payload.status != "ASSIGNED":
        raise HTTPException(
            status_code=422, detail="Assigned assets must be returned before changing their status"
        )
    if active_provision is None and payload.status == "ASSIGNED":
        raise HTTPException(
            status_code=422, detail="Use Provide Asset to mark an asset as assigned"
        )

    asset.assetCode = payload.assetCode.strip()
    asset.name = payload.name.strip()
    asset.categoryDefinitionId = payload.categoryDefinitionId

    # Derive category from definition if not provided explicitly
    if payload.category is not None:
        asset.category = payload.category
    elif payload.categoryDefinitionId:
        cat_result = await db.execute(
            select(AssetCategoryDefinition).where(
                AssetCategoryDefinition.id == payload.categoryDefinitionId,
                AssetCategoryDefinition.organizationId == ctx.organization.id,
            )
        )
        cat_def = cat_result.scalar_one_or_none()
        if cat_def:
            asset.category = cat_def.name.upper().replace(" ", "_")
        else:
            asset.category = "OTHER"
    else:
        asset.category = "OTHER"
    asset.serialNumber = payload.serialNumber.strip() if payload.serialNumber else None
    asset.brand = payload.brand.strip() if payload.brand else None
    asset.model = payload.model.strip() if payload.model else None
    asset.purchaseDate = payload.purchaseDate
    asset.purchasePrice = payload.purchasePrice
    asset.warrantyExpiryDate = payload.warrantyExpiryDate
    asset.condition = payload.condition
    # New assets always land in the register as available. Lifecycle transitions
    # into assigned or maintenance states must happen through those flows.
    asset.status = "AVAILABLE" if is_new else payload.status
    asset.location = payload.location.strip() if payload.location else None
    asset.notes = payload.notes.strip() if payload.notes else None
    asset.quantity = payload.quantity

    # asset.id is set via generate_uuid() at creation — no flush needed

    # Handle custom fields
    if not is_new:
        result = await db.execute(
            select(AssetCustomFieldValue).where(AssetCustomFieldValue.assetId == asset.id)
        )
        existing_cfvs = result.scalars().all()
        incoming_cf_ids = {cf.fieldDefinitionId for cf in (payload.customFields or [])}
        for old in existing_cfvs:
            if old.fieldDefinitionId not in incoming_cf_ids:
                await db.delete(old)

    for cf in payload.customFields or []:
        if not is_new:
            result = await db.execute(
                select(AssetCustomFieldValue).where(
                    AssetCustomFieldValue.assetId == asset.id,
                    AssetCustomFieldValue.fieldDefinitionId == cf.fieldDefinitionId,
                )
            )
            existing_cfv = result.scalar_one_or_none()
            if existing_cfv:
                existing_cfv.value = cf.value
                continue

        db.add(
            AssetCustomFieldValue(
                assetId=asset.id,
                fieldDefinitionId=cf.fieldDefinitionId,
                value=cf.value,
            )
        )

    # Handle units
    if not is_new:
        result = await db.execute(select(AssetUnit).where(AssetUnit.assetId == asset.id))
        existing_units = result.scalars().all()
        existing_count = len(existing_units)
        incoming_units = payload.units or []
        if incoming_units:
            for i, unit_input in enumerate(incoming_units):
                if i < existing_count:
                    existing_units[i].serialNumber = (
                        unit_input.serialNumber.strip() if unit_input.serialNumber else None
                    )
                    existing_units[i].condition = payload.condition
                    existing_units[i].warrantyExpiryDate = payload.warrantyExpiryDate
                    if previous_warranty_expiry != payload.warrantyExpiryDate:
                        existing_units[i].lastWarrantyAlertSentAt = None
                        existing_units[i].reminderCompleted = False
                else:
                    db.add(
                        AssetUnit(
                            assetId=asset.id,
                            serialNumber=unit_input.serialNumber.strip()
                            if unit_input.serialNumber
                            else None,
                            status="AVAILABLE",
                            condition=payload.condition,
                            warrantyExpiryDate=payload.warrantyExpiryDate,
                            reminderCompleted=False,
                        )
                    )
            if len(incoming_units) < existing_count:
                for unit in existing_units[len(incoming_units) :]:
                    await db.delete(unit)
        else:
            if existing_count == 0:
                tracking_unit = await _ensure_tracking_unit(db, asset)
                tracking_unit.serialNumber = asset.serialNumber
            for unit in existing_units or ([tracking_unit] if existing_count == 0 else []):
                unit.condition = payload.condition
                unit.warrantyExpiryDate = payload.warrantyExpiryDate
                if previous_warranty_expiry != payload.warrantyExpiryDate:
                    unit.lastWarrantyAlertSentAt = None
                    unit.reminderCompleted = False
    else:
        for unit_input in payload.units or []:
            db.add(
                AssetUnit(
                    assetId=asset.id,
                    serialNumber=unit_input.serialNumber.strip()
                    if unit_input.serialNumber
                    else None,
                    status="AVAILABLE",
                    condition=payload.condition,
                    warrantyExpiryDate=payload.warrantyExpiryDate,
                    reminderCompleted=False,
                )
            )

    await db.commit()
    return await get_asset(db, ctx, asset.id)


# ── Bulk Asset Creation ──────────────────────────────────────────────────────


async def bulk_create_assets(
    db: AsyncSession,
    ctx: MemberContext,
    payload: BulkAssetCreateRequest,
) -> list[AssetDetailResponse]:
    org_id = ctx.organization.id
    quantity = len(payload.serialNumbers)

    if quantity == 0:
        raise HTTPException(status_code=422, detail="At least one serial number is required")

    # Validate category exists if provided
    category_name = "OTHER"
    if payload.categoryDefinitionId:
        cat = await _get_category_or_404(db, org_id, payload.categoryDefinitionId)
        category_name = cat.name.upper().replace(" ", "_")

    # Validate custom field definitions if provided
    if payload.customFields:
        for cf in payload.customFields:
            result = await db.execute(
                select(AssetCategoryFieldDefinition).where(
                    AssetCategoryFieldDefinition.id == cf.fieldDefinitionId,
                )
            )
            field_def = result.scalar_one_or_none()
            if field_def is None:
                raise HTTPException(
                    status_code=422,
                    detail=f"Custom field definition {cf.fieldDefinitionId} not found",
                )
            if field_def.isRequired and (not cf.value or not cf.value.strip()):
                raise HTTPException(
                    status_code=422, detail=f"Required field '{field_def.fieldName}' is missing"
                )

    asset_code = payload.assetCode.strip()
    condition = payload.condition
    location = payload.location.strip() if payload.location else None
    brand = payload.brand.strip() if payload.brand else None
    model = payload.model.strip() if payload.model else None

    created_assets: list[Asset] = []

    for serial in payload.serialNumbers:
        serial_val = serial.strip()
        asset = Asset(
            organizationId=org_id,
            assetCode=asset_code,
            name=payload.name.strip(),
            category=category_name,
            categoryDefinitionId=payload.categoryDefinitionId,
            serialNumber=serial_val,
            brand=brand,
            model=model,
            condition=condition,
            status="AVAILABLE",
            location=location,
            quantity=1,
        )
        db.add(asset)
        await db.flush()

        db.add(
            AssetUnit(
                assetId=asset.id,
                serialNumber=serial_val,
                status="AVAILABLE",
                condition=condition,
                warrantyExpiryDate=asset.warrantyExpiryDate,
                reminderCompleted=False,
            )
        )

        # Create custom field values
        for cf in payload.customFields or []:
            db.add(
                AssetCustomFieldValue(
                    assetId=asset.id,
                    fieldDefinitionId=cf.fieldDefinitionId,
                    value=cf.value,
                )
            )

        created_assets.append(asset)

    await db.commit()

    # Refresh and return
    results = []
    for asset in created_assets:
        result = await get_asset(db, ctx, asset.id)
        results.append(result)
    return results


# ── Available Groups ─────────────────────────────────────────────────────────


def _group_key(name: str, cat_def_id: str | None, asset_code: str) -> str:
    raw = f"{name}|{cat_def_id or ''}|{asset_code}"
    return hashlib.md5(raw.encode()).hexdigest()[:16]


async def list_available_groups(
    db: AsyncSession,
    ctx: MemberContext,
) -> list[AvailableAssetGroupResponse]:
    org_id = ctx.organization.id
    result = await db.execute(
        select(Asset)
        .where(
            Asset.organizationId == org_id,
            Asset.deletedAt.is_(None),
            Asset.status == "AVAILABLE",
        )
        .order_by(Asset.name.asc(), Asset.assetCode.asc())
    )
    assets = result.scalars().all()

    groups: dict[str, dict] = {}
    for asset in assets:
        key = _group_key(asset.name, asset.categoryDefinitionId, asset.assetCode)
        if key not in groups:
            cat_name = None
            if asset.categoryDefinitionId:
                cat_result = await db.execute(
                    select(AssetCategoryDefinition.name).where(
                        AssetCategoryDefinition.id == asset.categoryDefinitionId,
                    )
                )
                cat_name = cat_result.scalar_one_or_none()

            groups[key] = {
                "groupKey": key,
                "assetName": asset.name,
                "categoryName": cat_name,
                "categoryDefinitionId": asset.categoryDefinitionId,
                "assetCode": asset.assetCode,
                "brand": asset.brand,
                "model": asset.model,
                "availableQuantity": 0,
            }
        groups[key]["availableQuantity"] += 1

    return [
        AvailableAssetGroupResponse(**g)
        for g in sorted(groups.values(), key=lambda x: (-x["availableQuantity"], x["assetName"]))
    ]


# ── Issue Assets ─────────────────────────────────────────────────────────────


async def issue_assets(
    db: AsyncSession,
    ctx: MemberContext,
    payload: AssetIssueRequest,
) -> AssetIssueResponse:
    org_id = ctx.organization.id

    # Validate member exists
    await _get_member_or_404(db, org_id, payload.memberId)

    provider_id = payload.providedByMemberId or ctx.member.id
    if provider_id != payload.memberId:
        await _get_member_or_404(db, org_id, provider_id)

    # Lock available assets matching the group key
    # We need to find the group definition from any matching asset
    # Search all available assets and match by group key
    all_available = await db.execute(
        select(Asset)
        .where(
            Asset.organizationId == org_id,
            Asset.deletedAt.is_(None),
            Asset.status == "AVAILABLE",
        )
        .options(selectinload(Asset.units))
        .order_by(Asset.name.asc(), Asset.assetCode.asc())
        .with_for_update(skip_locked=True)
    )
    candidates = all_available.scalars().all()

    # Filter by group key
    matching = []
    for asset in candidates:
        key = _group_key(asset.name, asset.categoryDefinitionId, asset.assetCode)
        if key == payload.groupKey:
            matching.append(asset)

    if len(matching) < payload.quantity:
        raise HTTPException(
            status_code=422,
            detail=f"Only {len(matching)} assets available in this group, but {payload.quantity} requested",
        )

    to_issue = matching[: payload.quantity]
    issued_ids: list[str] = []
    assignment_ids: list[str] = []

    for asset in to_issue:
        tracking_unit = await _ensure_tracking_unit(db, asset)
        assignment = AssetAssignment(
            assetId=asset.id,
            assetUnitId=tracking_unit.id,
            memberId=payload.memberId,
            providedByMemberId=provider_id,
            providedDate=datetime.now(UTC),
            conditionWhileProviding=payload.conditionWhileProviding,
            provideNotes=payload.notes.strip() if payload.notes else None,
        )
        db.add(assignment)
        await db.flush()

        asset.status = "ASSIGNED"
        tracking_unit.status = "ASSIGNED"
        tracking_unit.currentHolderMemberId = payload.memberId
        tracking_unit.condition = payload.conditionWhileProviding

        issued_ids.append(asset.id)
        assignment_ids.append(assignment.id)

    await db.commit()
    return AssetIssueResponse(issuedAssetIds=issued_ids, assignmentIds=assignment_ids)


async def delete_asset(db: AsyncSession, ctx: MemberContext, asset_id: str) -> None:
    asset = await _get_asset_or_404(db, ctx.organization.id, asset_id)
    active_provision = next(
        (record for record in asset.provisions if record.returnDate is None), None
    )
    if active_provision is not None:
        raise HTTPException(status_code=422, detail="Return the provided asset before archiving it")
    asset.deletedAt = datetime.now(UTC)
    await db.commit()


# ── Unit-based Provide/Return ──────────────────────────────────────────────────


async def return_asset(
    db: AsyncSession,
    ctx: MemberContext,
    asset_id: str,
    payload: AssetReturnRequest,
) -> AssetDetailResponse:
    asset = await _get_asset_or_404(db, ctx.organization.id, asset_id)

    # Filter provisions - if assetUnitId specified, match that specific unit
    provisions_pool = asset.provisions
    if payload.assetUnitId:
        provisions_pool = [p for p in provisions_pool if p.assetUnitId == payload.assetUnitId]

    provision = next((record for record in provisions_pool if record.returnDate is None), None)
    if provision is None:
        raise HTTPException(status_code=422, detail="This asset is not currently assigned")
    if provision.memberId != payload.memberId:
        raise HTTPException(
            status_code=422, detail="The selected employee does not hold this asset"
        )

    receiver_id = payload.receivedByMemberId or ctx.member.id
    await _get_member_or_404(db, ctx.organization.id, receiver_id)

    next_status = payload.nextStatus
    if next_status is None:
        next_status = (
            "AVAILABLE"
            if payload.returnedCondition in {"NEW", "GOOD", "FAIR"}
            else "IN_MAINTENANCE"
        )

    provision.returnDate = payload.returnDate or datetime.now(UTC)
    provision.returnedCondition = payload.returnedCondition
    provision.receivedByMemberId = receiver_id
    provision.returnNotes = payload.returnNotes.strip() if payload.returnNotes else None
    if provision.handoverRequestedAt and provision.handoverCompletedAt is None:
        provision.handoverCompletedAt = provision.returnDate
    if payload.returnNotes:
        provision.handoverConditionNotes = payload.returnNotes.strip()

    asset.condition = payload.returnedCondition
    unit_id = payload.assetUnitId or provision.assetUnitId

    # Update the specific unit
    if unit_id:
        unit = next((u for u in (asset.units or []) if u.id == unit_id), None)
        if unit:
            unit.status = next_status
            unit.condition = payload.returnedCondition
            unit.currentHolderMemberId = None

        if asset.units and len(asset.units) > 0:
            asset.status = _derive_asset_status(asset.units)
        else:
            asset.status = next_status
    else:
        asset.status = next_status

    await db.commit()
    return await get_asset(db, ctx, asset.id)


async def request_asset_return(
    db: AsyncSession,
    ctx: MemberContext,
    asset_id: str,
) -> dict:
    asset = await _get_asset_or_404(db, ctx.organization.id, asset_id)

    provision = next(
        (record for record in asset.provisions if record.returnDate is None), None
    )
    if provision is None:
        raise HTTPException(
            status_code=422,
            detail="This asset is not currently assigned to anyone",
        )

    if provision.handoverRequestedAt is not None:
        raise HTTPException(
            status_code=422,
            detail="Return has already been requested for this asset",
        )

    provision.handoverRequestedAt = datetime.now(UTC)

    employee = await _get_member_or_404(db, ctx.organization.id, provision.memberId)

    notification = AssetNotification(
        organizationId=ctx.organization.id,
        assetId=asset.id,
        assetUnitId=provision.assetUnitId,
        memberId=provision.memberId,
        type="RETURN_REQUESTED",
        title="Asset return requested",
        message=(
            f"Your {asset.name} ({asset.assetCode}) return has been requested "
            f"by admin. Please hand over the device at your earliest convenience."
        ),
    )
    db.add(notification)

    await db.commit()
    return {"success": True, "message": f"Return requested for {asset.name}"}


async def get_returned_assets(
    db: AsyncSession,
    ctx: MemberContext,
) -> list[ReturnedAssetSummary]:
    result = await db.execute(
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
            AssetAssignment.returnDate.is_not(None),
        )
        .options(
            joinedload(AssetAssignment.asset),
            joinedload(AssetAssignment.member).joinedload(Member.user),
        )
        .order_by(AssetAssignment.returnDate.desc())
    )
    assignments = result.unique().scalars().all()

    items: list[ReturnedAssetSummary] = []
    for assignment in assignments:
        asset = assignment.asset
        member = assignment.member

        employee_name = member.user.name if member and member.user else None
        employee_email = member.user.email if member and member.user else None

        is_temp = assignment.replacementAssignmentId is not None

        ticket_info = None
        if asset:
            ticket_result = await db.execute(
                select(AssetMaintenanceLog)
                .where(
                    AssetMaintenanceLog.assetId == asset.id,
                    AssetMaintenanceLog.organizationId == ctx.organization.id,
                )
                .order_by(AssetMaintenanceLog.createdAt.desc())
                .limit(1)
            )
            ticket_info = ticket_result.unique().scalars().first()

        items.append(
            ReturnedAssetSummary(
                id=assignment.id,
                assetId=asset.id if asset else "",
                assetName=asset.name if asset else "",
                assetCode=asset.assetCode if asset else "",
                serialNumber=asset.serialNumber if asset else None,
                category=asset.category if asset else "",
                brand=asset.brand if asset else None,
                condition=asset.condition if asset else None,
                employeeMemberId=assignment.memberId,
                employeeName=employee_name,
                employeeEmail=employee_email,
                providedDate=assignment.providedDate.isoformat(),
                returnDate=assignment.returnDate.isoformat(),
                returnedCondition=assignment.returnedCondition,
                returnNotes=assignment.returnNotes,
                isTemporaryReplacement=is_temp,
                replacementAssetName=None,
                hasTicket=ticket_info is not None,
                ticketId=ticket_info.ticketId if ticket_info else None,
                maintenanceType=ticket_info.maintenanceType if ticket_info else None,
                maintenanceStatus=ticket_info.status if ticket_info else None,
                issueDescription=ticket_info.issueDescription if ticket_info else None,
            )
        )

    return items


# ── Replacements ────────────────────────────────────────────────────────────────


async def list_member_tickets(
    db: AsyncSession, ctx: MemberContext, member_id: str
) -> list[MemberTicketSummary]:
    result = await db.execute(
        select(AssetMaintenanceLog)
        .where(
            AssetMaintenanceLog.organizationId == ctx.organization.id,
            AssetMaintenanceLog.loggedByMemberId == member_id,
            AssetMaintenanceLog.status.in_(["OPEN", "IN_PROGRESS"]),
        )
        .order_by(AssetMaintenanceLog.createdAt.desc())
    )
    logs = result.scalars().all()
    return [
        MemberTicketSummary(
            id=log.id,
            ticketId=log.ticketId,
            maintenanceType=log.maintenanceType,
            issueDescription=log.issueDescription,
            status=log.status,
            replacementDecision=log.replacementDecision,
        )
        for log in logs
    ]


async def list_replacements(db: AsyncSession, ctx: MemberContext) -> list[ReplacementRecord]:
    result = await db.execute(
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
            AssetAssignment.replacementAssignmentId.is_not(None),
        )
        .options(
            joinedload(AssetAssignment.asset),
            joinedload(AssetAssignment.member).joinedload(Member.user),
            joinedload(AssetAssignment.receivedByMember).joinedload(Member.user),
        )
        .order_by(AssetAssignment.createdAt.desc())
    )
    old_assignments = result.unique().scalars().all()

    records: list[ReplacementRecord] = []
    for old in old_assignments:
        new_assignment = None
        if old.replacementAssignmentId:
            new_result = await db.execute(
                select(AssetAssignment)
                .where(AssetAssignment.id == old.replacementAssignmentId)
                .options(
                    joinedload(AssetAssignment.asset),
                    joinedload(AssetAssignment.providedByMember).joinedload(Member.user),
                )
            )
            new_assignment = new_result.unique().scalar_one_or_none()

        old_asset = old.asset
        new_asset = new_assignment.asset if new_assignment else None

        employee = old.member
        employee_name = employee.user.name if employee and employee.user else None
        employee_email = employee.user.email if employee and employee.user else None

        provider = new_assignment.providedByMember if new_assignment else None
        provider_name = provider.user.name if provider and provider.user else None

        ticket_info = None
        if old_asset:
            ticket_result = await db.execute(
                select(AssetMaintenanceLog)
                .where(
                    AssetMaintenanceLog.assetId == old_asset.id,
                    AssetMaintenanceLog.organizationId == ctx.organization.id,
                )
                .order_by(AssetMaintenanceLog.createdAt.desc())
                .limit(1)
            )
            ticket_info = ticket_result.unique().scalars().first()

        records.append(
            ReplacementRecord(
                id=old.id,
                employeeMemberId=old.memberId,
                employeeName=employee_name,
                employeeEmail=employee_email,
                originalAssetId=old_asset.id if old_asset else "",
                originalAssetName=old_asset.name if old_asset else "",
                originalAssetCode=old_asset.assetCode if old_asset else "",
                originalSerial=old_asset.serialNumber if old_asset else None,
                originalUnitStatus=old_asset.status if old_asset else None,
                replacementAssignmentId=new_assignment.id if new_assignment else None,
                replacementAssetId=new_asset.id if new_asset else "",
                replacementAssetName=new_asset.name if new_asset else "",
                replacementAssetCode=new_asset.assetCode if new_asset else "",
                replacementSerial=new_asset.serialNumber if new_asset else None,
                replacementMode=ticket_info.replacementDecision if ticket_info else "",
                expectedReturnDate=(
                    new_assignment.expectedReturnDate.isoformat()
                    if new_assignment and new_assignment.expectedReturnDate
                    else None
                ),
                returnReminderSent=new_assignment.returnReminderSent if new_assignment else False,
                providedByMemberId=provider.id if provider else None,
                providedByName=provider_name,
                ticketId=ticket_info.ticketId if ticket_info else None,
                maintenanceType=ticket_info.maintenanceType if ticket_info else None,
                maintenanceStatus=ticket_info.status if ticket_info else None,
                issueDescription=ticket_info.issueDescription if ticket_info else None,
                returnDate=old.returnDate.isoformat() if old.returnDate else "",
                swapCompletedAt=old.createdAt.isoformat(),
            )
        )

    return records


async def provide_replacement(
    db: AsyncSession,
    ctx: MemberContext,
    payload: ReplacementProvideRequest,
) -> ReplacementRecord:
    employee = await db.execute(
        select(Member)
        .where(Member.id == payload.employeeMemberId, Member.organizationId == ctx.organization.id)
    )
    employee_member = employee.scalar_one_or_none()
    if employee_member is None:
        raise HTTPException(status_code=404, detail="Employee not found")

    ticket_result = await db.execute(
        select(AssetMaintenanceLog)
        .where(
            AssetMaintenanceLog.ticketId == payload.ticketId,
            AssetMaintenanceLog.organizationId == ctx.organization.id,
            AssetMaintenanceLog.loggedByMemberId == payload.employeeMemberId,
        )
        .options(
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.provisions).joinedload(
                AssetAssignment.member
            ).joinedload(Member.user),
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.units),
        )
    )
    ticket = ticket_result.unique().scalar_one_or_none()
    if ticket is None:
        raise HTTPException(status_code=404, detail="Ticket not found for this employee")
    if ticket.status not in ("OPEN", "IN_PROGRESS"):
        raise HTTPException(status_code=422, detail="Ticket is not open or in progress")
    if ticket.asset is None:
        raise HTTPException(status_code=422, detail="Ticket has no associated asset")

    asset = ticket.asset
    active_assignment = next(
        (rec for rec in asset.provisions if rec.returnDate is None and rec.memberId == payload.employeeMemberId),
        None,
    )
    if active_assignment is None:
        raise HTTPException(status_code=422, detail="Employee does not have this asset currently assigned")

    preview = await _build_swap_preview(db, ctx.organization.id, ticket, asset)
    selected_option = next(
        (opt for opt in preview.options if opt.mode == payload.replacementMode and opt.available), None
    )
    if selected_option is None or not selected_option.assetUnitIds:
        raise HTTPException(
            status_code=422,
            detail=f"No available {payload.replacementMode.lower().replace('_', ' ')} units found",
        )

    replacement_unit_id = selected_option.assetUnitIds[0]
    unit_result = await db.execute(
        select(AssetUnit)
        .join(Asset, Asset.id == AssetUnit.assetId)
        .where(
            AssetUnit.id == replacement_unit_id,
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
        )
        .options(joinedload(AssetUnit.asset).joinedload(Asset.units))
    )
    replacement_unit = unit_result.unique().scalar_one_or_none()
    if replacement_unit is None or replacement_unit.asset is None:
        raise HTTPException(status_code=404, detail="Replacement unit not found")
    if _normalize_asset_status(replacement_unit.status) != "AVAILABLE":
        raise HTTPException(status_code=422, detail="Replacement unit is no longer available")

    replacement_asset = replacement_unit.asset
    now = datetime.now(UTC)

    source_unit = None
    if ticket.assetUnitId:
        source_unit = next((unit for unit in asset.units if unit.id == ticket.assetUnitId), None)

    new_assignment = AssetAssignment(
        assetId=replacement_asset.id,
        assetUnitId=replacement_unit.id,
        memberId=payload.employeeMemberId,
        providedByMemberId=ctx.member.id,
        providedDate=now,
        conditionWhileProviding="GOOD",
        provideNotes=payload.notes.strip() if payload.notes else None,
    )
    if payload.replacementMode == "TEMPORARY_BACKUP" and payload.expectedReturnDate:
        try:
            new_assignment.expectedReturnDate = datetime.fromisoformat(payload.expectedReturnDate)
        except (ValueError, TypeError):
            raise HTTPException(status_code=422, detail="Invalid expectedReturnDate format")

    db.add(new_assignment)
    await db.flush()

    active_assignment.replacementAssignmentId = new_assignment.id
    active_assignment.handoverRequestedAt = now

    active_assignment.returnDate = now
    active_assignment.returnedCondition = ticket.conditionBeforeMaintenance or asset.condition
    active_assignment.receivedByMemberId = ctx.member.id
    active_assignment.returnNotes = (
        payload.notes.strip() if payload.notes else "Revoked automatically during replacement"
    )
    if active_assignment.handoverCompletedAt is None:
        active_assignment.handoverCompletedAt = now

    if source_unit is not None:
        source_unit.status = "IN_MAINTENANCE"
        source_unit.currentHolderMemberId = None
        source_unit.condition = ticket.conditionBeforeMaintenance or source_unit.condition
        asset.status = _derive_asset_status(asset.units)
    else:
        asset.status = "IN_MAINTENANCE"

    replacement_unit.status = "ASSIGNED"
    replacement_unit.currentHolderMemberId = payload.employeeMemberId
    replacement_unit.condition = "GOOD"
    replacement_asset.status = _derive_asset_status(replacement_asset.units)

    ticket.status = "COMPLETED"
    ticket.replacementDecision = payload.replacementMode
    ticket.replacementAssetUnitId = replacement_unit.id
    if payload.notes and payload.notes.strip():
        note = payload.notes.strip()
        ticket.notes = f"{ticket.notes}\n\n[Replacement] {note}" if ticket.notes else f"[Replacement] {note}"

    await db.commit()

    await create_notification_batch(db, [
        NotificationCreateInput(
            organization_id=ctx.organization.id,
            member_id=payload.employeeMemberId,
            type="REPLACEMENT_PROVIDED",
            category="assets",
            title="Your asset has been replaced",
            message=(
                f"A {payload.replacementMode.lower().replace('_', ' ')} replacement "
                f"({replacement_asset.name}) has been assigned to you. Ticket {ticket.ticketId}."
            ),
            action_url=f"/{ctx.organization.slug}/assets",
            entity_type="REPLACEMENT",
            entity_id=ticket.id,
            metadata={"replacementMode": payload.replacementMode},
        )
    ])

    employee_user = employee_member.user
    if employee_user and employee_user.email:
        await send_asset_replacement_notification(
            to_email=employee_user.email,
            recipient_name=employee_user.name or employee_user.email,
            org_slug=ctx.organization.slug,
            ticket_id=ticket.ticketId,
            replacement_mode=payload.replacementMode,
            new_asset_name=replacement_asset.name,
            new_asset_code=replacement_asset.assetCode,
            admin_name=ctx.member.user.name if ctx.member.user else "Admin",
        )

    return ReplacementRecord(
        id=active_assignment.id,
        employeeMemberId=payload.employeeMemberId,
        employeeName=employee_user.name if employee_user else None,
        employeeEmail=employee_user.email if employee_user else None,
        originalAssetId=asset.id,
        originalAssetName=asset.name,
        originalAssetCode=asset.assetCode,
        originalSerial=asset.serialNumber,
        originalUnitStatus=source_unit.status if source_unit else asset.status,
        replacementAssignmentId=new_assignment.id,
        replacementAssetId=replacement_asset.id,
        replacementAssetName=replacement_asset.name,
        replacementAssetCode=replacement_asset.assetCode,
        replacementSerial=replacement_asset.serialNumber,
        replacementMode=payload.replacementMode,
        expectedReturnDate=(
            new_assignment.expectedReturnDate.isoformat()
            if new_assignment.expectedReturnDate
            else None
        ),
        returnReminderSent=False,
        providedByMemberId=ctx.member.id,
        providedByName=ctx.member.user.name if ctx.member.user else None,
        ticketId=ticket.ticketId,
        maintenanceType=ticket.maintenanceType,
        maintenanceStatus=ticket.status,
        issueDescription=ticket.issueDescription,
        returnDate=active_assignment.returnDate.isoformat() if active_assignment.returnDate else "",
        swapCompletedAt=datetime.now(UTC).isoformat(),
    )


async def raise_replacement_appraisal(
    db: AsyncSession,
    ctx: MemberContext,
    payload: ReplacementRaiseAppraisalRequest,
) -> dict:
    from app.modules.procurement.schema import AssetPurchaseRequisitionCreateRequest

    employee = await db.execute(
        select(Member)
        .where(Member.id == payload.employeeMemberId, Member.organizationId == ctx.organization.id)
    )
    employee_member = employee.scalar_one_or_none()
    if employee_member is None:
        raise HTTPException(status_code=404, detail="Employee not found")

    ticket_result = await db.execute(
        select(AssetMaintenanceLog)
        .where(
            AssetMaintenanceLog.ticketId == payload.ticketId,
            AssetMaintenanceLog.organizationId == ctx.organization.id,
            AssetMaintenanceLog.loggedByMemberId == payload.employeeMemberId,
        )
        .options(joinedload(AssetMaintenanceLog.asset))
    )
    ticket = ticket_result.unique().scalar_one_or_none()
    if ticket is None:
        raise HTTPException(status_code=404, detail="Ticket not found for this employee")

    proc_payload = AssetPurchaseRequisitionCreateRequest(
        requestType="REPLACEMENT",
        maintenanceTicketId=ticket.id,
        justification=f"Replacement requested for ticket {ticket.ticketId}: {ticket.issueDescription[:200]}",
        estimatedQuantity=1,
        estimatedUnitCost=0,
        estimatedTotalCost=0,
        urgency="STANDARD",
        assetName=ticket.asset.name if ticket.asset else None,
        assetCode=ticket.asset.assetCode if ticket.asset else None,
        replacementReason=payload.replacementMode,
        notes=payload.notes,
    )
    result = await _create_proc_req(db, ctx, proc_payload)
    return {"requisitionId": result.id, "message": "Appraisal raised to finance manager for approval"}


async def set_replacement_return_date(
    db: AsyncSession,
    ctx: MemberContext,
    assignment_id: str,
    payload: SetReturnDateRequest,
) -> dict:
    result = await db.execute(
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            AssetAssignment.id == assignment_id,
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
            AssetAssignment.returnDate.is_(None),
        )
    )
    assignment = result.unique().scalar_one_or_none()
    if assignment is None:
        raise HTTPException(status_code=404, detail="Active replacement assignment not found")

    try:
        assignment.expectedReturnDate = datetime.fromisoformat(payload.expectedReturnDate)
    except (ValueError, TypeError):
        raise HTTPException(status_code=422, detail="Invalid date format")

    assignment.returnReminderSent = False
    await db.commit()
    return {"success": True}


async def run_temp_replacement_reminder_scan(db: AsyncSession) -> dict[str, int]:
    tomorrow = date.today() + timedelta(days=1)
    notified = 0

    result = await db.execute(
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.deletedAt.is_(None),
            AssetAssignment.returnDate.is_(None),
            AssetAssignment.expectedReturnDate.is_not(None),
            AssetAssignment.returnReminderSent.is_(False),
        )
        .options(
            joinedload(AssetAssignment.asset),
            joinedload(AssetAssignment.member).joinedload(Member.user),
        )
    )
    assignments = result.unique().scalars().all()

    for assignment in assignments:
        expected = assignment.expectedReturnDate
        if expected is None:
            continue
        if expected.date() != tomorrow:
            continue

        member = assignment.member
        user = member.user if member else None
        if not user or not user.email:
            continue

        org_result = await db.execute(
            select(Organization).where(Organization.id == member.organizationId)
        )
        org = org_result.scalar_one_or_none()
        org_slug = org.slug if org else "app"

        await create_notification_batch(db, [
            NotificationCreateInput(
                organization_id=member.organizationId,
                member_id=member.id,
                type="TEMP_REPLACEMENT_DUE_REMINDER",
                category="assets",
                title="Your temporary replacement is due tomorrow",
                message=(
                    f"The temporary asset {assignment.asset.name if assignment.asset else ''} "
                    f"assigned to you is due for return by {expected.strftime('%b %d, %Y')}."
                ),
                action_url=f"/{org_slug}/assets",
                entity_type="REPLACEMENT",
                entity_id=assignment.id,
            )
        ])

        await send_temp_replacement_reminder(
            to_email=user.email,
            recipient_name=user.name or user.email,
            org_slug=org_slug,
            asset_name=assignment.asset.name if assignment.asset else "Asset",
            expected_return_date=expected.strftime("%b %d, %Y"),
        )

        assignment.returnReminderSent = True
        notified += 1

    if notified:
        await db.commit()
    return {"notified": notified}


# ── Maintenance ────────────────────────────────────────────────────────────────


async def create_maintenance_record(
    db: AsyncSession,
    ctx: MemberContext,
    asset_id: str,
    payload: AssetMaintenanceCreateRequest,
) -> AssetDetailResponse:
    asset = await _get_asset_or_404(db, ctx.organization.id, asset_id)

    if payload.assetUnitId:
        unit = next((u for u in (asset.units or []) if u.id == payload.assetUnitId), None)
        if unit:
            unit.status = "IN_MAINTENANCE"

    db.add(
        AssetMaintenanceLog(
            ticketId=await _generate_ticket_id(db),
            organizationId=ctx.organization.id,
            ticketMode="ASSET_ISSUE",
            assetId=asset.id,
            assetUnitId=payload.assetUnitId,
            loggedByMemberId=ctx.member.id,
            category=payload.category.strip() if payload.category else None,
            subject=payload.subject.strip() if payload.subject else None,
            attachmentsMetadata=payload.attachmentsMetadata,
            maintenanceType=payload.maintenanceType,
            issueDescription=payload.issueDescription.strip(),
            serviceDate=payload.serviceDate,
            expectedCompletionDate=payload.expectedCompletionDate,
            estimatedDowntimeHours=payload.estimatedDowntimeHours,
            operationalCriticalityTier=payload.operationalCriticalityTier,
            cost=payload.cost,
            status=payload.status,
            conditionBeforeMaintenance=payload.conditionBeforeMaintenance or asset.condition,
            notes=payload.notes.strip() if payload.notes else None,
        )
    )

    asset.status = "IN_MAINTENANCE"

    await db.commit()
    return await get_asset(db, ctx, asset.id)


async def create_helpdesk_ticket(
    db: AsyncSession,
    ctx: MemberContext,
    payload: HelpdeskTicketCreateRequest,
) -> MyTicketResponse:
    asset: Asset | None = None
    if payload.assetId is not None:
        asset = await _get_asset_or_404(db, ctx.organization.id, payload.assetId)
        if payload.assetUnitId:
            unit = next((u for u in (asset.units or []) if u.id == payload.assetUnitId), None)
            if unit is None:
                raise HTTPException(status_code=404, detail="Asset unit not found")

    ticket = AssetMaintenanceLog(
        ticketId=await _generate_ticket_id(db),
        organizationId=ctx.organization.id,
        ticketMode=payload.ticketMode,
        assetId=asset.id if asset else None,
        assetUnitId=payload.assetUnitId if asset else None,
        loggedByMemberId=ctx.member.id,
        category=payload.category.strip() if payload.category else None,
        subject=payload.subject.strip(),
        attachmentsMetadata=payload.attachmentsMetadata,
        maintenanceType=payload.maintenanceType,
        issueDescription=payload.issueDescription.strip(),
        serviceDate=payload.serviceDate or date.today(),
        expectedCompletionDate=payload.expectedCompletionDate,
        estimatedDowntimeHours=payload.estimatedDowntimeHours,
        operationalCriticalityTier=payload.operationalCriticalityTier,
        status="OPEN",
        conditionBeforeMaintenance=(
            payload.conditionBeforeMaintenance or (asset.condition if asset is not None else None)
        ),
        notes=payload.notes.strip() if payload.notes else None,
    )
    db.add(ticket)

    await db.commit()

    result = await db.execute(
        select(AssetMaintenanceLog)
        .where(AssetMaintenanceLog.id == ticket.id)
        .options(joinedload(AssetMaintenanceLog.asset))
    )
    created = result.unique().scalar_one()
    await _notify_general_helpdesk_admins(db, ctx, created)
    await db.commit()
    return _my_ticket_response(created)


async def withdraw_helpdesk_ticket(
    db: AsyncSession,
    ctx: MemberContext,
    ticket_id: str,
) -> MyTicketResponse:
    result = await db.execute(
        select(AssetMaintenanceLog)
        .where(
            AssetMaintenanceLog.id == ticket_id,
            AssetMaintenanceLog.organizationId == ctx.organization.id,
            AssetMaintenanceLog.loggedByMemberId == ctx.member.id,
        )
        .options(
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.units),
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.provisions),
        )
    )
    log = result.unique().scalar_one_or_none()
    if log is None:
        raise HTTPException(status_code=404, detail="Ticket not found")
    if log.status == "CANCELLED":
        raise HTTPException(status_code=409, detail="Ticket is already withdrawn")
    if log.status == "COMPLETED":
        raise HTTPException(status_code=409, detail="Completed tickets cannot be withdrawn")

    _mark_ticket_cancelled(log, ctx)
    log.status = "CANCELLED"

    if log.asset is not None:
        active_ticket_result = await db.execute(
            select(AssetMaintenanceLog.id)
            .where(
                AssetMaintenanceLog.organizationId == ctx.organization.id,
                AssetMaintenanceLog.id != log.id,
                AssetMaintenanceLog.status.in_(("OPEN", "IN_PROGRESS")),
                (
                    AssetMaintenanceLog.assetUnitId == log.assetUnitId
                    if log.assetUnitId
                    else AssetMaintenanceLog.assetId == log.assetId
                ),
            )
            .limit(1)
        )
        has_other_active_ticket = active_ticket_result.scalar_one_or_none() is not None

        if not has_other_active_ticket:
            asset = log.asset
            active_assignment = next(
                (record for record in asset.provisions if record.returnDate is None),
                None,
            )
            resumed_status = "ASSIGNED" if active_assignment is not None else "AVAILABLE"

            if log.assetUnitId:
                unit = next((u for u in (asset.units or []) if u.id == log.assetUnitId), None)
                if unit is not None:
                    unit.status = resumed_status
                asset.status = _derive_asset_status(asset.units) if asset.units else resumed_status
            else:
                asset.status = resumed_status

    await db.commit()
    refreshed_result = await db.execute(
        select(AssetMaintenanceLog)
        .where(
            AssetMaintenanceLog.id == log.id,
            AssetMaintenanceLog.organizationId == ctx.organization.id,
            AssetMaintenanceLog.loggedByMemberId == ctx.member.id,
        )
        .options(joinedload(AssetMaintenanceLog.asset))
    )
    refreshed_log = refreshed_result.unique().scalar_one()
    return _my_ticket_response(refreshed_log)


async def update_maintenance_record(
    db: AsyncSession,
    ctx: MemberContext,
    asset_id: str,
    maintenance_id: str,
    payload: AssetMaintenanceUpdateRequest,
) -> AssetDetailResponse:
    asset = await _get_asset_or_404(db, ctx.organization.id, asset_id)
    log = await _get_maintenance_or_404(db, asset_id, maintenance_id)
    previous_status = log.status

    log.status = payload.status
    log.expectedCompletionDate = payload.expectedCompletionDate or log.expectedCompletionDate
    log.completedDate = payload.completedDate
    log.conditionAfterMaintenance = payload.conditionAfterMaintenance
    if payload.notes:
        log.notes = payload.notes.strip()

    if payload.status in {"OPEN", "IN_PROGRESS"}:
        if asset.units and len(asset.units) > 0:
            asset.status = _derive_asset_status(asset.units)
        else:
            asset.status = "IN_MAINTENANCE"
    elif payload.status == "COMPLETED":
        if payload.conditionAfterMaintenance:
            asset.condition = payload.conditionAfterMaintenance
        if _is_return_request_ticket(log):
            active_assignment = next(
                (record for record in asset.provisions if record.returnDate is None),
                None,
            )
            if active_assignment is not None:
                resolved_return_date = (
                    datetime.combine(payload.completedDate, datetime.min.time(), tzinfo=UTC)
                    if payload.completedDate
                    else datetime.now(UTC)
                )
                active_assignment.returnDate = resolved_return_date
                active_assignment.returnedCondition = (
                    payload.conditionAfterMaintenance or active_assignment.conditionWhileProviding
                )
                active_assignment.receivedByMemberId = ctx.member.id
                active_assignment.returnNotes = payload.notes.strip() if payload.notes else "Returned via maintenance workflow"
                if active_assignment.handoverRequestedAt and active_assignment.handoverCompletedAt is None:
                    active_assignment.handoverCompletedAt = resolved_return_date

                resolved_unit_id = log.assetUnitId or active_assignment.assetUnitId
                if resolved_unit_id:
                    unit = next((u for u in (asset.units or []) if u.id == resolved_unit_id), None)
                    if unit:
                        unit.status = payload.nextAssetStatus or "AVAILABLE"
                        unit.condition = payload.conditionAfterMaintenance or unit.condition
                        unit.currentHolderMemberId = None
                    asset.status = _derive_asset_status(asset.units) if asset.units else (payload.nextAssetStatus or "AVAILABLE")
                else:
                    asset.status = payload.nextAssetStatus or "AVAILABLE"
            await db.commit()
            await _notify_general_helpdesk_requester_if_completed(db, ctx, log, previous_status)
            await db.commit()
            return await get_asset(db, ctx, asset.id)
        if payload.nextAssetStatus:
            if log.assetUnitId:
                unit = next((u for u in (asset.units or []) if u.id == log.assetUnitId), None)
                if unit:
                    unit.status = payload.nextAssetStatus
                    unit.condition = payload.conditionAfterMaintenance or unit.condition
                if asset.units and len(asset.units) > 0:
                    asset.status = _derive_asset_status(asset.units)
                else:
                    asset.status = payload.nextAssetStatus
            else:
                asset.status = payload.nextAssetStatus
    elif payload.status == "CANCELLED":
        _mark_ticket_cancelled(log, ctx)
        if log.assetUnitId:
            unit = next((u for u in (asset.units or []) if u.id == log.assetUnitId), None)
            if unit:
                unit.status = payload.nextAssetStatus or "DAMAGED"
            if asset.units and len(asset.units) > 0:
                asset.status = _derive_asset_status(asset.units)
            else:
                asset.status = payload.nextAssetStatus or "DAMAGED"
        else:
            asset.status = payload.nextAssetStatus or "DAMAGED"

    await db.commit()
    await _notify_general_helpdesk_requester_if_completed(db, ctx, log, previous_status)
    await db.commit()
    return await get_asset(db, ctx, asset.id)


async def update_maintenance_record_by_id(
    db: AsyncSession,
    ctx: MemberContext,
    maintenance_id: str,
    payload: AssetMaintenanceUpdateRequest,
) -> None:
    result = await db.execute(
        select(AssetMaintenanceLog)
        .where(
            AssetMaintenanceLog.id == maintenance_id,
            AssetMaintenanceLog.organizationId == ctx.organization.id,
        )
        .options(
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.units),
        )
    )
    log = result.unique().scalar_one_or_none()
    if log is None:
        raise HTTPException(status_code=404, detail="Maintenance record not found")

    if log.assetId:
        await update_maintenance_record(db, ctx, log.assetId, maintenance_id, payload)
        return

    previous_status = log.status
    log.status = payload.status
    log.expectedCompletionDate = payload.expectedCompletionDate or log.expectedCompletionDate
    log.completedDate = payload.completedDate
    log.conditionAfterMaintenance = payload.conditionAfterMaintenance
    if payload.notes:
        log.notes = payload.notes.strip()

    if payload.status == "CANCELLED":
        _mark_ticket_cancelled(log, ctx)

    await db.commit()
    await _notify_general_helpdesk_requester_if_completed(db, ctx, log, previous_status)
    await db.commit()


# ── Meta ───────────────────────────────────────────────────────────────────────


async def get_asset_meta(db: AsyncSession, ctx: MemberContext) -> AssetMetaResponse:
    await _ensure_laptop_os_field(db, ctx.organization.id)
    scope = getattr(ctx, "scope", "organization")
    members_query = (
        select(Member.id, User.name, User.email)
        .join(User, User.id == Member.userId)
        .where(Member.organizationId == ctx.organization.id)
        .order_by(User.name.asc().nullslast(), User.email.asc())
    )
    if scope == "self":
        members_query = members_query.where(Member.id == ctx.member.id)
    result = await db.execute(members_query)
    members = [
        AssetLookupOption(id=member_id, label=name or email or member_id, email=email)
        for member_id, name, email in result.all()
    ]

    categories_result = await db.execute(
        select(AssetCategoryDefinition)
        .where(
            AssetCategoryDefinition.organizationId == ctx.organization.id,
            AssetCategoryDefinition.isActive.is_(True),
        )
        .options(joinedload(AssetCategoryDefinition.fields))
        .order_by(AssetCategoryDefinition.name.asc())
    )
    categories_raw = categories_result.unique().scalars().all()

    categories = [
        AssetCategoryResponse(
            id=c.id,
            name=c.name,
            assetCode=c.assetCode,
            description=c.description,
            isActive=c.isActive,
            fields=[
                CategoryFieldDefinitionResponse(
                    id=f.id,
                    categoryId=f.categoryId,
                    fieldName=f.fieldName,
                    fieldType=f.fieldType,
                    fieldOptions=f.fieldOptions,
                    isRequired=f.isRequired,
                    displayOrder=f.displayOrder,
                )
                for f in (c.fields or [])
            ],
        )
        for c in categories_raw
    ]

    return AssetMetaResponse(
        members=members,
        categories=categories,
        statuses=sorted(ASSET_STATUSES),
        conditions=sorted(ASSET_CONDITIONS),
        maintenanceTypes=sorted(MAINTENANCE_TYPES),
        maintenanceStatuses=sorted(MAINTENANCE_STATUSES),
        ticketModes=sorted(TICKET_MODES),
        reportTypes=sorted(REPORT_TYPES),
    )


# ── Reports ────────────────────────────────────────────────────────────────────


async def export_asset_report(
    db: AsyncSession,
    ctx: MemberContext,
    request: AssetReportRequest,
) -> str:
    report_type = request.reportType

    if report_type in {
        "ALL_ASSETS",
        "AVAILABLE_ASSETS",
        "PROVIDED_ASSETS",
        "DAMAGED_ASSETS",
        "OFFBOARDING_PENDING_RETURN",
    }:
        status_filter = None
        if report_type == "AVAILABLE_ASSETS":
            status_filter = "AVAILABLE"
        elif report_type == "PROVIDED_ASSETS":
            status_filter = "ASSIGNED"
        elif report_type == "DAMAGED_ASSETS":
            status_filter = "DAMAGED"
        elif report_type == "OFFBOARDING_PENDING_RETURN":
            status_filter = "PENDING_RETURN"

        dataset = await list_assets(
            db,
            ctx,
            AssetFilters(
                status=status_filter,
                currentHolderMemberId=request.memberId,
                page=1,
                page_size=500,
            ),
        )
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(
            [
                "Asset Code",
                "Asset Name",
                "Category",
                "Status",
                "Condition",
                "Holder",
                "Location",
                "Units Total",
                "Units Available",
                "Units Provided",
            ]
        )
        for item in dataset.items:
            us = item.unitSummary
            writer.writerow(
                [
                    item.assetCode,
                    item.name,
                    item.category,
                    item.status,
                    item.condition,
                    item.currentHolderName or "",
                    item.location or "",
                    us.total if us else "",
                    us.available if us else "",
                    us.provided if us else "",
                ]
            )
        return buffer.getvalue()

    if report_type in {"RETURNED_ASSETS", "EMPLOYEE_ASSET_REPORT"}:
        query = (
            select(AssetAssignment)
            .join(Asset, Asset.id == AssetAssignment.assetId)
            .where(Asset.organizationId == ctx.organization.id, Asset.deletedAt.is_(None))
            .options(
                joinedload(AssetAssignment.asset),
                joinedload(AssetAssignment.member).joinedload(Member.user),
                joinedload(AssetAssignment.providedByMember).joinedload(Member.user),
                joinedload(AssetAssignment.receivedByMember).joinedload(Member.user),
            )
            .order_by(AssetAssignment.providedDate.desc())
        )
        if report_type == "RETURNED_ASSETS":
            query = query.where(AssetAssignment.returnDate.is_not(None))
        if request.memberId:
            query = query.where(AssetAssignment.memberId == request.memberId)
        rows = (await db.execute(query)).unique().scalars().all()
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(
            [
                "Asset Code",
                "Asset Name",
                "Employee",
                "Provided Date",
                "Return Date",
                "Returned Condition",
            ]
        )
        for row in rows:
            writer.writerow(
                [
                    row.asset.assetCode if row.asset else "",
                    row.asset.name if row.asset else "",
                    row.member.user.name if row.member and row.member.user else "",
                    row.providedDate.isoformat(),
                    row.returnDate.isoformat() if row.returnDate else "",
                    row.returnedCondition or "",
                ]
            )
        return buffer.getvalue()

    if report_type == "MAINTENANCE_HISTORY":
        rows = (
            (
                await db.execute(
                    select(AssetMaintenanceLog)
                    .outerjoin(Asset, Asset.id == AssetMaintenanceLog.assetId)
                    .where(
                        AssetMaintenanceLog.organizationId == ctx.organization.id,
                        or_(Asset.id.is_(None), Asset.deletedAt.is_(None)),
                    )
                    .options(joinedload(AssetMaintenanceLog.asset))
                    .order_by(AssetMaintenanceLog.createdAt.desc())
                )
            )
            .unique()
            .scalars()
            .all()
        )
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(
            [
                "Ticket ID",
                "Asset Code",
                "Asset Name",
                "Maintenance Type",
                "Status",
                "Service Date",
                "Completed Date",
                "Cost",
                "Issue Description",
            ]
        )
        for row in rows:
            writer.writerow(
                [
                    row.ticketId,
                    row.asset.assetCode if row.asset else "",
                    row.asset.name if row.asset else "",
                    row.maintenanceType,
                    row.status,
                    row.serviceDate.isoformat(),
                    row.completedDate.isoformat() if row.completedDate else "",
                    f"{_to_float(row.cost):.2f}" if row.cost is not None else "",
                    row.issueDescription,
                ]
            )
        return buffer.getvalue()

    raise HTTPException(status_code=422, detail="Unsupported report type")


async def export_asset_report_pdf(
    db: AsyncSession,
    ctx: MemberContext,
    request: AssetReportRequest,
) -> bytes:
    csv_text = await export_asset_report(db, ctx, request)
    return _csv_to_pdf_bytes(request.reportType, csv_text)


async def export_asset_report_xlsx(
    db: AsyncSession,
    ctx: MemberContext,
    request: AssetReportRequest,
) -> bytes:
    csv_text = await export_asset_report(db, ctx, request)
    return _csv_to_xlsx_bytes(request.reportType, csv_text)


# ── Dashboard ─────────────────────────────────────────────────────────────────


STATUS_COLORS: dict[str, str] = {
    "AVAILABLE": "#00874a",
    "ASSIGNED": "#2563eb",
    "IN_MAINTENANCE": "#d97706",
    "PENDING_RETURN": "#9a6700",
    "DAMAGED": "#dc2626",
    "LOST": "#7c3aed",
    "RETIRED": "#6b7280",
    "DISPOSED": "#9ca3af",
}


async def get_dashboard(db: AsyncSession, ctx: MemberContext) -> AssetDashboardResponse:
    org_id = ctx.organization.id

    total = await db.scalar(
        select(func.count()).where(
            Asset.organizationId == org_id,
            Asset.deletedAt.is_(None),
        )
    )
    total = total or 0

    status_rows = await db.execute(
        select(Asset.status, func.count().label("cnt"))
        .where(Asset.organizationId == org_id, Asset.deletedAt.is_(None))
        .group_by(Asset.status)
    )
    status_map: dict[str, int] = {}
    for raw_status, count in status_rows.all():
        normalized_status = _normalize_asset_status(raw_status) or raw_status
        status_map[normalized_status] = status_map.get(normalized_status, 0) + count

    def _c(name: str) -> int:
        return status_map.get(name, 0)

    raw_statuses = sorted(status_map.keys(), key=lambda s: -status_map[s])

    status_distribution = [
        AssetStatusCount(name=s, value=status_map[s], color=STATUS_COLORS.get(s, "#6b7280"))
        for s in raw_statuses
    ]

    monthly_rows = await db.execute(
        select(
            extract("year", Asset.createdAt).label("year"),
            extract("month", Asset.createdAt).label("month"),
            func.count().label("cnt"),
        )
        .where(Asset.organizationId == org_id, Asset.deletedAt.is_(None))
        .group_by(
            extract("year", Asset.createdAt),
            extract("month", Asset.createdAt),
        )
        .order_by(
            extract("year", Asset.createdAt).desc(),
            extract("month", Asset.createdAt).desc(),
        )
        .limit(12)
    )
    monthly_trends = [
        MonthlyTrend(month=f"{row.year}-{int(row.month):02d}", count=row.cnt)
        for row in monthly_rows
    ]
    monthly_trends.reverse()

    recent_activity: list[RecentActivityItem] = []

    provisions = await db.execute(
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.organizationId == org_id,
            Asset.deletedAt.is_(None),
        )
        .options(
            joinedload(AssetAssignment.asset),
            joinedload(AssetAssignment.member).joinedload(Member.user),
        )
        .order_by(AssetAssignment.providedDate.desc())
        .limit(5)
    )
    for p in provisions.unique().scalars().all():
        recent_activity.append(
            RecentActivityItem(
                type="ASSIGNED",
                assetName=p.asset.name if p.asset else "",
                memberName=p.member.user.name if p.member and p.member.user else None,
                date=p.providedDate.isoformat() if p.providedDate else "",
                detail=p.provideNotes.strip() if p.provideNotes else None,
            )
        )

    returns = await db.execute(
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.organizationId == org_id,
            Asset.deletedAt.is_(None),
            AssetAssignment.returnDate.is_not(None),
        )
        .options(
            joinedload(AssetAssignment.asset),
            joinedload(AssetAssignment.member).joinedload(Member.user),
        )
        .order_by(AssetAssignment.returnDate.desc())
        .limit(5)
    )
    for r in returns.unique().scalars().all():
        recent_activity.append(
            RecentActivityItem(
                type="RETURNED",
                assetName=r.asset.name if r.asset else "",
                memberName=r.member.user.name if r.member and r.member.user else None,
                date=r.returnDate.isoformat() if r.returnDate else "",
                detail=r.returnNotes.strip() if r.returnNotes else None,
            )
        )

    maintenance = await db.execute(
        select(AssetMaintenanceLog)
        .outerjoin(Asset, Asset.id == AssetMaintenanceLog.assetId)
        .where(
            AssetMaintenanceLog.organizationId == org_id,
            or_(Asset.id.is_(None), Asset.deletedAt.is_(None)),
        )
        .options(
            joinedload(AssetMaintenanceLog.asset),
        )
        .order_by(AssetMaintenanceLog.createdAt.desc())
        .limit(5)
    )
    for m in maintenance.unique().scalars().all():
        recent_activity.append(
            RecentActivityItem(
                type="MAINTENANCE",
                assetName=m.asset.name if m.asset else (m.subject or "General help request"),
                memberName=None,
                date=m.createdAt.isoformat() if m.createdAt else "",
                detail=m.issueDescription.strip() if m.issueDescription else None,
            )
        )

    recent_activity.sort(key=lambda x: x.date, reverse=True)
    recent_activity = recent_activity[:8]

    open_tickets_result = await db.execute(
        select(AssetMaintenanceLog)
        .outerjoin(Asset, Asset.id == AssetMaintenanceLog.assetId)
        .where(
            AssetMaintenanceLog.organizationId == org_id,
            or_(Asset.id.is_(None), Asset.deletedAt.is_(None)),
            AssetMaintenanceLog.status.in_(["OPEN", "IN_PROGRESS"]),
        )
        .options(
            joinedload(AssetMaintenanceLog.asset),
        )
        .order_by(AssetMaintenanceLog.createdAt.desc())
        .limit(5)
    )
    open_tickets = open_tickets_result.unique().scalars().all()

    recent_tickets = [
        TicketAlertItem(
            id=t.id,
            ticketId=t.ticketId,
            ticketMode=t.ticketMode,
            assetName=t.asset.name if t.asset else None,
            category=t.category,
            subject=t.subject,
            maintenanceType=t.maintenanceType,
            status=t.status,
            issueDescription=t.issueDescription.strip() if t.issueDescription else "",
            createdAt=t.createdAt.isoformat() if t.createdAt else "",
        )
        for t in open_tickets
    ]

    return AssetDashboardResponse(
        totalAssets=total,
        availableCount=_c("AVAILABLE"),
        providedCount=_c("ASSIGNED"),
        maintenanceCount=_c("IN_MAINTENANCE") + _c("PENDING_RETURN"),
        damagedCount=_c("DAMAGED"),
        retiredCount=_c("RETIRED"),
        openTicketCount=len(open_tickets),
        statusDistribution=status_distribution,
        monthlyTrends=monthly_trends,
        recentActivity=recent_activity,
        recentTickets=recent_tickets,
    )


async def get_brand_model_analytics(
    db: AsyncSession, ctx: MemberContext
) -> AssetBrandModelAnalyticsResponse:
    result = await db.execute(
        select(Asset)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
        )
        .options(
            joinedload(Asset.units),
            joinedload(Asset.customFieldValues).joinedload(AssetCustomFieldValue.fieldDefinition),
        )
        .order_by(Asset.name.asc(), Asset.model.asc(), Asset.createdAt.asc())
    )
    assets = result.unique().scalars().all()

    grouped_rows: dict[str, dict] = {}

    for asset in assets:
        if not _is_laptop_asset(asset):
            continue

        brand = _resolve_asset_brand(asset)
        model = _resolve_asset_model(asset, brand)
        row_key = hashlib.md5(f"{brand}|{model}".encode()).hexdigest()[:16]
        row = grouped_rows.setdefault(
            row_key,
            {
                "rowKey": row_key,
                "brand": brand,
                "model": model,
                "totalStock": 0,
                "inOfficeStock": 0,
                "providedStock": 0,
                "maintenanceOrDamagedStock": 0,
                "temporaryLaptopStockDepth": 0,
                "lowStockAlert": False,
                "assetIds": set(),
                "unitIds": set(),
                "serialNumbers": set(),
            },
        )

        inventory_units = asset.units or [None]
        for unit in inventory_units:
            current_status = _normalize_asset_status(
                unit.status if unit is not None else asset.status
            )
            serial_number = (
                _normalized_token(unit.serialNumber) if unit is not None else _normalized_token(asset.serialNumber)
            )

            row["totalStock"] += 1
            row["assetIds"].add(asset.id)
            if unit is not None:
                row["unitIds"].add(unit.id)
            if serial_number:
                row["serialNumbers"].add(serial_number)

            if current_status == "AVAILABLE":
                row["inOfficeStock"] += 1
                if _is_temporary_laptop(asset, brand, model):
                    row["temporaryLaptopStockDepth"] += 1
            elif current_status == "ASSIGNED":
                row["providedStock"] += 1
            elif current_status in {"IN_MAINTENANCE", "PENDING_RETURN", "DAMAGED"}:
                row["maintenanceOrDamagedStock"] += 1

        row["lowStockAlert"] = row["inOfficeStock"] == 0

    rows = [
        AssetBrandModelAnalyticsRow(
            rowKey=item["rowKey"],
            brand=item["brand"],
            model=item["model"],
            totalStock=item["totalStock"],
            inOfficeStock=item["inOfficeStock"],
            providedStock=item["providedStock"],
            maintenanceOrDamagedStock=item["maintenanceOrDamagedStock"],
            temporaryLaptopStockDepth=item["temporaryLaptopStockDepth"],
            lowStockAlert=item["lowStockAlert"],
            assetIds=sorted(item["assetIds"]),
            unitIds=sorted(item["unitIds"]),
            serialNumbers=sorted(item["serialNumbers"]),
        )
        for item in grouped_rows.values()
    ]
    rows.sort(key=lambda item: (item.brand.lower(), item.model.lower()))

    return AssetBrandModelAnalyticsResponse(
        rows=rows,
        temporaryLaptopStockDepth=sum(row.temporaryLaptopStockDepth for row in rows),
    )


async def get_os_distribution_analytics(
    db: AsyncSession, ctx: MemberContext
) -> AssetOsDistributionResponse:
    await _ensure_laptop_os_field(db, ctx.organization.id)

    scope = getattr(ctx, "scope", "organization")
    query = (
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
            AssetAssignment.returnDate.is_(None),
            Asset.status.in_(("ASSIGNED", "PROVIDED")),
        )
        .options(
            joinedload(AssetAssignment.asset)
            .joinedload(Asset.customFieldValues)
            .joinedload(AssetCustomFieldValue.fieldDefinition)
        )
    )
    if scope == "self":
        query = query.where(AssetAssignment.memberId == ctx.member.id)

    result = await db.execute(query)
    assignments = result.unique().scalars().all()

    os_to_members: dict[str, set[str]] = {}
    total_laptop_users: set[str] = set()
    for assignment in assignments:
        asset = assignment.asset
        if asset is None or not _is_laptop_asset(asset):
            continue

        os_value = "Unknown"
        for field_value in asset.customFieldValues or []:
            field_name = _normalize_key(
                field_value.fieldDefinition.fieldName if field_value.fieldDefinition else None
            )
            if field_name == "os":
                os_value = _normalize_os_name(field_value.value)
                break

        total_laptop_users.add(assignment.memberId)
        os_to_members.setdefault(os_value, set()).add(assignment.memberId)

    total_users = len(total_laptop_users)
    rows = [
        AssetOsDistributionRow(
            osName=os_name,
            headcount=len(member_ids),
            percentage=round((len(member_ids) / total_users) * 100, 1) if total_users else 0,
            memberIds=sorted(member_ids),
        )
        for os_name, member_ids in os_to_members.items()
    ]
    rows.sort(key=lambda row: (-row.headcount, row.osName.lower()))

    return AssetOsDistributionResponse(rows=rows, totalLaptopUsers=total_users)


async def get_upcoming_warranty_feed(
    db: AsyncSession, ctx: MemberContext
) -> WarrantyExpirationFeedResponse:
    today = date.today()
    window_end = today + timedelta(days=7)
    scope = getattr(ctx, "scope", "organization")

    query = (
        select(Asset)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
        )
        .options(
            joinedload(Asset.units),
            joinedload(Asset.provisions)
            .joinedload(AssetAssignment.member)
            .joinedload(Member.user),
        )
        .order_by(Asset.name.asc(), Asset.assetCode.asc())
    )
    result = await db.execute(query)
    assets = result.unique().scalars().all()

    items: list[WarrantyExpirationFeedItem] = []
    for asset in assets:
        active_assignment = next(
            (record for record in asset.provisions if record.returnDate is None),
            None,
        )
        if active_assignment is None:
            continue
        if scope == "self" and active_assignment.memberId != ctx.member.id:
            continue

        units = asset.units or [None]
        for unit in units:
            warranty_expiry = unit.warrantyExpiryDate if unit is not None else asset.warrantyExpiryDate
            if warranty_expiry is None or not (today <= warranty_expiry <= window_end):
                continue
            if unit is not None and unit.currentHolderMemberId and unit.currentHolderMemberId != active_assignment.memberId:
                continue

            items.append(
                WarrantyExpirationFeedItem(
                    assetId=asset.id,
                    assetUnitId=unit.id if unit is not None else None,
                    assetCode=asset.assetCode,
                    assetName=asset.name,
                    serialNumber=unit.serialNumber if unit is not None else asset.serialNumber,
                    brand=asset.brand,
                    model=asset.model,
                    category=asset.category,
                    employeeMemberId=active_assignment.memberId,
                    employeeName=active_assignment.member.user.name
                    if active_assignment.member and active_assignment.member.user
                    else None,
                    employeeEmail=active_assignment.member.user.email
                    if active_assignment.member and active_assignment.member.user
                    else None,
                    warrantyExpiryDate=warranty_expiry,
                    daysUntilExpiry=(warranty_expiry - today).days,
                    hasReminderSent=bool(unit.lastWarrantyAlertSentAt) if unit is not None else False,
                )
            )

    items.sort(key=lambda item: (item.daysUntilExpiry, item.warrantyExpiryDate, item.assetName.lower()))
    return WarrantyExpirationFeedResponse(items=items, total=len(items))


async def run_warranty_tracker_scan(db: AsyncSession) -> dict[str, int]:
    today = date.today()
    target_date = today + timedelta(days=7)
    sent_count = 0
    notification_count = 0

    result = await db.execute(
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.deletedAt.is_(None),
            AssetAssignment.returnDate.is_(None),
        )
        .options(
            joinedload(AssetAssignment.asset).joinedload(Asset.units),
            joinedload(AssetAssignment.member).joinedload(Member.user),
            joinedload(AssetAssignment.asset).joinedload(Asset.provisions),
        )
    )
    active_assignments = result.unique().scalars().all()
    organization_cache: dict[str, Organization] = {}
    admin_cache: dict[str, list[Member]] = {}

    for assignment in active_assignments:
        asset = assignment.asset
        if asset is None:
            continue

        tracking_unit = None
        if assignment.assetUnitId:
            tracking_unit = next((unit for unit in asset.units if unit.id == assignment.assetUnitId), None)
        if tracking_unit is None:
            tracking_unit = await _ensure_tracking_unit(db, asset)
            tracking_unit.currentHolderMemberId = assignment.memberId

        tracking_unit.warrantyExpiryDate = tracking_unit.warrantyExpiryDate or asset.warrantyExpiryDate
        if tracking_unit.warrantyExpiryDate != target_date:
            continue
        if tracking_unit.reminderCompleted:
            continue

        org_id = asset.organizationId
        if org_id not in organization_cache:
            organization_cache[org_id] = await db.get(Organization, org_id)
        organization = organization_cache.get(org_id)
        if organization is None:
            continue

        if org_id not in admin_cache:
            admin_cache[org_id] = await _members_with_asset_admin_scope(db, org_id)
        admin_members = admin_cache[org_id]

        assigned_member = assignment.member
        assigned_user = assigned_member.user if assigned_member else None
        holder_name = (
            assigned_user.name
            or assigned_user.email
            or assignment.memberId
            if assigned_user is not None
            else assignment.memberId
        )
        message = (
            f"The warranty for {asset.name} ({asset.assetCode}) expires on "
            f"{target_date.isoformat()}. Prepare a renewal, laptop refresh, or asset swap."
        )

        for admin_member in admin_members:
            admin_user = admin_member.user
            if admin_user is None:
                continue
            db.add(
                AssetNotification(
                    organizationId=org_id,
                    assetId=asset.id,
                    assetUnitId=tracking_unit.id,
                    memberId=admin_member.id,
                    type="WARRANTY_EXPIRING_7_DAYS",
                    title="Warranty expiring in 7 days",
                    message=message,
                )
            )
            notification_count += 1
            if admin_user.email:
                await send_asset_warranty_expiry_alert(
                    to_email=admin_user.email,
                    recipient_name=admin_user.name or admin_user.email,
                    recipient_role="Admin",
                    organization_name=organization.name,
                    org_slug=organization.slug,
                    asset_name=asset.name,
                    asset_code=asset.assetCode,
                    serial_number=tracking_unit.serialNumber or asset.serialNumber,
                    model=asset.model,
                    holder_name=holder_name,
                    warranty_expiry_date=target_date,
                )

        if assigned_member is not None:
            db.add(
                AssetNotification(
                    organizationId=org_id,
                    assetId=asset.id,
                    assetUnitId=tracking_unit.id,
                    memberId=assigned_member.id,
                    type="WARRANTY_EXPIRING_7_DAYS",
                    title="Your assigned hardware warranty expires in 7 days",
                    message=message,
                )
            )
            notification_count += 1
            if assigned_user and assigned_user.email:
                await send_asset_warranty_expiry_alert(
                    to_email=assigned_user.email,
                    recipient_name=assigned_user.name or assigned_user.email,
                    recipient_role="Employee",
                    organization_name=organization.name,
                    org_slug=organization.slug,
                    asset_name=asset.name,
                    asset_code=asset.assetCode,
                    serial_number=tracking_unit.serialNumber or asset.serialNumber,
                    model=asset.model,
                    holder_name=holder_name,
                    warranty_expiry_date=target_date,
                )

        tracking_unit.lastWarrantyAlertSentAt = datetime.now(UTC)
        tracking_unit.reminderCompleted = True
        sent_count += 1

    await db.commit()
    return {
        "matched_assets": len(active_assignments),
        "alerts_sent": sent_count,
        "notifications_created": notification_count,
    }


async def list_my_tickets(db: AsyncSession, ctx: MemberContext) -> list[MyTicketResponse]:
    result = await db.execute(
        select(AssetMaintenanceLog)
        .outerjoin(Asset, Asset.id == AssetMaintenanceLog.assetId)
        .where(
            AssetMaintenanceLog.organizationId == ctx.organization.id,
            or_(Asset.id.is_(None), Asset.deletedAt.is_(None)),
            AssetMaintenanceLog.loggedByMemberId == ctx.member.id,
        )
        .options(joinedload(AssetMaintenanceLog.asset))
        .order_by(AssetMaintenanceLog.createdAt.desc())
    )
    logs = result.unique().scalars().all()
    return [_my_ticket_response(log) for log in logs]


async def list_tickets(db: AsyncSession, ctx: MemberContext) -> list[MaintenanceTicketResponse]:
    result = await db.execute(
        select(AssetMaintenanceLog)
        .outerjoin(Asset, Asset.id == AssetMaintenanceLog.assetId)
        .where(
            AssetMaintenanceLog.organizationId == ctx.organization.id,
            or_(Asset.id.is_(None), Asset.deletedAt.is_(None)),
        )
        .options(
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.provisions).joinedload(
                AssetAssignment.member
            ).joinedload(Member.user),
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.units),
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.customFieldValues).joinedload(
                AssetCustomFieldValue.fieldDefinition
            ),
            joinedload(AssetMaintenanceLog.loggedByMember).joinedload(Member.user),
            joinedload(AssetMaintenanceLog.cancelledByMember).joinedload(Member.user),
        )
        .order_by(AssetMaintenanceLog.createdAt.desc())
    )
    logs = result.unique().scalars().all()
    responses: list[MaintenanceTicketResponse] = []
    for log in logs:
        response = _maintenance_ticket_response(log)
        if log.asset is not None:
            response.swapPreview = await _build_swap_preview(
                db, ctx.organization.id, log, log.asset
            )
        responses.append(response)
    return responses


async def get_swap_preview(
    db: AsyncSession, ctx: MemberContext, maintenance_id: str
) -> AssetSwapPreviewResponse:
    result = await db.execute(
        select(AssetMaintenanceLog)
        .where(
            AssetMaintenanceLog.id == maintenance_id,
            AssetMaintenanceLog.organizationId == ctx.organization.id,
        )
        .options(
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.provisions).joinedload(
                AssetAssignment.member
            ).joinedload(Member.user),
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.units),
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.customFieldValues).joinedload(
                AssetCustomFieldValue.fieldDefinition
            ),
        )
    )
    log = result.unique().scalar_one_or_none()
    if log is None or log.asset is None:
        raise HTTPException(status_code=404, detail="Maintenance record not found")

    return await _build_swap_preview(db, ctx.organization.id, log, log.asset)


async def revoke_and_swap_asset(
    db: AsyncSession,
    ctx: MemberContext,
    maintenance_id: str,
    payload: AssetRevokeSwapRequest,
) -> AssetSwapExecutionResponse:
    if payload.maintenanceId != maintenance_id:
        raise HTTPException(status_code=422, detail="Maintenance ID mismatch")

    result = await db.execute(
        select(AssetMaintenanceLog)
        .where(
            AssetMaintenanceLog.id == maintenance_id,
            AssetMaintenanceLog.organizationId == ctx.organization.id,
        )
        .options(
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.provisions).joinedload(
                AssetAssignment.member
            ).joinedload(Member.user),
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.units),
            joinedload(AssetMaintenanceLog.asset).joinedload(Asset.customFieldValues).joinedload(
                AssetCustomFieldValue.fieldDefinition
            ),
        )
    )
    log = result.unique().scalar_one_or_none()
    if log is None or log.asset is None:
        raise HTTPException(status_code=404, detail="Maintenance record not found")

    asset = log.asset
    preview = await _build_swap_preview(db, ctx.organization.id, log, asset)
    if preview.requiresReplacementValidation and preview.recommendedMode is None:
        raise HTTPException(
            status_code=422,
            detail="No replacement inventory is currently available for this incident",
        )

    active_assignment = next((record for record in asset.provisions if record.returnDate is None), None)
    if active_assignment is None:
        raise HTTPException(status_code=422, detail="The malfunctioning asset is not currently assigned")

    replacement_result = await db.execute(
        select(AssetUnit)
        .join(Asset, Asset.id == AssetUnit.assetId)
        .where(
            AssetUnit.id == payload.replacementAssetUnitId,
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
        )
        .options(
            joinedload(AssetUnit.asset).joinedload(Asset.units),
            joinedload(AssetUnit.asset).joinedload(Asset.customFieldValues).joinedload(
                AssetCustomFieldValue.fieldDefinition
            ),
        )
    )
    replacement_unit = replacement_result.unique().scalar_one_or_none()
    if replacement_unit is None or replacement_unit.asset is None:
        raise HTTPException(status_code=404, detail="Replacement unit not found")
    if _normalize_asset_status(replacement_unit.status) != "AVAILABLE":
        raise HTTPException(status_code=422, detail="Replacement unit is no longer available")

    selected_option = next(
        (option for option in preview.options if option.mode == payload.replacementMode), None
    )
    if selected_option is None or payload.replacementAssetUnitId not in selected_option.assetUnitIds:
        raise HTTPException(
            status_code=422,
            detail="Replacement unit does not match the validated inventory option",
        )

    provider_id = payload.providedByMemberId or ctx.member.id
    await _get_member_or_404(db, ctx.organization.id, provider_id)

    source_unit = None
    if log.assetUnitId:
        source_unit = next((unit for unit in asset.units if unit.id == log.assetUnitId), None)

    if source_unit is not None and payload.replacementAssetUnitId == source_unit.id:
        raise HTTPException(
            status_code=422,
            detail="Replacement unit cannot be the same as the malfunctioning unit",
        )

    replacement_asset = replacement_unit.asset
    now = datetime.now(UTC)
    new_assignment = AssetAssignment(
        assetId=replacement_asset.id,
        assetUnitId=replacement_unit.id,
        memberId=active_assignment.memberId,
        providedByMemberId=provider_id,
        providedDate=now,
        conditionWhileProviding=payload.replacementConditionWhileProviding,
        provideNotes=payload.notes.strip() if payload.notes else None,
    )
    db.add(new_assignment)
    await db.flush()

    active_assignment.replacementAssignmentId = new_assignment.id
    active_assignment.handoverRequestedAt = now
    if payload.notes:
        active_assignment.handoverConditionNotes = payload.notes.strip()

    if payload.revokeStatus != "PENDING_RETURN":
        active_assignment.returnDate = now
        active_assignment.returnedCondition = log.conditionBeforeMaintenance or asset.condition
        active_assignment.receivedByMemberId = ctx.member.id
        active_assignment.returnNotes = (
            payload.notes.strip() if payload.notes else "Revoked automatically during asset swap"
        )
        if active_assignment.handoverCompletedAt is None:
            active_assignment.handoverCompletedAt = now

    if source_unit is not None:
        source_unit.status = payload.revokeStatus
        source_unit.currentHolderMemberId = (
            active_assignment.memberId if payload.revokeStatus == "PENDING_RETURN" else None
        )
        source_unit.condition = log.conditionBeforeMaintenance or source_unit.condition
        asset.status = _derive_asset_status(asset.units)
    else:
        asset.status = payload.revokeStatus

    replacement_unit.status = "ASSIGNED"
    replacement_unit.currentHolderMemberId = active_assignment.memberId
    replacement_unit.condition = payload.replacementConditionWhileProviding
    replacement_asset.status = _derive_asset_status(replacement_asset.units)

    log.status = "COMPLETED"
    log.replacementDecision = payload.replacementMode
    log.replacementAssetUnitId = replacement_unit.id
    if payload.notes and payload.notes.strip():
        swap_note = payload.notes.strip()
        log.notes = f"{log.notes}\n\n[Swap] {swap_note}" if log.notes else f"[Swap] {swap_note}"

    await db.commit()

    assigned_member_name = None
    if active_assignment.member and active_assignment.member.user:
        assigned_member_name = active_assignment.member.user.name or active_assignment.member.user.email

    return AssetSwapExecutionResponse(
        maintenanceId=maintenance_id,
        revokedAssetId=asset.id,
        revokedAssetUnitId=source_unit.id if source_unit else None,
        revokedStatus=payload.revokeStatus,
        replacementAssetId=replacement_asset.id,
        replacementAssetUnitId=replacement_unit.id,
        replacementMode=payload.replacementMode,
        assignmentId=new_assignment.id,
        assignedMemberId=active_assignment.memberId,
        assignedMemberName=assigned_member_name,
    )
