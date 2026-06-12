"""Asset inline export service — XLSX, PDF, CSV per sub-domain."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import and_, extract, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import contains_eager, joinedload, selectinload

from app.models.asset import Asset
from app.models.asset_assignment import AssetAssignment
from app.models.asset_unit import AssetUnit
from app.models.member import Member
from app.models.user import User
from app.modules.assets.schema import AssetExportRequest
from app.shared.deps.organization_member import MemberContext

_HEADER_FILL = "#1F4E78"
_THIN_BORDER = Side(style="thin", color="D9E2EC")
_GRID_BORDER = Border(left=_THIN_BORDER, right=_THIN_BORDER, top=_THIN_BORDER, bottom=_THIN_BORDER)
_REPORT_HEADERS: list[str] = []


def _date_str(value: date | datetime | None) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return value.isoformat()


def _holder_name(asset: Asset) -> str:
    active = [p for p in (asset.provisions or []) if p.returnDate is None]
    if active and active[0].member and active[0].member.user:
        return active[0].member.user.name
    return ""


# ── XLSX helpers ───────────────────────────────────────────────────────────────


def _write_xlsx(headers: list[str], rows: list[list[str]], sheet_name: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color=_HEADER_FILL.lstrip("#"), end_color=_HEADER_FILL.lstrip("#"), fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = _GRID_BORDER

    body_font = Font(size=10)
    body_alignment = Alignment(vertical="center", wrap_text=False)
    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = body_font
            cell.alignment = body_alignment
            cell.border = _GRID_BORDER

    for col_idx in range(1, len(headers) + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        max_len = max(
            len(str(headers[col_idx - 1])),
            max((len(str(row[col_idx - 1])) for row in rows), default=0),
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ── PDF helpers ────────────────────────────────────────────────────────────────


def _write_pdf(
    title: str,
    headers: list[str],
    rows: list[list[str]],
    col_weights: list[float] | None = None,
    metadata: list[str] | None = None,
) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), topMargin=15 * mm, bottomMargin=15 * mm
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Title"], fontSize=14, spaceAfter=8 * mm)

    elements: list = [Paragraph(title, title_style)]

    if metadata:
        meta_style = ParagraphStyle(
            "Meta",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10,
            textColor=colors.HexColor("#475569"),
            spaceAfter=3 * mm
        )
        for line in metadata:
            elements.append(Paragraph(line, meta_style))
        elements.append(Spacer(1, 4 * mm))

    # Wrap headers and row cells in Paragraphs to support text auto-wrap
    header_style = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        textColor=colors.white,
        alignment=0 # Left aligned
    )
    body_style = ParagraphStyle(
        "TableBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1D1D1F"),
        alignment=0 # Left aligned
    )

    formatted_headers = [Paragraph(h, header_style) for h in headers]
    formatted_rows = []
    for r in rows:
        formatted_rows.append([Paragraph(str(cell) if cell is not None else "", body_style) for cell in r])

    table_data = [formatted_headers] + formatted_rows
    col_count = len(headers)
    page_width = landscape(A4)[0] - 30 * mm

    if col_weights:
        total_weight = sum(col_weights)
        col_widths = [w / total_weight * page_width for w in col_weights]
    else:
        col_widths = [page_width / max(col_count, 1)] * col_count

    style = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(_HEADER_FILL)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9E2EC")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ])

    pdf_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    pdf_table.setStyle(style)
    elements.append(pdf_table)
    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


# ── CSV helpers ────────────────────────────────────────────────────────────────


def _write_csv(headers: list[str], rows: list[list[str]]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return buf.getvalue()


# ── Domain: Register (monthly new assets) ──────────────────────────────────────


async def generate_register_export(
    db: AsyncSession,
    ctx: MemberContext,
    request: AssetExportRequest,
) -> bytes | str:
    start = datetime.fromisoformat(request.startDate)
    end = datetime.fromisoformat(request.endDate).replace(hour=23, minute=59, second=59)
    query = (
        select(Asset)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
            Asset.createdAt >= start,
            Asset.createdAt <= end,
        )
        .options(
            selectinload(Asset.provisions).selectinload(AssetAssignment.member).selectinload(Member.user),
        )
        .order_by(Asset.createdAt.desc())
    )
    if request.employeeIds:
        query = query.where(
            Asset.provisions.any(
                AssetAssignment.memberId.in_(request.employeeIds),
                AssetAssignment.returnDate.is_(None),
            )
        )
    result = await db.execute(query)
    assets = result.unique().scalars().all()

    # Brand column added after Asset Name
    headers = ["Asset Code", "Asset Name", "Brand", "Category", "Status", "Condition", "Holder", "Location", "Created Date"]
    rows = [
        [
            a.assetCode,
            a.name,
            a.brand or "",
            a.category or "",
            a.status,
            a.condition,
            _holder_name(a),
            a.location or "",
            _date_str(a.createdAt),
        ]
        for a in assets
    ]

    if request.format == "csv":
        return _write_csv(headers, rows)
    elif request.format == "pdf":
        month_name = start.strftime("%B %Y")
        metadata = [
            f"Month: {month_name}",
            f"No. of Assets Registered: {len(assets)}",
        ]
        col_weights = [1.0, 1.2, 1.1, 1.2, 1.0, 0.9, 1.5, 1.2, 1.0]
        return _write_pdf("Asset Register Report", headers, rows, col_weights=col_weights, metadata=metadata)
    else:
        return _write_xlsx(headers, rows, "Asset Register")


# ── Domain: Issued (ASSIGNED assets) ───────────────────────────────────────────


async def generate_issued_export(
    db: AsyncSession,
    ctx: MemberContext,
    request: AssetExportRequest,
) -> bytes | str:
    start = datetime.fromisoformat(request.startDate)
    end = datetime.fromisoformat(request.endDate).replace(hour=23, minute=59, second=59)
    query = (
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
            AssetAssignment.returnDate.is_(None),
            AssetAssignment.providedDate >= start,
            AssetAssignment.providedDate <= end,
        )
        .options(
            contains_eager(AssetAssignment.asset),
            joinedload(AssetAssignment.member).joinedload(Member.user),
        )
        .order_by(AssetAssignment.providedDate.desc())
    )
    if request.employeeIds:
        query = query.where(AssetAssignment.memberId.in_(request.employeeIds))
    result = await db.execute(query)
    assignments = result.unique().scalars().all()

    headers = ["Asset Code", "Asset Name", "Category", "Employee", "Provided Date", "Condition", "Notes"]
    rows = [
        [
            a.asset.assetCode if a.asset else "",
            a.asset.name if a.asset else "",
            a.asset.category if a.asset else "",
            a.member.user.name if a.member and a.member.user else "",
            _date_str(a.providedDate),
            a.conditionWhileProviding or "",
            a.provideNotes or "",
        ]
        for a in assignments
    ]

    if request.format == "csv":
        return _write_csv(headers, rows)
    elif request.format == "pdf":
        month_name = start.strftime("%B %Y")
        metadata = [
            f"Month: {month_name}",
            f"No. of Assets Issued: {len(assignments)}",
        ]
        col_weights = [1.0, 1.2, 1.2, 1.5, 1.1, 0.9, 3.0]
        return _write_pdf("Issued Assets Report", headers, rows, col_weights=col_weights, metadata=metadata)
    else:
        return _write_xlsx(headers, rows, "Issued Assets")


# ── Domain: Returned ───────────────────────────────────────────────────────────


async def generate_returned_export(
    db: AsyncSession,
    ctx: MemberContext,
    request: AssetExportRequest,
) -> bytes | str:
    start = datetime.fromisoformat(request.startDate)
    end = datetime.fromisoformat(request.endDate).replace(hour=23, minute=59, second=59)
    query = (
        select(AssetAssignment)
        .join(Asset, Asset.id == AssetAssignment.assetId)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
            AssetAssignment.returnDate.is_not(None),
            AssetAssignment.returnDate >= start,
            AssetAssignment.returnDate <= end,
        )
        .options(
            contains_eager(AssetAssignment.asset),
            joinedload(AssetAssignment.member).joinedload(Member.user),
            joinedload(AssetAssignment.receivedByMember).joinedload(Member.user),
        )
        .order_by(AssetAssignment.returnDate.desc())
    )
    if request.employeeIds:
        query = query.where(AssetAssignment.memberId.in_(request.employeeIds))
    result = await db.execute(query)
    assignments = result.unique().scalars().all()

    headers = [
        "Asset Code", "Asset Name", "Category", "Employee",
        "Provided Date", "Return Date", "Returned Condition", "Return Notes",
    ]
    rows = [
        [
            a.asset.assetCode if a.asset else "",
            a.asset.name if a.asset else "",
            a.asset.category if a.asset else "",
            a.member.user.name if a.member and a.member.user else "",
            _date_str(a.providedDate),
            _date_str(a.returnDate),
            a.returnedCondition or "",
            a.returnNotes or "",
        ]
        for a in assignments
    ]

    if request.format == "csv":
        return _write_csv(headers, rows)
    elif request.format == "pdf":
        month_name = start.strftime("%B %Y")
        metadata = [
            f"Month: {month_name}",
            f"No. of Assets Returned: {len(assignments)}",
        ]
        col_weights = [1.0, 1.2, 1.2, 1.5, 1.1, 1.1, 1.1, 2.5]
        return _write_pdf("Returned Assets Report", headers, rows, col_weights=col_weights, metadata=metadata)
    else:
        return _write_xlsx(headers, rows, "Returned Assets")


# ── Domain: Inventory ──────────────────────────────────────────────────────────


async def generate_inventory_export(
    db: AsyncSession,
    ctx: MemberContext,
    request: AssetExportRequest,
) -> bytes | str:
    query = (
        select(Asset)
        .where(
            Asset.organizationId == ctx.organization.id,
            Asset.deletedAt.is_(None),
        )
        .options(
            joinedload(Asset.units),
        )
        .order_by(Asset.brand, Asset.model)
    )
    result = await db.execute(query)
    assets = result.unique().scalars().all()

    headers = [
        "Brand", "Model", "Asset Code", "Asset Name", "Category",
        "Total Stock", "Available", "Provided", "In Maintenance", "Damaged",
        "Status", "Condition", "Low Stock Alert",
    ]
    rows = []
    for a in assets:
        units = a.units or []
        total = len(units)
        available = sum(1 for u in units if u.status == "AVAILABLE")
        provided = sum(1 for u in units if u.status == "ASSIGNED")
        in_maint = sum(1 for u in units if u.status == "IN_MAINTENANCE")
        damaged = sum(1 for u in units if u.status in ("DAMAGED", "LOST"))
        low_stock = "Yes" if total > 0 and available <= 0 else "No"
        rows.append([
            a.brand or "",
            a.model or "",
            a.assetCode,
            a.name,
            a.category or "",
            str(total),
            str(available),
            str(provided),
            str(in_maint),
            str(damaged),
            a.status,
            a.condition,
            low_stock,
        ])

    if request.format == "csv":
        return _write_csv(headers, rows)
    elif request.format == "pdf":
        metadata = [
            f"No. of Inventory Assets: {len(assets)}",
        ]
        col_weights = [1.1, 1.1, 1.0, 1.2, 1.2, 0.8, 0.8, 0.8, 0.8, 0.8, 1.0, 1.0, 0.9]
        return _write_pdf("Inventory Report", headers, rows, col_weights=col_weights, metadata=metadata)
    else:
        return _write_xlsx(headers, rows, "Inventory")
