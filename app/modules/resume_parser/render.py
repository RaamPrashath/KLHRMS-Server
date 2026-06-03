from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Inches, Pt

from app.modules.resume_parser.schema import TemplateType

MODULE_DIR = Path(__file__).resolve().parent
DOCS_DIR = MODULE_DIR / "docs"

BULLET_CHAR = "-"
FONT_SIZE_NORMAL = Pt(11)
SPACE_0PT = Pt(0)
SPACE_6PT = Pt(6)
SPACE_12PT = Pt(12)
LEFT_MARGIN = Cm(1.91)
RIGHT_MARGIN = Cm(1.91)
TOP_MARGIN = Cm(2.54)
BOTTOM_MARGIN = Cm(2.54)
RIGHT_ALIGN_TAB_POS = Inches(5.9)


def create_resume_docx(resume_data: dict[str, Any], template_type: TemplateType) -> bytes:
    template = _template_path(template_type)
    doc = Document(str(template)) if template.exists() else Document()
    _clear_body(doc)
    if template_type == "ncs":
        _render_ncs(doc, resume_data)
    else:
        _render_kovan_like(doc, resume_data)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _template_path(template_type: TemplateType) -> Path:
    if template_type == "ncs":
        return DOCS_DIR / "ncs_template.docx"
    if template_type == "rsc":
        return DOCS_DIR / "rsc_template.docx"
    return DOCS_DIR / "template.docx"


def _clear_body(doc: Document) -> None:
    for paragraph in doc.paragraphs[:]:
        element = paragraph._element
        element.getparent().remove(element)
    for table in doc.tables[:]:
        element = table._element
        element.getparent().remove(element)


def _set_margins(doc: Document) -> None:
    for section in doc.sections:
        section.left_margin = LEFT_MARGIN
        section.right_margin = RIGHT_MARGIN
        section.top_margin = TOP_MARGIN
        section.bottom_margin = BOTTOM_MARGIN


def _ensure_bullet_style(doc: Document) -> None:
    if "List Bullet" in doc.styles:
        return
    style = doc.styles.add_style("List Bullet", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.paragraph_format.space_before = SPACE_0PT
    style.paragraph_format.space_after = SPACE_0PT


def _paragraph(doc: Document, text: str = "", *, bold: bool = False, italic: bool = False) -> Any:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = FONT_SIZE_NORMAL
    paragraph.paragraph_format.space_before = SPACE_0PT
    paragraph.paragraph_format.space_after = SPACE_0PT
    return paragraph


def _heading(doc: Document, text: str, level: int = 1) -> None:
    if not text:
        return
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(18) if level == 0 else Pt(14 if level == 1 else 12)
    paragraph.paragraph_format.space_before = SPACE_12PT if level else SPACE_6PT
    paragraph.paragraph_format.space_after = SPACE_6PT


def _bullet(doc: Document, text: str) -> None:
    if not text:
        return
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(f"{BULLET_CHAR} {text}")
    run.font.size = FONT_SIZE_NORMAL
    paragraph.paragraph_format.left_indent = Cm(0.64)
    paragraph.paragraph_format.space_before = SPACE_0PT
    paragraph.paragraph_format.space_after = SPACE_0PT


def _right_date(paragraph: Any, dates: str | None) -> None:
    if not dates:
        return
    paragraph.paragraph_format.tab_stops.clear_all()
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        RIGHT_ALIGN_TAB_POS,
        WD_TAB_ALIGNMENT.RIGHT,
    )
    paragraph.add_run("\t")
    run = paragraph.add_run(dates.upper())
    run.bold = True
    run.font.size = FONT_SIZE_NORMAL


def _render_kovan_like(doc: Document, data: dict[str, Any]) -> None:
    _set_margins(doc)
    _ensure_bullet_style(doc)
    personal = data.get("personal_info") if isinstance(data.get("personal_info"), dict) else {}
    name = str(personal.get("name") or "").strip()
    if name:
        _heading(doc, name, level=0)

    summary = _string_list(data.get("summary"))
    if summary:
        _heading(doc, "Professional Summary")
        for point in summary:
            _bullet(doc, point)

    experiences = data.get("experience") if isinstance(data.get("experience"), list) else []
    projects = data.get("projects") if isinstance(data.get("projects"), list) else []
    if experiences or projects:
        _heading(doc, "Professional Experience")
        for exp in experiences:
            if not isinstance(exp, dict):
                continue
            company = str(exp.get("company") or "N/A")
            paragraph = _paragraph(doc, company.upper(), bold=True)
            _right_date(paragraph, str(exp.get("dates") or ""))
            title = str(exp.get("title") or "").strip()
            if title:
                _paragraph(doc, title, italic=True)
            _render_project_details(doc, exp)
            for responsibility in _string_list(exp.get("responsibilities")):
                _bullet(doc, responsibility)
        for project in projects:
            if isinstance(project, dict):
                _render_project_details(doc, project)
                for responsibility in _string_list(project.get("responsibilities")):
                    _bullet(doc, responsibility)

    education = data.get("education") if isinstance(data.get("education"), list) else []
    if education:
        _heading(doc, "Education")
        for item in education:
            if not isinstance(item, dict):
                continue
            degree = str(item.get("degree") or "").strip()
            school = str(item.get("school") or "").strip()
            dates = str(item.get("dates") or "").strip()
            paragraph = _paragraph(doc, degree.upper(), bold=True)
            _right_date(paragraph, dates)
            if school:
                _paragraph(doc, school)

    awards = _string_list(data.get("awards"))
    if awards:
        _heading(doc, "Awards & Certifications")
        for award in awards:
            _bullet(doc, award)


def _render_project_details(doc: Document, item: dict[str, Any]) -> None:
    details = [
        ("Client", item.get("client")),
        ("Project", item.get("project")),
        ("Description", item.get("project_desc")),
    ]
    technologies = _string_list(item.get("technologies"))
    if technologies:
        details.append(("Technologies", ", ".join(technologies)))
    for label, value in details:
        text = str(value or "").strip()
        if not text:
            continue
        prefix = "" if label == "Description" else f"{label}: "
        _paragraph(doc, f"{prefix}{text}", bold=label == "Project")


def _render_ncs(doc: Document, data: dict[str, Any]) -> None:
    _ensure_bullet_style(doc)
    personal = data.get("personal_info") if isinstance(data.get("personal_info"), dict) else {}
    name = str(personal.get("name") or "").strip()
    experiences = data.get("experience") if isinstance(data.get("experience"), list) else []
    title = ""
    if experiences and isinstance(experiences[0], dict):
        title = str(experiences[0].get("title") or "").strip()

    heading = doc.add_paragraph()
    heading.add_run(title or name or "Resume")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        run.font.name = "Garamond"
        run.font.size = Pt(13)
        run.bold = True

    summary = _string_list(data.get("summary"))
    if summary:
        midpoint = max(1, len(summary) // 2)
        for chunk in (summary[:midpoint], summary[midpoint:]):
            if chunk:
                paragraph = doc.add_paragraph(" ".join(chunk))
                paragraph.paragraph_format.space_before = Pt(12)
                paragraph.paragraph_format.space_after = Pt(12)

    skills = _string_list(data.get("skills"))
    if skills:
        _heading(doc, "KEY SKILLS")
        for skill in skills:
            _bullet(doc, f"Experience with {skill}")
        _heading(doc, "TECHNICAL SKILLS")
        _paragraph(doc, ", ".join(skills))

    _render_kovan_like(doc, {**data, "summary": [], "skills": []})


def _string_list(value: object) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if str(item or "").strip()]
