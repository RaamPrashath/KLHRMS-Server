from __future__ import annotations

import csv
import html
import io
import re
from collections import defaultdict
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.modules.attendance_report.schema import (
    AttendanceReportExportEmployee,
    AttendanceReportExportRequest,
    AttendanceReportRow,
)

_HEADER_FILL = "1F4E78"
_TITLE_FILL = "D9EAF7"
_SUBTOTAL_FILL = "D9EAF7"
_THIN_BORDER = Side(style="thin", color="D9E2EC")
_GRID_BORDER = Border(left=_THIN_BORDER, right=_THIN_BORDER, top=_THIN_BORDER, bottom=_THIN_BORDER)
_REPORT_HEADERS = [
    "Date",
    "Day",
    "Project",
    "Description",
    "Hours",
    "Tickets on Hold / Pending Clarification",
]


def _employee_name(employee: AttendanceReportExportEmployee) -> str:
    return employee.name.strip() or employee.email or employee.id


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", " ", name).strip() or "Employee"
    base = cleaned[:31]
    candidate = base
    counter = 2
    while candidate in used:
        suffix = f" {counter}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate)
    return candidate


def _date_label(value: date) -> str:
    return value.isoformat()


def _day_label(value: date) -> str:
    return value.strftime("%A")


def _short_day_label(value: date) -> str:
    return value.strftime("%a")

def _is_weekend(value: date) -> bool:
    return value.weekday() >= 5


def _display_hours(hours: float | None, force8: bool) -> float | str:
    if hours is None:
        return ""
    normalized = min(float(hours), 8.0) if force8 else float(hours)
    return round(normalized, 2)


def _numeric_hours(hours: float | None, force8: bool) -> float:
    value = _display_hours(hours, force8)
    return float(value) if isinstance(value, int | float) else 0.0


def _row_description(row: AttendanceReportRow) -> str:
    return row.clockOutDescription or ""


def _project_label(row: AttendanceReportRow) -> str:
    if row.projectName and row.taskName:
        return f"{row.projectName} - {row.taskName}"
    return row.projectName or row.taskName or ""


def _rows_by_employee(rows: list[AttendanceReportRow]) -> dict[str, list[AttendanceReportRow]]:
    grouped: dict[str, list[AttendanceReportRow]] = defaultdict(list)
    for row in rows:
        grouped[row.employeeId].append(row)
    for employee_rows in grouped.values():
        employee_rows.sort(key=lambda item: item.date)
    return grouped


def _rows_by_client_employee(rows: list[AttendanceReportRow]) -> dict[str, dict[str, list[AttendanceReportRow]]]:
    """Group rows by (clientName, employeeId) for client-grouped timesheet."""
    grouped: dict[str, dict[str, list[AttendanceReportRow]]] = defaultdict(lambda: defaultdict(list))
    for row in rows:
        client = row.clientName or ""
        grouped[client][row.employeeId].append(row)
    for client_rows in grouped.values():
        for emp_rows in client_rows.values():
            emp_rows.sort(key=lambda item: item.date)
    return grouped


def _client_employee_matrix(
    rows: list[AttendanceReportRow],
    dates: list[date],
    employees: list[AttendanceReportExportEmployee],
) -> list[tuple[str, str, str, dict[str, dict[date, AttendanceReportRow]]]]:
    """
    Build (clientName, employeeId, employeeName, date_map) tuples
    ordered by client name then employee name.
    """
    client_employee_map: dict[str, dict[str, dict[date, AttendanceReportRow]]] = defaultdict(lambda: defaultdict(dict))
    employee_names: dict[str, str] = {}
    for row in rows:
        client = row.clientName or ""
        client_employee_map[client][row.employeeId][row.date] = row
        employee_names[row.employeeId] = row.employeeName

    result: list[tuple[str, str, str, dict[str, dict[date, AttendanceReportRow]]]] = []
    for client in sorted(client_employee_map.keys()):
        emp_map = client_employee_map[client]
        emp_ids = sorted(emp_map.keys(), key=lambda eid: employee_names.get(eid, "").lower())
        for eid in emp_ids:
            result.append((client, eid, employee_names.get(eid, ""), emp_map[eid]))
    return result


def _employees_for_export(payload: AttendanceReportExportRequest) -> list[AttendanceReportExportEmployee]:
    if payload.employees:
        return sorted(payload.employees, key=lambda item: _employee_name(item).lower())
    seen: dict[str, AttendanceReportExportEmployee] = {}
    for row in payload.rows:
        seen[row.employeeId] = AttendanceReportExportEmployee(
            id=row.employeeId,
            name=row.employeeName,
            email=row.employeeEmail,
        )
    return sorted(seen.values(), key=lambda item: _employee_name(item).lower())


def _apply_report_sheet_layout(ws, last_row: int) -> None:
    widths = [14, 14, 22, 64, 12, 34]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for row in ws.iter_rows(min_row=1, max_row=last_row, max_col=len(_REPORT_HEADERS)):
        for cell in row:
            cell.border = _GRID_BORDER


def generate_report_xlsx(payload: AttendanceReportExportRequest) -> bytes:
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    grouped = _rows_by_employee(payload.rows)
    used_sheet_names: set[str] = set()
    employees = _employees_for_export(payload)

    if not employees:
        employees = [AttendanceReportExportEmployee(id="empty", name="No selected employees")]

    for employee in employees:
        employee_rows = grouped.get(employee.id, [])
        ws = wb.create_sheet(_safe_sheet_name(_employee_name(employee), used_sheet_names))

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_REPORT_HEADERS))
        title_cell = ws.cell(row=1, column=1, value=f"Name: {_employee_name(employee)}")
        title_cell.font = Font(bold=True, size=12, color="1D1D1F")
        title_cell.fill = PatternFill("solid", fgColor=_TITLE_FILL)
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 24

        for col, header in enumerate(_REPORT_HEADERS, start=1):
            cell = ws.cell(row=2, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 24

        for row_index, row in enumerate(employee_rows, start=3):
            values = [
                _date_label(row.date),
                _day_label(row.date),
                _project_label(row),
                _row_description(row),
                _display_hours(row.totalHours, payload.force8),
                "",
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_index, column=col, value=value)
                cell.alignment = Alignment(
                    horizontal="left" if col in {3, 4, 6} else "center",
                    vertical="top",
                    wrap_text=col in {4, 6},
                )
                if col == 5 and value != "":
                    cell.number_format = "0.00"
                if _is_weekend(row.date):
                    cell.font = Font(color="FF0000")
                    cell.fill = PatternFill("solid", fgColor="FFCCCC")
            ws.row_dimensions[row_index].height = max(24, min(96, 18 + (_row_description(row).count("\n") + 1) * 12))

        total_row = len(employee_rows) + 3
        ws.cell(row=total_row, column=4, value="Total Hours:")
        ws.cell(row=total_row, column=4).font = Font(bold=True)
        if employee_rows:
            ws.cell(row=total_row, column=5, value=f"=SUM(E3:E{total_row - 1})")
        else:
            ws.cell(row=total_row, column=5, value="")
        ws.cell(row=total_row, column=5).font = Font(bold=True)
        ws.cell(row=total_row, column=5).number_format = "0.00"
        for col in range(1, len(_REPORT_HEADERS) + 1):
            cell = ws.cell(row=total_row, column=col)
            cell.fill = PatternFill("solid", fgColor=_SUBTOTAL_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        _apply_report_sheet_layout(ws, total_row)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_timesheet_maps(
    payload: AttendanceReportExportRequest,
) -> tuple[list[date], list[AttendanceReportExportEmployee], dict[str, dict[date, AttendanceReportRow]]]:
    dates = payload.dateColumns or sorted({row.date for row in payload.rows})
    employees = _employees_for_export(payload)
    grouped: dict[str, dict[date, AttendanceReportRow]] = defaultdict(dict)
    for row in payload.rows:
        grouped[row.employeeId][row.date] = row
    return dates, employees, grouped


def generate_timesheet_xlsx(payload: AttendanceReportExportRequest, db_rows: list[AttendanceReportRow] | None = None) -> bytes:
    # Use database rows if provided (spans the entire year), otherwise fall back to payload.rows
    rows_to_use = db_rows if db_rows is not None else payload.rows

    # Determine which months to generate based on dates in the payload
    dates_in_payload = payload.dateColumns or [row.date for row in payload.rows]
    if dates_in_payload:
        date_set = set(dates_in_payload)
        months_to_generate = set()
        for d in dates_in_payload:
            months_to_generate.add((d.year, d.month))
        months_sorted = sorted(months_to_generate, reverse=True)
    else:
        today = date.today()
        months_sorted = [(today.year, today.month)]

    # Group the rows by month key: (year, month)
    rows_by_month = defaultdict(list)
    for row in rows_to_use:
        rows_by_month[(row.date.year, row.date.month)].append(row)

    # Map employee_id to clientName from the entire year's data
    employee_client_map = {}
    for row in rows_to_use:
        if row.employeeId not in employee_client_map and row.clientName:
            employee_client_map[row.employeeId] = row.clientName

    # Default any remaining employees to ""
    for emp in payload.employees:
        if emp.id not in employee_client_map:
            employee_client_map[emp.id] = ""

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    import calendar

    # Define client group colors
    CLIENT_COLORS = ["FCE4D6", "F2F2F2", "FFF2CC", "E2EFDA", "D9E1F2"]

    # Find unique clients across all rows and sort them (putting empty/General at the bottom)
    unique_clients = sorted(list(set(employee_client_map.values())))
    sorted_clients = sorted(unique_clients, key=lambda c: (c == "" or c == "General", c.lower()))
    client_colors_map = {client: CLIENT_COLORS[idx % len(CLIENT_COLORS)] for idx, client in enumerate(sorted_clients)}

    # Fonts
    HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="000000")
    HEADER_WEEKEND_FONT = Font(name="Segoe UI", size=10, bold=True, color="C00000")
    EMPLOYEE_FONT = Font(name="Segoe UI", size=10, bold=True, color="000000")
    TOTAL_FONT = Font(name="Segoe UI", size=10, bold=True, color="000000")
    DATA_FONT = Font(name="Segoe UI", size=10, color="000000")

    CODE_FONT_L = Font(name="Segoe UI", size=10, bold=True, color="C00000")
    CODE_FONT_FH = Font(name="Segoe UI", size=10, bold=True, color="1F618D")
    CODE_FONT_CO = Font(name="Segoe UI", size=10, bold=True, color="B8860B")
    CODE_FONT_H = Font(name="Segoe UI", size=10, bold=True, color="5B2C6F")

    # Fills
    HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
    HEADER_WEEKEND_FILL = PatternFill("solid", fgColor="FADBD8")
    WEEKEND_FILL = PatternFill("solid", fgColor="FDF2F4")
    GREEN_FILL = PatternFill("solid", fgColor="E2EFDA")

    FILL_L = PatternFill("solid", fgColor="FADBD8")
    FILL_FH = PatternFill("solid", fgColor="D4E6F1")
    FILL_CO = PatternFill("solid", fgColor="FEF3CD")
    FILL_H = PatternFill("solid", fgColor="E8DAEF")

    # Borders
    _THIN_BORDER_SIDE = Side(style="thin", color="D9D9D9")
    _GRID_BORDER = Border(left=_THIN_BORDER_SIDE, right=_THIN_BORDER_SIDE, top=_THIN_BORDER_SIDE, bottom=_THIN_BORDER_SIDE)

    # Generate sheets only for months in the requested date range
    for year, m in months_sorted:
        month_name = calendar.month_name[m]
        sheet_title = f"{month_name} {year}"
        ws = wb.create_sheet(title=sheet_title)

        _, num_days = calendar.monthrange(year, m)
        all_month_dates = [date(year, m, d) for d in range(1, num_days + 1)]
        dates = [d for d in all_month_dates if d in date_set] if dates_in_payload else all_month_dates

        # Group employees by client name
        client_employees = defaultdict(list)
        for emp in payload.employees:
            client = employee_client_map[emp.id]
            client_employees[client].append(emp)

        # Build matrix of rows for this sheet
        month_rows = rows_by_month[(year, m)]
        matrix = []
        for client in sorted_clients:
            emps = sorted(client_employees[client], key=lambda e: e.name.lower())
            for emp in emps:
                emp_date_map = {}
                for row in month_rows:
                    if row.employeeId == emp.id:
                        emp_date_map[row.date] = row
                matrix.append((client, emp.id, emp.name, emp_date_map))

        total_cols = len(dates) + 4  # Client Name + Employee + Total + days + Total

        # Row 1: Title
        title_value = f"{month_name} {year} MONTHLY REPORT"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        title_cell = ws.cell(row=1, column=1, value=title_value)
        title_cell.font = Font(bold=True, size=14, color="1D1D1F")
        title_cell.fill = PatternFill("solid", fgColor="EEEEEE")
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 22

        # Row 2: Headers
        ws.cell(row=2, column=1, value="Client Name")
        ws.cell(row=2, column=2, value="Employee")
        ws.cell(row=2, column=3, value="Total")
        for offset, day in enumerate(dates, start=4):
            ws.cell(row=2, column=offset, value=day.day)
        ws.cell(row=2, column=total_cols, value="Total")
        ws.row_dimensions[2].height = 24

        # Row 3: Days
        ws.cell(row=3, column=1, value="Day:")
        ws.cell(row=3, column=2, value="")
        ws.cell(row=3, column=3, value="Hours")
        for offset, day in enumerate(dates, start=4):
            ws.cell(row=3, column=offset, value=_short_day_label(day))
        ws.cell(row=3, column=total_cols, value="Hours")
        ws.row_dimensions[3].height = 24

        # Weekend column indices (1-indexed)
        weekend_cols = {4 + idx for idx, day in enumerate(dates) if _is_weekend(day)}

        # Style headers
        for col in range(1, total_cols + 1):
            for r in (2, 3):
                cell = ws.cell(row=r, column=col)
                cell.border = _GRID_BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if col in weekend_cols:
                    cell.font = HEADER_WEEKEND_FONT
                    cell.fill = HEADER_WEEKEND_FILL
                else:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL

        # Populate data rows
        row_index = 4
        for client, emp_id, emp_name, date_map in matrix:
            client_label = client or "General"
            ws.cell(row=row_index, column=1, value=client_label)
            ws.cell(row=row_index, column=2, value=emp_name or emp_id)

            client_color = client_colors_map.get(client or "", "F2F2F2")
            client_fill = PatternFill("solid", fgColor=client_color)

            # Formulas for Totals
            first_day_col = get_column_letter(4)
            last_day_col = get_column_letter(3 + len(dates))

            # Column C: Total Hours formula
            total_cell_c = ws.cell(row=row_index, column=3, value=f"=SUM({first_day_col}{row_index}:{last_day_col}{row_index})")
            total_cell_c.number_format = "0"
            total_cell_c.font = TOTAL_FONT
            total_cell_c.fill = client_fill
            total_cell_c.alignment = Alignment(horizontal="center", vertical="center")
            total_cell_c.border = _GRID_BORDER

            # Last Column: Total Hours formula
            total_cell_last = ws.cell(row=row_index, column=total_cols, value=f"=C{row_index}")
            total_cell_last.number_format = "0"
            total_cell_last.font = TOTAL_FONT
            total_cell_last.fill = client_fill
            total_cell_last.alignment = Alignment(horizontal="center", vertical="center")
            total_cell_last.border = _GRID_BORDER

            # Client Name & Employee styling
            for col_idx in (1, 2):
                cell = ws.cell(row=row_index, column=col_idx)
                cell.border = _GRID_BORDER
                cell.fill = client_fill
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = EMPLOYEE_FONT if col_idx == 2 else DATA_FONT

            # Day cells
            for offset, day in enumerate(dates, start=4):
                report_row = date_map.get(day)
                cell = ws.cell(row=row_index, column=offset)
                cell.border = _GRID_BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center")

                # Default values and styles
                val = ""
                cell_font = DATA_FONT
                cell_fill = client_fill

                is_weekend_day = _is_weekend(day)
                if is_weekend_day:
                    cell_fill = WEEKEND_FILL

                if report_row:
                    leave_name = report_row.leaveTypeName
                    entry_type = report_row.entryType
                    hours = report_row.totalHours

                    if leave_name or entry_type == "LEAVE":
                        val = "L"
                        cell_font = CODE_FONT_L
                        cell_fill = FILL_L
                    elif entry_type == "FLOATING_HOLIDAY":
                        val = "FH"
                        cell_font = CODE_FONT_FH
                        cell_fill = FILL_FH
                    elif entry_type == "COMP_OFF":
                        val = "CO"
                        cell_font = CODE_FONT_CO
                        cell_fill = FILL_CO
                    elif entry_type == "HOLIDAY":
                        val = "H"
                        cell_font = CODE_FONT_H
                        cell_fill = FILL_H
                    elif hours is not None and hours > 0:
                        norm_hours = min(hours, 8.0) if payload.force8 else hours
                        val = int(norm_hours) if norm_hours.is_integer() else round(norm_hours, 1)

                        if not is_weekend_day:
                            cell_fill = GREEN_FILL

                cell.value = val
                cell.font = cell_font
                cell.fill = cell_fill

            ws.row_dimensions[row_index].height = 20
            row_index += 1

        # Freeze panes & column dimensions
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 10
        for col in range(4, total_cols):
            ws.column_dimensions[get_column_letter(col)].width = 6
        ws.column_dimensions[get_column_letter(total_cols)].width = 10

        ws.freeze_panes = "D4"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _paragraph(value: str, style: ParagraphStyle) -> Paragraph:
    escaped = html.escape(value or "").replace("\n", "<br/>")
    return Paragraph(escaped or "&nbsp;", style)


def generate_report_pdf(payload: AttendanceReportExportRequest) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("ReportCell", parent=styles["Normal"], fontSize=8, leading=10)
    title_style = ParagraphStyle("ReportTitle", parent=styles["Heading2"], fontSize=12, leading=14)

    grouped = _rows_by_employee(payload.rows)
    story = []
    employees = _employees_for_export(payload)
    for employee_index, employee in enumerate(employees):
        if employee_index:
            story.append(PageBreak())
        story.append(Paragraph(f"Name: {_employee_name(employee)}", title_style))
        if payload.periodLabel:
            story.append(Paragraph(html.escape(payload.periodLabel), normal))
        story.append(Spacer(1, 4 * mm))

        table_data: list[list[object]] = [_REPORT_HEADERS]
        weekend_rows: list[int] = []
        for idx, row in enumerate(grouped.get(employee.id, []), start=1):
            if _is_weekend(row.date):
                weekend_rows.append(idx)
            table_data.append(
                [
                    _date_label(row.date),
                    _day_label(row.date),
                    _project_label(row),
                    _paragraph(_row_description(row), normal),
                    _display_hours(row.totalHours, payload.force8),
                    "",
                ]
            )
        total = sum(_numeric_hours(row.totalHours, payload.force8) for row in grouped.get(employee.id, []))
        table_data.append(["", "", "", "Total Hours:", f"{total:.2f}" if total else "", ""])

        style_cmds: list[object] = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_HEADER_FILL}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E2EC")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ALIGN", (0, 0), (2, -1), "CENTER"),
            ("ALIGN", (4, 1), (4, -1), "CENTER"),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor(f"#{_SUBTOTAL_FILL}")),
            ("FONTNAME", (3, -1), (4, -1), "Helvetica-Bold"),
        ]
        for r in weekend_rows:
            style_cmds.append(("TEXTCOLOR", (0, r), (-1, r), colors.red))
        table = Table(
            table_data,
            colWidths=[24 * mm, 24 * mm, 38 * mm, 108 * mm, 18 * mm, 54 * mm],
            repeatRows=1,
        )
        table.setStyle(TableStyle(style_cmds))
        story.append(table)

    if not story:
        story.append(Paragraph("No selected employees", title_style))
    doc.build(story)
    return buffer.getvalue()


def generate_timesheet_pdf(payload: AttendanceReportExportRequest) -> bytes:
    dates = payload.dateColumns or sorted({row.date for row in payload.rows})
    matrix = _client_employee_matrix(payload.rows, dates, payload.employees)
    page_size = landscape(A3) if len(dates) > 10 else landscape(A4)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=8 * mm,
        rightMargin=8 * mm,
        topMargin=10 * mm,
        bottomMargin=8 * mm,
    )
    styles = getSampleStyleSheet()
    story = [Paragraph(payload.title, styles["Title"])]
    if payload.periodLabel:
        story.append(Paragraph(html.escape(payload.periodLabel), styles["Normal"]))
    story.append(Spacer(1, 4 * mm))

    # Build table data: Client Name | Employee | Total Hours | dates... | Total Hours
    table_data: list[list[object]] = [
        ["Client Name", "Employee", "Total Hours", *[str(day.day) for day in dates], "Total"],
        ["Day:", "", "", *[_short_day_label(day) for day in dates], "Hours"],
    ]
    weekend_cols: list[int] = []
    for idx, day in enumerate(dates):
        if _is_weekend(day):
            weekend_cols.append(3 + idx)  # offset by Client+Employee+Total

    for client, emp_id, emp_name, date_map in matrix:
        values = []
        total = 0.0
        for day in dates:
            report_row = date_map.get(day)
            hours = _numeric_hours(report_row.totalHours, payload.force8) if report_row else 0.0
            total += hours
            values.append(f"{hours:.2f}" if hours else "")
        table_data.append([client or "—", emp_name or emp_id, f"{total:.2f}" if total else "", *values, f"{total:.2f}" if total else ""])

    usable_width = page_size[0] - 16 * mm
    client_width = 28 * mm
    employee_width = 42 * mm
    total_width = 18 * mm
    day_width = max((usable_width - client_width - employee_width - total_width - total_width) / max(len(dates), 1), 7 * mm)
    col_widths = [client_width, employee_width, total_width, *([day_width] * len(dates)), total_width]

    style_cmds: list[object] = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_HEADER_FILL}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor(f"#{_TITLE_FILL}")),
        ("FONTNAME", (0, 0), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9E2EC")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("ALIGN", (0, 2), (1, -1), "LEFT"),
        ("FONTNAME", (0, 2), (1, -1), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for c in weekend_cols:
        style_cmds.append(("TEXTCOLOR", (c, 0), (c, -1), colors.red))
    table = Table(table_data, colWidths=col_widths, repeatRows=2)
    table.setStyle(TableStyle(style_cmds))
    story.append(table)
    doc.build(story)
    return buffer.getvalue()


def generate_report_csv(payload: AttendanceReportExportRequest) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Employee", "Date", "Day", "Project", "Description", "Hours"])

    grouped = _rows_by_employee(payload.rows)
    employees = _employees_for_export(payload)

    if not employees:
        employees = [AttendanceReportExportEmployee(id="empty", name="No selected employees")]

    for employee in employees:
        employee_rows = grouped.get(employee.id, [])
        for row in employee_rows:
            writer.writerow([
                _employee_name(employee),
                _date_label(row.date),
                _day_label(row.date),
                _project_label(row),
                _row_description(row),
                _display_hours(row.totalHours, payload.force8),
            ])

    return output.getvalue().encode("utf-8-sig")


def generate_timesheet_csv(payload: AttendanceReportExportRequest) -> bytes:
    dates = payload.dateColumns or sorted({row.date for row in payload.rows})
    matrix = _client_employee_matrix(payload.rows, dates, payload.employees)
    output = io.StringIO()
    writer = csv.writer(output)

    headers = ["Client Name", "Employee", "Total Hours", *[str(day.day) for day in dates], "Total Hours"]
    writer.writerow(headers)
    writer.writerow(["Day:", "", "", *[_short_day_label(day) for day in dates], "Hours"])

    for client, emp_id, emp_name, date_map in matrix:
        row_values = [client or "—", emp_name or emp_id]
        total = 0.0
        for day in dates:
            report_row = date_map.get(day)
            hours = _numeric_hours(report_row.totalHours, payload.force8) if report_row else 0.0
            total += hours
            row_values.append(f"{hours:.2f}" if hours else "")
        row_values.insert(2, f"{total:.2f}" if total else "")
        row_values.append(f"{total:.2f}" if total else "")
        writer.writerow(row_values)

    return output.getvalue().encode("utf-8-sig")


def generate_attendance_report_export(payload: AttendanceReportExportRequest, db_rows: list[AttendanceReportRow] | None = None) -> bytes:
    if payload.mode == "report":
        if payload.format == "csv":
            return generate_report_csv(payload)
        if payload.format == "xlsx":
            return generate_report_xlsx(payload)
        return generate_report_pdf(payload)
    if payload.format == "csv":
        return generate_timesheet_csv(payload)
    if payload.format == "xlsx":
        return generate_timesheet_xlsx(payload, db_rows=db_rows)
    return generate_timesheet_pdf(payload)
